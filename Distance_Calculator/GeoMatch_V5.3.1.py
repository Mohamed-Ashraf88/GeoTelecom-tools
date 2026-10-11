"""GeoMatch — Spatial Matching & Proximity Analysis.
Version 5.3.1: integrated Help & Support menu, product details, quick guide and contact windows.
Author: Mohamed Ashraf.
Python 3.12 x64. Install: pip install PySide6-Essentials==6.11.2 pandas==3.0.1 numpy==2.5.3 pyproj==3.8.0 openpyxl==3.1.5 scikit-learn==1.9.1 requests==2.34.2 xlrd==2.0.2 geopandas==1.1.4 pyogrio==0.13.0 shapely==2.1.2 scipy==1.18.1 rasterio==1.5.1
Run: python GeoMatch_V5.3.1.py
Check startup: python GeoMatch_V5.3.1.py --smoke
Qt imports are confined to the GUI parent; calculations use spawned processes.
"""

from __future__ import annotations

import json
import logging
import math
import multiprocessing as mp
import os
import queue
import re
import sys
import tempfile
import time
import traceback
from dataclasses import asdict, dataclass, replace
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

TITLE = "GeoMatch — Spatial Matching & Proximity Analysis"
VERSION = "5.3.1"
BUILD_NUMBER = "20260928.1"
BUILD_DATE = "28 September 2026"
SUPPORT_EMAIL = "Mohamed--Ashraf@outlook.com"
SUPPORT_LINKEDIN = "https://www.linkedin.com/in/mohamed---ashraf/"
MATCH_MODES = ("Nearest N", "All within radius", "Nearest eligible", "Mutual nearest")
AUTHOR = "Mohamed Ashraf"
FIELDS = (
    "Site ID",
    "Latitude",
    "Longitude",
    "Antenna AGL (m)",
    "Terrain elevation (m)",
    "Absolute antenna elevation (m)",
    "Match attribute",
)
HEIGHT_FIELDS = FIELDS[3:6]
RESULT_COLUMNS = (
    "Source row",
    "Source site",
    "Target row",
    "Target site",
    "Rank",
    "Ground distance (m)",
    "Distance (m)",
    "Bearing (deg)",
    "Distance mode",
    "Match status",
    "Terrain status",
    "Source latitude",
    "Source longitude",
    "Target latitude",
    "Target longitude",
)
METHODS = ("WGS84 geodesic", "Haversine (spherical)")
MODES = ("2D ground", "3D approximate slant")


class Cancelled(Exception):
    pass


def checkpoint(cancel):
    if cancel is not None and cancel.is_set():
        raise Cancelled("Run stopped")


def normalize_header(value):
    return " ".join(
        str(value).strip().lower().replace("_", " ").replace("-", " ").split()
    )


ALIASES = {
    "Match attribute": ("match attribute", "category", "asset type"),
    "Site ID": (
        "point name",
        "pointname",
        "site id",
        "site name",
        "sitename",
        "site",
        "id",
        "name",
        "station",
        "cell",
        "node",
    ),
    "Latitude": ("latitude", "lat", "lat deg", "y", "northing"),
    "Longitude": ("longitude", "lon", "long", "lng", "lon deg", "x", "easting"),
    "Antenna AGL (m)": (
        "antenna agl (m)",
        "antenna agl",
        "antenna height",
        "ant height",
        "height agl",
        "agl",
    ),
    "Terrain elevation (m)": (
        "terrain elevation (m)",
        "terrain elevation",
        "ground elevation",
    ),
    "Absolute antenna elevation (m)": (
        "absolute antenna elevation (m)",
        "absolute antenna elevation",
        "antenna elevation amsl",
    ),
}


def suggest_mapping(headers):
    """Exact normalized aliases only. Ambiguities require a user choice."""
    mapping, conflicts = {}, {}
    for field, aliases in ALIASES.items():
        candidates = [h for h in headers if normalize_header(h) in aliases]
        mapping[field] = candidates[0] if len(candidates) == 1 else ""
        if len(candidates) > 1:
            conflicts[field] = candidates
    return mapping, conflicts


def validate_mapping(headers, mapping):
    for field in FIELDS[:3]:
        if not mapping.get(field):
            raise ValueError(f"Choose a column for {field}.")
    selected = [v for key, v in mapping.items() if v and key != "Match attribute"]
    if len(set(selected)) != len(selected):
        raise ValueError("Each input column may be mapped only once.")
    missing = [v for v in mapping.values() if v and v not in headers]
    if missing:
        raise ValueError(f"Mapped columns are no longer present: {', '.join(missing)}")


@dataclass(frozen=True)
class FileSpec:
    path: str
    mapping: tuple
    signature: tuple
    sheet: str | None = None
    rows: int = 0
    crs: str = "EPSG:4326"
    sha256: str = ""
    origin: str = "file"
    edits: tuple = ()
    excluded: tuple = ()


@dataclass(frozen=True)
class RunSpec:
    source: FileSpec
    target: FileSpec
    method: str = METHODS[0]
    mode: str = MODES[0]
    top_n: int = 1
    radius_m: float = 0.0
    ignore_same_location: bool = False
    flag_outliers: bool = False
    terrain: bool = False
    match_mode: str = "Nearest N"
    match_attributes: bool = False

    def validate(self):
        if self.match_mode not in MATCH_MODES:
            raise ValueError("Choose a supported matching mode.")
        if self.match_mode == "All within radius" and self.radius_m <= 0:
            raise ValueError(
                "All-within-radius matching requires a positive distance limit; turn Unlimited off."
            )
        if self.method not in METHODS or self.mode not in MODES:
            raise ValueError("Choose a supported calculation method and distance mode.")
        if not isinstance(self.top_n, int) or not 1 <= self.top_n <= 10:
            raise ValueError("Targets per source must be a whole number from 1 to 10.")
        if not math.isfinite(self.radius_m) or self.radius_m < 0:
            raise ValueError("Maximum distance must be a finite, nonnegative number.")


def signature(path):
    stat = Path(path).stat()
    return stat.st_size, stat.st_mtime_ns


def radius_in_metres(value, unit, unlimited):
    if unlimited:
        return 0.0
    if unit not in ("m", "km"):
        raise ValueError("Choose metres (m) or kilometres (km).")
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise ValueError(
            "Enter a numeric maximum distance, or select Unlimited."
        ) from None
    metres = number * (1000 if unit == "km" else 1)
    if not math.isfinite(metres) or metres <= 0:
        raise ValueError(
            "Maximum distance must be greater than zero and finite, or select Unlimited."
        )
    return metres


def parse_coordinate_pair(text):
    text = text.strip().replace("−", "-")
    parts = (
        re.split(r"[,;\t\n]+", text) if re.search(r"[,;\t\n]", text) else text.split()
    )
    if len(parts) != 2:
        raise ValueError(
            "Paste latitude, longitude. Separate DMS coordinates with a comma or semicolon."
        )
    lat, lon = parse_angle(parts[0], "latitude"), parse_angle(parts[1], "longitude")
    if (
        not math.isfinite(lat)
        or not math.isfinite(lon)
        or abs(lat) > 90
        or abs(lon) > 180
    ):
        raise ValueError(
            "Latitude must be within −90…90 and longitude within −180…180; both must be finite."
        )
    return lat, lon


def read_input(path, cancel=None, progress=None, sheet=None):
    """Read strings without dropping leading zeroes in site identifiers."""
    checkpoint(cancel)
    ext = Path(path).suffix.lower()
    if ext in (".csv", ".txt"):
        parts = []
        with pd.read_csv(
            path,
            sep=None,
            engine="python",
            dtype=str,
            keep_default_na=False,
            chunksize=10000,
        ) as reader:
            for chunk in reader:
                checkpoint(cancel)
                parts.append(chunk)
                if progress:
                    progress(
                        f"Reading {Path(path).name}: {sum(len(p) for p in parts):,} rows"
                    )
        frame = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()
    elif ext in (".xlsx", ".xls"):
        frame = pd.read_excel(
            path,
            dtype=str,
            keep_default_na=False,
            engine="openpyxl" if ext == ".xlsx" else "xlrd",
            sheet_name=sheet if sheet is not None else 0,
        )
    elif ext == ".geojson":
        with open(path, encoding="utf-8-sig") as stream:
            data = json.load(stream)
        if data.get("type") != "FeatureCollection":
            raise ValueError("GeoJSON must contain a FeatureCollection of points.")
        if data.get("crs"):
            raise ValueError(
                "Legacy GeoJSON CRS declarations are unsupported. Re-export as WGS84 RFC 7946 GeoJSON."
            )
        records = []
        for feature in data.get("features", []):
            checkpoint(cancel)
            props = dict(feature.get("properties") or {})
            geom = feature.get("geometry") or {}
            coords = geom.get("coordinates") or []
            # Keep non-point features as rejected rows, never as guessed coordinates.
            props["Latitude"] = (
                coords[1] if geom.get("type") == "Point" and len(coords) >= 2 else ""
            )
            props["Longitude"] = (
                coords[0] if geom.get("type") == "Point" and len(coords) >= 2 else ""
            )
            records.append(props)
        frame = pd.DataFrame(records)
    else:
        raise ValueError("Choose an XLSX, XLS, CSV, TXT, or GeoJSON file.")
    checkpoint(cancel)
    frame.columns = [str(c) for c in frame.columns]
    if not frame.columns.is_unique:
        raise ValueError(
            "Input contains duplicate column names. Rename them before importing."
        )
    if frame.empty:
        raise ValueError(f"{Path(path).name} contains no data rows.")
    return frame


def validate_frame(raw, mapping, spec, cancel=None, file=None):
    validate_mapping(raw.columns, mapping)
    checkpoint(cancel)
    data = pd.DataFrame(index=raw.index)
    data["Input row"] = np.arange(len(raw)) + 2
    for field in FIELDS:
        if mapping.get(field):
            data[field] = raw[mapping[field]]
    data["Site ID"] = data["Site ID"].fillna("").astype(str).str.strip()
    reasons = pd.Series("", index=data.index, dtype=object)

    def reject(mask, reason):
        mask = pd.Series(mask, index=data.index) & reasons.eq("")
        reasons.loc[mask] = reason

    reject(data["Site ID"].eq(""), "Missing site ID")
    (
        latitudes,
        longitudes,
        coordinate_reasons,
        coordinate_warnings,
        coordinate_details,
    ) = normalize_coordinates(raw, mapping, file or spec.source, cancel)
    data["Latitude"], data["Longitude"] = latitudes, longitudes
    for index, problem in zip(data.index, coordinate_reasons):
        if problem and not reasons.loc[index]:
            reasons.loc[index] = problem
    if spec.match_attributes:
        if not mapping.get("Match attribute"):
            raise ValueError(
                "Map a Match attribute column in both inputs to enable attribute matching."
            )
        data["Match attribute"] = (
            data["Match attribute"].fillna("").astype(str).str.strip()
        )
        reject(
            data["Match attribute"].eq(""),
            "Missing match attribute (blank values never match)",
        )
    for col, bound in (("Latitude", 90), ("Longitude", 180)):
        data[col] = pd.to_numeric(data[col], errors="coerce")
        reject(
            ~np.isfinite(data[col]) | ~data[col].between(-bound, bound),
            f"Invalid {col.lower()} (WGS84 degrees required)",
        )
    for col in HEIGHT_FIELDS:
        if col in data:
            data[col] = pd.to_numeric(data[col], errors="coerce")
    if spec.mode == MODES[1]:
        absolute = data.get(
            "Absolute antenna elevation (m)", pd.Series(np.nan, index=data.index)
        )
        agl = data.get("Antenna AGL (m)", pd.Series(np.nan, index=data.index))
        terrain = data.get("Terrain elevation (m)", pd.Series(np.nan, index=data.index))
        valid_agl = np.isfinite(agl) & (agl >= 0)
        data["Resolved absolute (m)"] = absolute.where(
            np.isfinite(absolute), (terrain + agl).where(valid_agl)
        )
        reject(
            ~np.isfinite(data["Resolved absolute (m)"]),
            "3D requires absolute antenna elevation, or terrain elevation + nonnegative AGL",
        )
    if spec.terrain:
        agl = data.get("Antenna AGL (m)", pd.Series(np.nan, index=data.index))
        reject(
            ~np.isfinite(agl) | (agl < 0),
            "Terrain check requires a finite, nonnegative antenna AGL height",
        )
    valid_indices = data.index[reasons.eq("")]
    duplicate_indices = data.loc[valid_indices].index[
        data.loc[valid_indices, "Site ID"].duplicated(keep="first")
    ]
    reasons.loc[duplicate_indices] = "Duplicate site ID (first valid row retained)"
    valid = data.loc[reasons.eq("")].copy()
    valid["Flag"] = ""
    checkpoint(cancel)
    if spec.flag_outliers and len(valid) > 10:
        from sklearn.ensemble import IsolationForest

        scores = IsolationForest(contamination=0.05, random_state=42).fit_predict(
            valid[["Latitude", "Longitude"]]
        )
        valid.loc[scores == -1, "Flag"] = (
            "Statistical outlier — retained unless excluded during review"
        )
    checkpoint(cancel)
    audit = raw.copy()
    # Prefix metadata to preserve every original input column.
    for col in ("GeoMatch input row", "GeoMatch reason", "GeoMatch disposition"):
        while col in audit.columns:
            audit = audit.rename(columns={col: "Original " + col})
    audit.insert(0, "GeoMatch input row", np.arange(len(raw)) + 2)
    audit["GeoMatch reason"] = reasons
    audit["GeoMatch disposition"] = np.where(reasons.eq(""), "Retained", "Rejected")
    flagged = valid.index[valid["Flag"].ne("")]
    audit.loc[flagged, "GeoMatch reason"] = valid.loc[flagged, "Flag"]
    audit.loc[flagged, "GeoMatch disposition"] = "Flagged — retained"
    # Preserve raw user columns if they happen to use our audit names.
    for name, values in (
        ("GeoMatch WGS84 latitude", latitudes),
        ("GeoMatch WGS84 longitude", longitudes),
        ("GeoMatch coordinate warning", coordinate_warnings),
    ):
        while name in audit:
            audit = audit.rename(columns={name: "Original " + name})
        audit[name] = values
    valid, audit = valid.reset_index(drop=True), audit.reset_index(drop=True)
    valid.attrs["coordinate_operation"] = coordinate_details
    audit.attrs["coordinate_operation"] = coordinate_details
    return valid, audit


def distances(lat, lon, tlats, tlons, method):
    if method == METHODS[0]:
        from pyproj import Geod

        fwd, _, ground = Geod(ellps="WGS84").inv(
            np.full(len(tlons), lon), np.full(len(tlats), lat), tlons, tlats
        )
        return np.asarray(ground), (np.asarray(fwd) + 360) % 360
    if method != METHODS[1]:
        raise ValueError("Unknown distance method")
    a1, a2 = np.radians(lat), np.radians(tlats)
    dl = np.radians(tlons - lon)
    a = np.clip(
        np.sin((a2 - a1) / 2) ** 2 + np.cos(a1) * np.cos(a2) * np.sin(dl / 2) ** 2, 0, 1
    )
    ground = 6371000.0 * 2 * np.arctan2(np.sqrt(a), np.sqrt(1 - a))
    bearing = (
        np.degrees(
            np.arctan2(
                np.sin(dl) * np.cos(a2),
                np.cos(a1) * np.sin(a2) - np.sin(a1) * np.cos(a2) * np.cos(dl),
            )
        )
        % 360
    )
    return ground, bearing


def terrain_check(s, t, cancel=None):
    """Retains the existing optional 5 GHz / 60% Fresnel / K=4/3 check."""
    checkpoint(cancel)
    from pyproj import Geod
    import requests

    geod = Geod(ellps="WGS84")
    _, _, ground = geod.inv(
        s["Longitude"], s["Latitude"], t["Longitude"], t["Latitude"]
    )
    if ground < 10:
        return "Not evaluated: ground separation under 10 m"
    count = max(3, min(int(ground / 10), 500))
    pts = (
        [(s["Longitude"], s["Latitude"])]
        + geod.npts(
            s["Longitude"], s["Latitude"], t["Longitude"], t["Latitude"], count - 2
        )
        + [(t["Longitude"], t["Latitude"])]
    )
    try:
        response = requests.post(
            "https://api.open-elevation.com/api/v1/lookup",
            json={
                "locations": [{"latitude": lat, "longitude": lon} for lon, lat in pts]
            },
            timeout=(5, 15),
        )
        response.raise_for_status()
        heights = np.asarray(
            [p["elevation"] for p in response.json()["results"]], dtype=float
        )
        if len(heights) != len(pts) or not np.isfinite(heights).all():
            return "Unavailable: invalid terrain response"
    except (requests.RequestException, ValueError, KeyError, TypeError):
        return "Unavailable: terrain service error"
    checkpoint(cancel)
    start = heights[0] + s["Antenna AGL (m)"]
    end = heights[-1] + t["Antenna AGL (m)"]
    for i in range(1, len(pts) - 1):
        checkpoint(cancel)
        _, _, d1 = geod.inv(*pts[0], *pts[i])
        d2 = ground - d1
        fresnel = 17.32 * math.sqrt(
            max(0, (d1 / 1000) * (d2 / 1000) / (5 * ground / 1000))
        )
        bulge = d1 * d2 / (2 * (4 / 3) * 6371000)
        if start + d1 / ground * (end - start) - heights[i] - bulge < 0.6 * fresnel:
            return "Blocked (sampled terrain)"
    return "Clear (sampled terrain only)"


def input_fingerprint(path, cancel=None):
    import hashlib

    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            checkpoint(cancel)
            digest.update(block)
    return digest.hexdigest()


def horizontal_crs(value):
    from pyproj import network

    network.set_network_enabled(False)
    from pyproj import CRS

    try:
        crs = CRS.from_user_input(value)
    except Exception as error:
        raise ValueError(
            "Choose a valid horizontal CRS, for example EPSG:4326 or EPSG:32636."
        ) from error
    if not (crs.is_projected or crs.is_geographic) or len(crs.axis_info) != 2:
        raise ValueError(
            "Use a two-dimensional geographic or projected CRS. Heights are handled separately."
        )
    if crs.is_geographic and any(
        "degree" not in axis.unit_name.lower() for axis in crs.axis_info
    ):
        raise ValueError(
            "Geographic input must use degree-based axes; projected input uses its CRS units."
        )
    return crs


def search_crs(query):
    from pyproj.database import query_crs_info
    from pyproj.enums import PJType

    tokens = query.lower().replace("epsg:", "").split()
    if not tokens:
        raise ValueError(
            "Enter an EPSG code or place/projection name, such as 32636 or WGS 84 UTM 36N."
        )
    return [
        (f"EPSG:{entry.code}", entry.name)
        for entry in query_crs_info(
            auth_name="EPSG", pj_types=[PJType.GEOGRAPHIC_2D_CRS, PJType.PROJECTED_CRS]
        )
        if all(t in (entry.code + " " + entry.name).lower() for t in tokens)
    ][:100]


def parse_angle(value, axis):
    """Decimal degrees, signed DMS or hemispheric DMS; no silent axis swapping."""
    text = str(value).strip().upper().replace("−", "-")
    try:
        number = float(text)
        if not math.isfinite(number):
            raise ValueError("Coordinate is not finite")
        return number
    except ValueError:
        pass
    hemisphere = re.findall(r"[NSEW]", text)
    if len(hemisphere) > 1 or (
        hemisphere and hemisphere[0] not in ("NS" if axis == "latitude" else "EW")
    ):
        raise ValueError(f"Hemisphere does not match {axis}")
    if hemisphere and not (
        text.startswith(hemisphere[0]) or text.endswith(hemisphere[0])
    ):
        raise ValueError("Hemisphere must be at the start or end")
    body = re.sub(r"[NSEW]", "", text).strip()
    body = re.sub(r"[°º′'″\"’:dms]", " ", body, flags=re.I)
    parts = body.split()
    if not 1 <= len(parts) <= 3:
        raise ValueError("Use decimal degrees or degrees minutes seconds")
    if any(not re.fullmatch(r"[+-]?\d+(?:\.\d*)?|[+-]?\.\d+", p) for p in parts):
        raise ValueError("Invalid coordinate notation")
    numbers = list(map(float, parts))
    if len(numbers) > 1 and (not 0 <= numbers[1] < 60 or numbers[0] != int(numbers[0])):
        raise ValueError("DMS minutes must be 0–59; degrees must be whole")
    if len(numbers) > 2 and (not 0 <= numbers[2] < 60 or numbers[1] != int(numbers[1])):
        raise ValueError("DMS seconds must be 0–59; minutes must be whole")
    negative = parts[0].startswith("-")
    if hemisphere:
        h_negative = hemisphere[0] in "SW"
        if negative and not h_negative or parts[0].startswith("+") and h_negative:
            raise ValueError("Sign conflicts with hemisphere")
        negative = h_negative
    magnitude = (
        abs(numbers[0])
        + (numbers[1] / 60 if len(numbers) > 1 else 0)
        + (numbers[2] / 3600 if len(numbers) > 2 else 0)
    )
    return -magnitude if negative else magnitude


def normalize_coordinates(raw, mapping, file, cancel=None):
    from pyproj import Transformer, network

    # No grid downloads or approximate datum substitutions hidden from the user.
    network.set_network_enabled(False)
    crs = horizontal_crs(file.crs if file else "EPSG:4326")
    transform = Transformer.from_crs(
        crs, "EPSG:4326", always_xy=True, allow_ballpark=False, only_best=True
    )
    latitudes, longitudes, reasons, warnings = [], [], [], []
    xcol, ycol = mapping["Longitude"], mapping["Latitude"]
    area = crs.area_of_use
    for xraw, yraw in zip(raw[xcol], raw[ycol]):
        checkpoint(cancel)
        try:
            if crs.is_geographic:
                x, y = parse_angle(xraw, "longitude"), parse_angle(yraw, "latitude")
                if abs(y) > 90 or abs(x) > 180:
                    hint = (
                        " Check the X/Y column mapping; coordinates may be swapped."
                        if abs(y) <= 180 and abs(x) <= 90
                        else ""
                    )
                    raise ValueError("Invalid latitude/longitude range." + hint)
            else:
                x, y = float(xraw), float(yraw)
                if not math.isfinite(x) or not math.isfinite(y):
                    raise ValueError("Projected coordinates must be finite numbers")
            lon, lat = transform.transform(x, y, errcheck=True)
            if (
                not math.isfinite(lat)
                or not math.isfinite(lon)
                or abs(lat) > 90
                or abs(lon) > 180
            ):
                raise ValueError("Transformed coordinates are outside WGS84 limits")
            warning = []
            if abs(lat) < 1e-12 and abs(lon) < 1e-12:
                warning.append("At 0°, 0°: confirm this is intentional")
            if area:
                within_lon = (
                    area.west <= lon <= area.east
                    if area.west <= area.east
                    else lon >= area.west or lon <= area.east
                )
                if not within_lon or not area.south <= lat <= area.north:
                    warning.append(
                        "Outside selected CRS area of use: confirm CRS and location"
                    )
            latitudes.append(lat)
            longitudes.append(lon)
            reasons.append("")
            warnings.append("; ".join(warning))
        except Exception as error:
            latitudes.append(np.nan)
            longitudes.append(np.nan)
            reasons.append(f"Coordinate check: {error}")
            warnings.append("")
    details = {
        "input_crs": crs.to_string(),
        "input_crs_wkt": crs.to_wkt(),
        "output_crs": "EPSG:4326",
        "axis_mapping": {
            "x_easting_or_longitude": xcol,
            "y_northing_or_latitude": ycol,
        },
        "input_axis_units": [a.unit_name for a in crs.axis_info],
        "operation": transform.description,
        "accuracy_metres": transform.accuracy,
        "height_transformation": "none; user-supplied common vertical reference required",
    }
    try:
        operation = transform.get_last_used_operation()
        details.update(
            operation=operation.description,
            operation_definition=operation.definition,
            accuracy_metres=operation.accuracy,
        )
    except Exception:
        pass
    return np.array(latitudes), np.array(longitudes), reasons, warnings, details


def parse_pasted_table(text):
    import io

    if not text.strip():
        raise ValueError(
            "Clipboard is empty. Copy a table including its column headings from Excel."
        )
    frame = pd.read_csv(io.StringIO(text), sep="\t", dtype=str, keep_default_na=False)
    if frame.empty or len(frame.columns) < 3:
        raise ValueError(
            "Paste a tab-separated table with a header and at least ID, X and Y columns."
        )
    return frame


def calculate_reference(source, target, spec, cancel=None, progress=None):
    spec.validate()
    if source.empty or target.empty:
        raise ValueError("Both files must contain at least one eligible site.")
    result = []
    targets = target.to_dict("records")
    sources = source.to_dict("records")
    tlats = target["Latitude"].to_numpy(dtype=float)
    tlons = target["Longitude"].to_numpy(dtype=float)

    def metrics(s, rows, lats, lons):
        ground, bearing = distances(
            s["Latitude"], s["Longitude"], lats, lons, spec.method
        )
        metric = (
            np.sqrt(
                ground**2
                + (
                    rows["Resolved absolute (m)"].to_numpy()
                    - s["Resolved absolute (m)"]
                )
                ** 2
            )
            if spec.mode == MODES[1]
            else ground
        )
        eligible = np.isfinite(metric)
        if spec.ignore_same_location:
            eligible &= ground > 0.01
        if spec.radius_m > 0:
            eligible &= metric <= spec.radius_m
        if spec.match_attributes:
            eligible &= rows["Match attribute"].to_numpy() == s["Match attribute"]
        indices = np.flatnonzero(eligible)
        ordered = indices[np.argsort(metric[indices], kind="stable")]
        return ground, bearing, metric, ordered

    reverse = {}
    if spec.match_mode == "Mutual nearest":
        slats, slons = source["Latitude"].to_numpy(dtype=float), source[
            "Longitude"
        ].to_numpy(dtype=float)
        for number, t in enumerate(targets):
            checkpoint(cancel)
            _, _, _, ordered = metrics(t, source, slats, slons)
            if len(ordered):
                reverse[number] = int(ordered[0])
            if progress:
                progress(
                    f"Checking mutual neighbours: {number + 1:,} / {len(targets):,}"
                )
    for number, s in enumerate(sources):
        checkpoint(cancel)
        ground, bearing, metric, ordered = metrics(s, target, tlats, tlons)
        selected = (
            ordered
            if spec.match_mode == "All within radius"
            else ordered[
                : (
                    1
                    if spec.match_mode in ("Nearest eligible", "Mutual nearest")
                    else spec.top_n
                )
            ]
        )
        if (
            spec.match_mode == "Mutual nearest"
            and len(selected)
            and reverse.get(int(selected[0])) != number
        ):
            selected = []
        common = {
            "Source row": s["Input row"],
            "Source site": s["Site ID"],
            "Distance mode": spec.mode,
            "Source latitude": s["Latitude"],
            "Source longitude": s["Longitude"],
        }
        if not len(selected):
            reason = (
                "No mutual nearest target"
                if spec.match_mode == "Mutual nearest"
                else (
                    "No eligible target (attribute/distance filters)"
                    if spec.match_attributes
                    else (
                        "No target within distance limit"
                        if spec.radius_m > 0
                        else "No eligible target"
                    )
                )
            )
            result.append(
                {**common, "Match status": reason, "Terrain status": "Not evaluated"}
            )
        for rank, idx in enumerate(selected, 1):
            checkpoint(cancel)
            t = targets[idx]
            terrain = (
                terrain_check(s, t, cancel)
                if spec.terrain and rank == 1
                else ("Not evaluated: rank > 1" if spec.terrain else "Not requested")
            )
            result.append(
                {
                    **common,
                    "Target row": t["Input row"],
                    "Target site": t["Site ID"],
                    "Rank": rank,
                    "Target latitude": t["Latitude"],
                    "Target longitude": t["Longitude"],
                    "Ground distance (m)": float(ground[idx]),
                    "Distance (m)": float(metric[idx]),
                    "Bearing (deg)": (
                        float(bearing[idx]) if ground[idx] > 0.01 else np.nan
                    ),
                    "Match status": "Matched",
                    "Terrain status": terrain,
                }
            )
        if progress and (
            number + 1 == len(source) or (number + 1) % max(1, len(source) // 100) == 0
        ):
            progress(
                f"Calculated {number + 1:,} / {len(source):,} source sites",
                (number + 1) / len(source),
            )
    return pd.DataFrame(result, columns=RESULT_COLUMNS)


class SpatialIndex:
    """Conservative unit-sphere chord bounds; exact ranking in the chosen metric.

    WGS84's minimum curvature radius exceeds 6,335,000 m. Its surface line
    element is therefore >= that sphere's line element at geodetic lat/lon.
    Sphere chord <= sphere arc <= ellipsoidal distance <= slant distance.
    A metric upper bound thus gives a safe chord search radius, even at poles
    or the date line. eps=0 and a small outward pad avoid pruning boundary ties.
    """

    def __init__(self, frame, spec):
        from scipy.spatial import cKDTree

        self.frame, self.spec = frame, spec
        self.lats = frame["Latitude"].to_numpy(float)
        self.lons = frame["Longitude"].to_numpy(float)
        self.xyz = self.unit(self.lats, self.lons)
        self.tree = cKDTree(self.xyz)
        self.heights = (
            frame["Resolved absolute (m)"].to_numpy(float)
            if spec.mode == MODES[1]
            else None
        )
        self.attributes = (
            frame["Match attribute"].to_numpy() if spec.match_attributes else None
        )
        self.lower_radius = 6335000.0 if spec.method == METHODS[0] else 6371000.0
        self.evaluated = 0

    @staticmethod
    def unit(lat, lon):
        lat, lon = np.radians(lat), np.radians(lon)
        return np.stack(
            (np.cos(lat) * np.cos(lon), np.cos(lat) * np.sin(lon), np.sin(lat)), axis=-1
        )

    def exact(self, row, ids):
        ids = np.asarray(ids, dtype=int)
        self.evaluated += len(ids)
        ground, bearing = distances(
            row["Latitude"],
            row["Longitude"],
            self.lats[ids],
            self.lons[ids],
            self.spec.method,
        )
        metric = (
            np.sqrt(ground**2 + (self.heights[ids] - row["Resolved absolute (m)"]) ** 2)
            if self.heights is not None
            else ground
        )
        keep = np.isfinite(metric)
        if self.spec.ignore_same_location:
            keep &= ground > 0.01
        if self.spec.radius_m:
            keep &= metric <= self.spec.radius_m
        if self.attributes is not None:
            keep &= self.attributes[ids] == row["Match attribute"]
        ids, ground, bearing, metric = (
            ids[keep],
            ground[keep],
            bearing[keep],
            metric[keep],
        )
        order = np.lexsort((ids, metric))
        return ids[order], ground[order], bearing[order], metric[order]

    def query(self, row, cancel=None):
        checkpoint(cancel)
        point = self.unit(row["Latitude"], row["Longitude"])
        count = (
            1
            if self.spec.match_mode in ("Nearest eligible", "Mutual nearest")
            else self.spec.top_n
        )
        bound = self.spec.radius_m
        if not bound:
            # Expand until enough eligible neighbours provide an upper bound.
            # Attribute mismatches / coincident rows never terminate search early.
            k = min(16, len(self.frame))
            while True:
                checkpoint(cancel)
                _, ids = self.tree.query(point, k=k, eps=0)
                found = self.exact(row, np.atleast_1d(ids))
                if len(found[0]) >= count:
                    bound = float(found[3][count - 1])
                    break
                if k == len(self.frame):
                    return tuple(a[:count] for a in found)
                k = min(len(self.frame), k * 2)
        radius = min(2.0, bound / self.lower_radius) + 1e-12
        ids = self.tree.query_ball_point(point, radius, eps=0, return_sorted=True)
        checkpoint(cancel)
        found = self.exact(row, ids)
        return (
            found
            if self.spec.match_mode == "All within radius"
            else tuple(a[:count] for a in found)
        )


def calculate_batches(
    source, target, spec, cancel=None, progress=None, batch_size=2000, target_index=None
):
    """Yield bounded result batches, including when one source has many matches."""
    spec.validate()
    if source.empty or target.empty:
        raise ValueError("Both files must contain at least one eligible site.")
    index = target_index or SpatialIndex(target, spec)
    reverse = {}
    if spec.match_mode == "Mutual nearest":
        source_index = SpatialIndex(source, spec)
        for i, values in enumerate(target.itertuples(index=False, name=None)):
            row = dict(zip(target.columns, values))
            found = source_index.query(row, cancel)[0]
            if len(found):
                reverse[i] = int(found[0])
            if progress and i % 100 == 0:
                progress(f"Checking mutual neighbours: {i + 1:,} / {len(target):,}")
    records = []
    for number, values in enumerate(source.itertuples(index=False, name=None)):
        checkpoint(cancel)
        s = dict(zip(source.columns, values))
        ids, ground, bearing, metric = index.query(s, cancel)
        if (
            spec.match_mode == "Mutual nearest"
            and len(ids)
            and reverse.get(int(ids[0])) != number
        ):
            ids = []
        common = {
            "Source row": s["Input row"],
            "Source site": s["Site ID"],
            "Distance mode": spec.mode,
            "Source latitude": s["Latitude"],
            "Source longitude": s["Longitude"],
        }
        if not len(ids):
            reason = (
                "No mutual nearest target"
                if spec.match_mode == "Mutual nearest"
                else (
                    "No eligible target (attribute/distance filters)"
                    if spec.match_attributes
                    else (
                        "No target within distance limit"
                        if spec.radius_m > 0
                        else "No eligible target"
                    )
                )
            )
            records.append(
                {**common, "Match status": reason, "Terrain status": "Not evaluated"}
            )
        for rank, idx in enumerate(ids, 1):
            checkpoint(cancel)
            t = target.iloc[int(idx)]
            terrain = (
                terrain_check(s, t, cancel)
                if spec.terrain and rank == 1
                else ("Not evaluated: rank > 1" if spec.terrain else "Not requested")
            )
            records.append(
                {
                    **common,
                    "Target row": t["Input row"],
                    "Target site": t["Site ID"],
                    "Rank": rank,
                    "Target latitude": t["Latitude"],
                    "Target longitude": t["Longitude"],
                    "Ground distance (m)": float(ground[rank - 1]),
                    "Distance (m)": float(metric[rank - 1]),
                    "Bearing (deg)": (
                        float(bearing[rank - 1]) if ground[rank - 1] > 0.01 else np.nan
                    ),
                    "Match status": "Matched",
                    "Terrain status": terrain,
                }
            )
            if len(records) >= batch_size:
                yield pd.DataFrame(records, columns=RESULT_COLUMNS)
                records.clear()
        if len(records) >= batch_size:
            yield pd.DataFrame(records, columns=RESULT_COLUMNS)
            records.clear()
        if progress and (
            number + 1 == len(source) or (number + 1) % max(1, len(source) // 100) == 0
        ):
            progress(
                f"Calculated {number + 1:,} / {len(source):,} source sites",
                (number + 1) / len(source),
            )
    if records:
        yield pd.DataFrame(records, columns=RESULT_COLUMNS)


def calculate(source, target, spec, cancel=None, progress=None):
    return pd.concat(
        calculate_batches(source, target, spec, cancel, progress), ignore_index=True
    )


class WorkingCopy:
    """Cell/exclusion deltas with undo/redo. Original DataFrame is never mutated."""

    def __init__(self, raw, file):
        self.original = raw
        self.frame = raw.copy()
        self.edits = {(int(r), str(c)): str(v) for r, c, v in file.edits}
        self.excluded = set(file.excluded)
        for (row, col), value in self.edits.items():
            self.frame.at[row, col] = value
        self.undo_stack, self.redo_stack = [], []

    def commit(self, changes):
        if not changes:
            return
        self.undo_stack.append(changes)
        self.redo_stack.clear()
        self.apply(changes, True)

    def apply(self, changes, forward):
        for row, col, before, after in changes:
            value = after if forward else before
            if col is None:
                (self.excluded.add if value else self.excluded.discard)(row)
            else:
                self.frame.at[row, col] = value
                if value == str(self.original.at[row, col]):
                    self.edits.pop((row, col), None)
                else:
                    self.edits[row, col] = value

    def edit(self, row, col, value):
        before, after = str(self.frame.at[row, col]), str(value)
        if before != after:
            self.commit([(row, col, before, after)])

    def exclude(self, rows, excluded):
        self.commit(
            [
                (r, None, r in self.excluded, excluded)
                for r in sorted(set(rows))
                if (r in self.excluded) != excluded
            ]
        )

    def undo(self):
        if self.undo_stack:
            changes = self.undo_stack.pop()
            self.apply(changes, False)
            self.redo_stack.append(changes)

    def redo(self):
        if self.redo_stack:
            changes = self.redo_stack.pop()
            self.apply(changes, True)
            self.undo_stack.append(changes)

    def file(self, original):
        return replace(
            original,
            edits=tuple((r, c, v) for (r, c), v in sorted(self.edits.items())),
            excluded=tuple(sorted(self.excluded)),
        )


def validate_working(raw, file, spec, cancel=None, offset=0, seen=None):
    raw = raw.reset_index(drop=True)
    edited = raw.copy()
    changes = [[] for _ in range(len(raw))]
    for row, col, value in file.edits:
        if offset <= row < offset + len(raw):
            if col not in raw:
                raise ValueError(f"Edited column missing: {col}")
            local = row - offset
            changes[local].append(
                {"column": col, "original": str(raw.at[local, col]), "working": value}
            )
            edited.at[local, col] = value
    excluded = {r - offset for r in file.excluded if offset <= r < offset + len(raw)}
    # Exclude before duplicate detection: an excluded first occurrence must not
    # suppress a later eligible occurrence with the same identifier.
    validation_input = edited.copy()
    id_col = dict(file.mapping)["Site ID"]
    for row in excluded:
        validation_input.at[row, id_col] = ""
    if seen is not None:
        for row, value in enumerate(validation_input[id_col].astype(str).str.strip()):
            if value in seen and row not in excluded:
                validation_input.at[row, id_col] = ""
    valid, audit = validate_frame(
        validation_input, dict(file.mapping), spec, cancel, file
    )
    # Restore working values in the audit; retain original changed values explicitly.
    for col in edited:
        audit[
            (
                col
                if col
                not in (
                    "GeoMatch input row",
                    "GeoMatch reason",
                    "GeoMatch disposition",
                    "GeoMatch WGS84 latitude",
                    "GeoMatch WGS84 longitude",
                    "GeoMatch coordinate warning",
                )
                else "Original " + col
            )
        ] = edited[col]
    for row in excluded:
        audit.at[row, "GeoMatch reason"] = "Explicit row exclusion in working copy"
        audit.at[row, "GeoMatch disposition"] = "Excluded by user"
    if seen is not None:
        for row, value in enumerate(edited[id_col].astype(str).str.strip()):
            if value in seen and row not in excluded:
                audit.at[row, "GeoMatch reason"] = (
                    "Duplicate site ID (first valid row retained)"
                )
        seen.update(valid["Site ID"].tolist())
    for name in ("GeoMatch edits",):
        while name in audit:
            audit.rename(columns={name: "Original " + name}, inplace=True)
    audit["GeoMatch edits"] = [
        json.dumps(v, ensure_ascii=False) if v else "" for v in changes
    ]
    valid["Input row"] += offset
    audit["GeoMatch input row"] += offset
    return valid, audit


def checked_raw(file, cancel=None, progress=None):
    if signature(file.path) != file.signature or (
        file.sha256 and input_fingerprint(file.path, cancel) != file.sha256
    ):
        raise ValueError(
            "Input changed after mapping. Import it again before editing or validating."
        )
    raw = read_input(file.path, cancel, progress, file.sheet)
    if file.sha256 and input_fingerprint(file.path, cancel) != file.sha256:
        raise ValueError("Input changed while reading. Import it again.")
    return raw


def ensure_separate_output(path, spec):
    if Path(path).resolve() in (
        Path(spec.source.path).resolve(),
        Path(spec.target.path).resolve(),
    ):
        raise ValueError(
            "Choose a separate output path; input files cannot be overwritten."
        )


def save_project(path, spec, exclude_flags=False, cancel=None, geo_settings=None):
    """Versioned JSON, no executable deserialization; clipboard data is embedded."""
    spec.validate()
    ensure_separate_output(path, spec)
    content = {
        "format": "GeoMatch project",
        "schema": 1,
        "version": VERSION,
        "run": asdict(spec),
        "exclude_flags": bool(exclude_flags),
        "clipboard": {},
        "geo_tools": geo_settings or {},
    }
    for role in ("source", "target"):
        file = getattr(spec, role)
        checkpoint(cancel)
        if input_fingerprint(file.path, cancel) != file.sha256:
            raise ValueError(
                f"{role.title()} input changed. Import again before saving a project."
            )
        if file.origin == "clipboard":
            content["clipboard"][role] = Path(file.path).read_bytes().decode("utf-8")
        content["run"][role]["path"] = str(Path(file.path).resolve())
    path = Path(path)
    temporary = path.with_name(f".geomatch-project-{os.getpid()}-{time.time_ns()}.tmp")
    try:
        with open(temporary, "x", encoding="utf-8") as stream:
            json.dump(content, stream, ensure_ascii=False, allow_nan=False)
        checkpoint(cancel)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return str(path)


def load_project(path, clipboard_dir, cancel=None):
    content = json.loads(Path(path).read_text(encoding="utf-8"))
    if content.get("format") != "GeoMatch project" or content.get("schema") != 1:
        raise ValueError("Unsupported GeoMatch project format.")
    run = dict(content["run"])
    for role in ("source", "target"):
        data = dict(run[role])
        if role in content.get("clipboard", {}):
            embedded = content["clipboard"][role]
            import hashlib

            if hashlib.sha256(embedded.encode("utf-8")).hexdigest() != data["sha256"]:
                raise ValueError("Embedded clipboard checksum does not match.")
            snapshot = Path(clipboard_dir) / f"Project_{role}_{time.time_ns()}.csv"
            snapshot.write_text(embedded, encoding="utf-8", newline="")
            data["path"] = str(snapshot)
            data["origin"] = "clipboard"
        checkpoint(cancel)
        if input_fingerprint(data["path"], cancel) != data["sha256"]:
            raise ValueError(
                f"{role.title()} contents differ from this saved project. Restore the original file or import it again."
            )
        data["signature"] = signature(data["path"])
        data["mapping"] = tuple(tuple(v) for v in data["mapping"])
        data["edits"] = tuple(tuple(v) for v in data.get("edits", ()))
        data["excluded"] = tuple(data.get("excluded", ()))
        if any(
            not isinstance(r, int)
            or not 0 <= r < data["rows"]
            or not isinstance(c, str)
            or not isinstance(v, str)
            for r, c, v in data["edits"]
        ):
            raise ValueError("Invalid working-copy edits in project.")
        if any(
            not isinstance(r, int) or not 0 <= r < data["rows"]
            for r in data["excluded"]
        ):
            raise ValueError("Invalid excluded row in project.")
        horizontal_crs(data["crs"])
        run[role] = FileSpec(**data)
    spec = RunSpec(**run)
    spec.validate()
    return spec, bool(content.get("exclude_flags", False))


def stream_large_run(
    path, spec, cancel=None, progress=None, chunk_size=10000, temporary=None
):
    """CSV/TXT source streaming; target and index stay resident, results go to ZIP.

    Global outlier fitting and mutual nearest require global source state and are
    deliberately refused here; the ordinary reviewed workflow supports them.
    """
    import zipfile
    import io

    spec.validate()
    ensure_separate_output(path, spec)
    if Path(spec.source.path).suffix.lower() not in (".csv", ".txt"):
        raise ValueError(
            "Large-file mode requires a CSV/TXT source. Export other source formats to CSV first."
        )
    if spec.flag_outliers or spec.match_mode == "Mutual nearest" or spec.terrain:
        raise ValueError(
            "Large-file mode supports Nearest N, Nearest eligible and All within radius, with outlier and online terrain checks off. Use the reviewed workflow for other modes."
        )
    if input_fingerprint(spec.source.path, cancel) != spec.source.sha256:
        raise ValueError("Source contents changed after mapping.")
    target_raw = checked_raw(spec.target, cancel, progress)
    target, target_audit = validate_working(target_raw, spec.target, spec, cancel)
    del target_raw
    if target.empty:
        raise ValueError("Target has no eligible sites.")
    index = SpatialIndex(target, spec)
    path = Path(path)
    temporary = (
        Path(temporary)
        if temporary
        else path.with_name(f".geomatch-stream-{os.getpid()}-{time.time_ns()}.zip")
    )
    rows = results = eligible = 0
    seen = set()
    try:
        with zipfile.ZipFile(
            temporary, "x", compression=zipfile.ZIP_DEFLATED
        ) as archive:
            # ZIP permits only one open write entry, so stream each table through
            # disk-backed temporary files; never accumulate all source/results.
            from contextlib import nullcontext

            with nullcontext():
                result_path, audit_path = Path(str(temporary) + ".results.csv"), Path(
                    str(temporary) + ".audit.csv"
                )
                with open(
                    result_path, "w", newline="", encoding="utf-8-sig"
                ) as result_out, open(
                    audit_path, "w", newline="", encoding="utf-8-sig"
                ) as audit_out:
                    pd.DataFrame(columns=RESULT_COLUMNS).to_csv(result_out, index=False)
                    with pd.read_csv(
                        spec.source.path,
                        sep=None,
                        engine="python",
                        dtype=str,
                        keep_default_na=False,
                        chunksize=chunk_size,
                    ) as reader:
                        for raw in reader:
                            checkpoint(cancel)
                            source, audit = validate_working(
                                raw, spec.source, spec, cancel, rows, seen
                            )
                            audit.to_csv(audit_out, index=False, header=rows == 0)
                            rows += len(raw)
                            del raw
                            eligible += len(source)
                            if not source.empty:
                                for batch in calculate_batches(
                                    source, target, spec, cancel, target_index=index
                                ):
                                    batch.to_csv(result_out, index=False, header=False)
                                    results += len(batch)
                            if progress:
                                progress(
                                    f"Streamed {rows:,} source rows / {results:,} result rows"
                                )
                if input_fingerprint(spec.source.path, cancel) != spec.source.sha256:
                    raise ValueError(
                        "Source changed during streaming; output was not replaced."
                    )
                for file in (result_path, audit_path):
                    checkpoint(cancel)
                    with open(file, "rb") as source, archive.open(
                        "Results.csv" if file == result_path else "Source review.csv",
                        "w",
                        force_zip64=True,
                    ) as out:
                        while block := source.read(1024 * 1024):
                            checkpoint(cancel)
                            out.write(block)
                with archive.open("Target review.csv", "w", force_zip64=True) as out:
                    with io.TextIOWrapper(
                        out, encoding="utf-8-sig", newline=""
                    ) as text:
                        target_audit.to_csv(text, index=False, chunksize=2000)
                archive.writestr(
                    "Run settings.json",
                    json.dumps(
                        {
                            "Product": TITLE,
                            "Version": VERSION,
                            "run": asdict(spec),
                            "source_rows": rows,
                            "eligible_source_rows": eligible,
                            "result_rows": results,
                            "target_coordinate_operation": target.attrs.get(
                                "coordinate_operation"
                            ),
                            "source_coordinate_operation": (
                                audit.attrs.get("coordinate_operation")
                                if rows
                                else None
                            ),
                            "notes": "CSV tables preserve text; import ID columns as Text in spreadsheet software. Original input files are unchanged. Heights in metres; no vertical transformation. Stable input-order ties.",
                        },
                        ensure_ascii=False,
                    ),
                )
        checkpoint(cancel)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
        Path(str(temporary) + ".results.csv").unlink(missing_ok=True)
        Path(str(temporary) + ".audit.csv").unlink(missing_ok=True)
    return {
        "path": str(path),
        "source_rows": rows,
        "eligible": eligible,
        "results": results,
    }


def gis_features(result, frames, cancel=None):
    layers = {"source_points": [], "target_points": [], "match_links": []}

    def clean(record):
        return {
            str(k): None if pd.isna(v) else v.item() if isinstance(v, np.generic) else v
            for k, v in record.items()
        }

    matched = result.loc[result["Match status"].eq("Matched")]
    source_counts = matched.groupby("Source row").size().to_dict()
    target_counts = matched.groupby("Target row").size().to_dict()
    source_status = (
        result.drop_duplicates("Source row")
        .set_index("Source row")["Match status"]
        .to_dict()
    )
    for layer, frame in zip(("source_points", "target_points"), frames):
        for row in frame.to_dict("records"):
            checkpoint(cancel)
            layers[layer].append(
                {
                    "type": "Feature",
                    "properties": {
                        **clean(row),
                        "GeoMatch layer": layer,
                        "GeoMatch link count": int(
                            (
                                source_counts
                                if layer == "source_points"
                                else target_counts
                            ).get(row["Input row"], 0)
                        ),
                        "GeoMatch match status": (
                            source_status.get(row["Input row"], "Unmatched")
                            if layer == "source_points"
                            else "Candidate target"
                        ),
                    },
                    "geometry": {
                        "type": "Point",
                        "coordinates": [row["Longitude"], row["Latitude"]],
                    },
                }
            )
    for row in result.to_dict("records"):
        checkpoint(cancel)
        if row["Match status"] != "Matched":
            continue
        x1, y1, x2, y2 = (
            row["Source longitude"],
            row["Source latitude"],
            row["Target longitude"],
            row["Target latitude"],
        )
        # Split visual connectors at the antimeridian rather than drawing across the world.
        if abs(x2 - x1) > 180:
            adjusted = x2 - 360 if x2 > x1 else x2 + 360
            edge = -180 if adjusted < x1 else 180
            mid_y = y1 + (y2 - y1) * (edge - x1) / (adjusted - x1)
            geometry = {
                "type": "MultiLineString",
                "coordinates": [[[x1, y1], [edge, mid_y]], [[-edge, mid_y], [x2, y2]]],
            }
        else:
            geometry = {
                "type": "MultiLineString",
                "coordinates": [[[x1, y1], [x2, y2]]],
            }
        layers["match_links"].append(
            {
                "type": "Feature",
                "properties": {**clean(row), "GeoMatch layer": "match_links"},
                "geometry": geometry,
            }
        )
    return layers


def export_gis(
    path, result, frames, audits, metadata, cancel=None, progress=None, temporary=None
):
    path = Path(path)
    if path.suffix.lower() not in (".geojson", ".gpkg"):
        raise ValueError("GIS export must use .geojson or .gpkg.")
    temporary = (
        Path(temporary)
        if temporary
        else path.with_name(f".geomatch-{os.getpid()}-{time.time_ns()}{path.suffix}")
    )
    if temporary.exists():
        raise ValueError("Temporary output path already exists.")
    try:
        layers = gis_features(result, frames, cancel)
        if path.suffix.lower() == ".geojson":
            # RFC 7946 WGS84 longitude/latitude; run metadata is a foreign member.
            content = {
                "type": "FeatureCollection",
                "geomatch_run": metadata,
                "features": [feature for layer in layers.values() for feature in layer],
            }
            with open(temporary, "x", encoding="utf-8") as stream:
                json.dump(content, stream, ensure_ascii=False, allow_nan=False)
        else:
            try:
                import geopandas as gpd
                import pyogrio
            except ImportError as error:
                raise ValueError(
                    "GeoPackage export needs geopandas and pyogrio. Install requirements-geomatch-5.3.txt, then restart."
                ) from error
            import sqlite3

            for layer, features in layers.items():
                checkpoint(cancel)
                if progress:
                    progress(f"Saving GIS layer: {layer}…")
                if features:
                    frame = gpd.GeoDataFrame.from_features(features, crs="EPSG:4326")
                else:
                    frame = gpd.GeoDataFrame(
                        {"geometry": []}, geometry="geometry", crs="EPSG:4326"
                    )
                frame.to_file(
                    temporary,
                    layer=layer,
                    driver="GPKG",
                    engine="pyogrio",
                    geometry_type=(
                        "MultiLineString" if layer == "match_links" else "Point"
                    ),
                )
            checkpoint(cancel)
            from contextlib import closing

            with closing(sqlite3.connect(temporary)) as database, database:
                tables = {
                    "run_settings": pd.DataFrame(
                        {
                            "Setting": list(metadata),
                            "Value": [str(v) for v in metadata.values()],
                        }
                    ),
                    "source_review": audits[0],
                    "target_review": audits[1],
                    "match_results": result,
                }
                for name, frame in tables.items():
                    checkpoint(cancel)
                    key = "geomatch_fid"
                    while key.casefold() in {str(c).casefold() for c in frame.columns}:
                        key = "_" + key
                    if len({str(c).casefold() for c in frame.columns}) != len(
                        frame.columns
                    ):
                        # SQLite names are case-insensitive; JSON preserves colliding source headers.
                        frame = pd.DataFrame(
                            {
                                "record_json": [
                                    json.dumps(
                                        {
                                            str(k): (
                                                None
                                                if pd.isna(v)
                                                else (
                                                    v.item()
                                                    if isinstance(v, np.generic)
                                                    else v
                                                )
                                            )
                                            for k, v in row.items()
                                        },
                                        ensure_ascii=False,
                                    )
                                    for row in frame.to_dict("records")
                                ]
                            }
                        )
                    quote = lambda value: '"' + str(value).replace('"', '""') + '"'
                    definitions = [f"{quote(key)} INTEGER PRIMARY KEY"]
                    definitions += [
                        f"{quote(c)} {'REAL' if pd.api.types.is_numeric_dtype(frame[c]) else 'TEXT'}"
                        for c in frame.columns
                    ]
                    database.execute(
                        f"CREATE TABLE {quote(name)} ({', '.join(definitions)})"
                    )
                    frame.to_sql(name, database, index=False, if_exists="append")
                    database.execute(
                        "INSERT INTO gpkg_contents(table_name,data_type,identifier,description) VALUES (?, 'attributes', ?, ?)",
                        (name, name, "GeoMatch report table"),
                    )
        checkpoint(cancel)
        os.replace(temporary, path)
    finally:
        for candidate in (
            temporary,
            Path(str(temporary) + "-wal"),
            Path(str(temporary) + "-shm"),
            Path(str(temporary) + "-journal"),
        ):
            candidate.unlink(missing_ok=True)


def export_workbook(
    path, result, audits, metadata, cancel=None, progress=None, temporary=None
):
    path = Path(path)
    if not path.parent.is_dir():
        raise ValueError("The output folder does not exist.")
    sheets = {
        "Results": result,
        "Source review": audits[0],
        "Target review": audits[1],
        "Run settings": pd.DataFrame(
            {"Setting": list(metadata), "Value": [str(v) for v in metadata.values()]}
        ),
    }
    if any(len(frame) > 1048575 for frame in sheets.values()):
        raise ValueError(
            "A sheet exceeds Excel's row limit. Reduce the run size before exporting."
        )
    if temporary is None:
        descriptor, temporary = tempfile.mkstemp(
            prefix=".sitespan-", suffix=".xlsx", dir=path.parent
        )
        os.close(descriptor)
    else:
        # The GUI knows this exact path so it can clean up a terminated export.
        with open(temporary, "xb"):
            pass
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.utils import get_column_letter

        workbook = Workbook(write_only=True)
        try:
            for name, frame in sheets.items():
                checkpoint(cancel)
                if progress:
                    progress(f"Saving {name.lower()}.")
                worksheet = workbook.create_sheet(name)
                worksheet.freeze_panes = "A2"
                worksheet.auto_filter.ref = (
                    f"A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}"
                )

                def cells(values):
                    for value in values:
                        if pd.isna(value):
                            value = None
                        elif isinstance(value, np.generic):
                            value = value.item()
                        cell = WriteOnlyCell(worksheet, value=value)
                        if isinstance(value, str):
                            cell.data_type = "s"
                        yield cell

                worksheet.append(list(cells(frame.columns)))
                for number, row in enumerate(frame.itertuples(index=False, name=None)):
                    if number % 1000 == 0:
                        checkpoint(cancel)
                    worksheet.append(list(cells(row)))
            workbook.save(temporary)
        finally:
            for sheet in workbook.worksheets:
                if not sheet.closed:
                    sheet.close()
                writer = getattr(sheet, "_writer", None)
                if writer and os.path.exists(writer.out):
                    writer.cleanup()
            workbook.close()
        checkpoint(cancel)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


GEO_TOOLS = (
    "Nearest line / road",
    "Polygon membership / proximity",
    "Survey comparison by ID",
    "Buffers and nearby counts",
    "DEM profiles of matched links",
    "Telecom sectors",
)


def positive_number(value, name, allow_zero=False):
    number = float(value)
    if not math.isfinite(number) or number < 0 or (not allow_zero and number == 0):
        raise ValueError(
            f"{name} must be finite and {'nonnegative' if allow_zero else 'greater than zero'}."
        )
    return number


def metric_projection(crs_text, frames):
    from pyproj import Transformer

    crs = horizontal_crs(crs_text)
    if not crs.is_projected or any(
        abs(axis.unit_conversion_factor - 1) > 1e-12 for axis in crs.axis_info[:2]
    ):
        raise ValueError(
            "Choose a projected analysis CRS with metre units (for example the appropriate local UTM zone), not latitude/longitude or feet."
        )
    if crs.to_epsg() == 3857:
        raise ValueError(
            "Web Mercator is unsuitable for this metric analysis. Choose a local projected CRS such as the correct UTM zone."
        )
    transform = Transformer.from_crs(
        "EPSG:4326", crs, always_xy=True, allow_ballpark=False, only_best=True
    )
    projected = []
    area = crs.area_of_use
    for frame in frames:
        lon, lat = frame["Longitude"].to_numpy(float), frame["Latitude"].to_numpy(float)
        if area is not None:
            within_lon = (
                (lon >= area.west) & (lon <= area.east)
                if area.west <= area.east
                else (lon >= area.west) | (lon <= area.east)
            )
            if not (within_lon & (lat >= area.south) & (lat <= area.north)).all():
                raise ValueError(
                    "Coordinates fall outside the analysis CRS area of use. Choose a suitable CRS or split the study area."
                )
        x, y = transform.transform(lon, lat, errcheck=True)
        xy = np.column_stack([x, y])
        if not np.isfinite(xy).all():
            raise ValueError("Analysis CRS conversion produced nonfinite coordinates.")
        projected.append(xy)
    return crs, transform, projected


def vector_features(path, layer, analysis_crs, wanted, cancel=None):
    import geopandas as gpd
    import pyogrio
    import shapely
    from pyproj import Transformer

    before = input_fingerprint(path, cancel)
    if Path(path).suffix.lower() not in (".gpkg", ".geojson", ".json"):
        raise ValueError("Choose a local GeoPackage or GeoJSON vector layer.")
    if Path(path).suffix.lower() == ".gpkg" and not layer:
        layers = [name for name, geometry in pyogrio.list_layers(path) if geometry]
        if len(layers) != 1:
            raise ValueError(
                "Enter the GeoPackage layer name. Available spatial layers: "
                + ", ".join(layers)
            )
        layer = layers[0]
    if Path(path).suffix.lower() in (".geojson", ".json"):
        with open(path, encoding="utf-8-sig") as stream:
            info = json.load(stream)
        if info.get("crs"):
            raise ValueError(
                "Legacy GeoJSON CRS declarations are unsupported; use WGS84 RFC 7946 GeoJSON or a GeoPackage with its CRS."
            )
    data = gpd.read_file(path, layer=layer or None, engine="pyogrio")
    checkpoint(cancel)
    if data.crs is None or data.empty:
        raise ValueError("The vector layer needs a CRS and at least one geometry.")
    source_crs = horizontal_crs(data.crs)
    allowed = (
        ("LineString", "MultiLineString")
        if wanted == "line"
        else ("Polygon", "MultiPolygon")
    )
    for i, geom in enumerate(data.geometry):
        checkpoint(cancel)
        if (
            geom is None
            or geom.is_empty
            or not geom.is_valid
            or geom.geom_type not in allowed
        ):
            raise ValueError(
                f"Vector feature {i + 1} is empty, invalid or not a {'line' if wanted == 'line' else 'polygon'}. Correct/filter the layer first."
            )
    # Check the actual vertices, not only centroids, against the analysis area.
    vertices = shapely.get_coordinates(data.geometry.to_numpy())
    to_wgs = Transformer.from_crs(
        source_crs, "EPSG:4326", always_xy=True, allow_ballpark=False, only_best=True
    )
    lon, lat = to_wgs.transform(vertices[:, 0], vertices[:, 1], errcheck=True)
    metric_projection(analysis_crs, [pd.DataFrame({"Longitude": lon, "Latitude": lat})])
    to_metric = Transformer.from_crs(
        source_crs, analysis_crs, always_xy=True, allow_ballpark=False, only_best=True
    )
    geometries = shapely.transform(
        shapely.force_2d(data.geometry.to_numpy()),
        to_metric.transform,
        interleaved=False,
    )
    if any(g.is_empty or not g.is_valid for g in geometries):
        raise ValueError("Vector geometry became invalid after projection.")
    if input_fingerprint(path, cancel) != before:
        raise ValueError("Vector file changed during reading.")
    attributes = data.drop(columns=data.geometry.name).reset_index(drop=True)
    return (
        geometries,
        attributes,
        {
            "path": str(path),
            "layer": layer,
            "sha256": before,
            "input_crs": source_crs.to_wkt(),
            "geometry_Z": "Ignored; analysis is 2D",
        },
    )


def nearest_vector(
    source, geometries, attributes, crs_text, polygon=False, cancel=None
):
    from shapely import Point, STRtree
    from shapely.ops import nearest_points

    crs, transform, coordinates = metric_projection(crs_text, [source])
    tree = STRtree(geometries)
    records, connectors = [], []
    for i, xy in enumerate(coordinates[0]):
        checkpoint(cancel)
        point = Point(xy)
        # query_nearest returns all equidistant matches; file order breaks ties.
        nearest = np.asarray(tree.query_nearest(point, all_matches=True), dtype=int)
        index = int(nearest.min())
        geom = geometries[index]
        foot = nearest_points(point, geom)[1]
        from shapely import LineString

        connectors.append(LineString([point, foot]) if not point.equals(foot) else None)
        row = source.iloc[i]
        record = {
            "Source row": row["Input row"],
            "Source site": row["Site ID"],
            "Feature row": index + 1,
            "Distance (m)": float(point.distance(geom)),
            "Nearest X (m)": foot.x,
            "Nearest Y (m)": foot.y,
            "Equal nearest features": len(nearest),
            "Analysis CRS": crs.to_string(),
        }
        if polygon:
            covering = sorted(int(v) for v in tree.query(point, predicate="covered_by"))
            record.update(
                {
                    "Relation": (
                        "Inside"
                        if geom.contains(point)
                        else "Boundary" if geom.covers(point) else "Outside"
                    ),
                    "Covering polygon count": len(covering),
                    "Covering feature rows": ";".join(str(v + 1) for v in covering),
                    "Boundary distance (m)": float(point.distance(geom.boundary)),
                }
            )
        for key, value in attributes.iloc[index].items():
            record["Feature " + str(key)] = value
        records.append(record)
    return pd.DataFrame(records), connectors, crs.to_string()


def resolved_elevation(row):
    absolute = row.get("Absolute antenna elevation (m)", np.nan)
    if pd.notna(absolute) and math.isfinite(float(absolute)):
        return float(absolute)
    terrain, agl = row.get("Terrain elevation (m)", np.nan), row.get(
        "Antenna AGL (m)", np.nan
    )
    if (
        pd.notna(terrain)
        and pd.notna(agl)
        and math.isfinite(float(terrain))
        and math.isfinite(float(agl))
        and agl >= 0
    ):
        return float(terrain + agl)
    return np.nan


def survey_comparison(
    source, target, crs_text, horizontal_tolerance, vertical_tolerance=None, cancel=None
):
    from shapely import Point

    ht = positive_number(horizontal_tolerance, "Horizontal tolerance", True)
    vt = (
        positive_number(vertical_tolerance, "Vertical tolerance", True)
        if vertical_tolerance is not None
        else None
    )
    crs, _, xy = metric_projection(crs_text, [source, target])
    source, target = source.reset_index(drop=True), target.reset_index(drop=True)
    observed = {str(row["Site ID"]): i for i, row in target.iterrows()}
    records, geometries, used = [], [], set()
    for i, baseline in source.reset_index(drop=True).iterrows():
        checkpoint(cancel)
        name = str(baseline["Site ID"])
        j = observed.get(name)
        record = {
            "Site ID": name,
            "Baseline row": baseline["Input row"],
            "Analysis CRS": crs.to_string(),
        }
        geometries.append(Point(xy[0][i]))
        if j is None:
            record["Status"] = "Missing in observed"
        else:
            used.add(j)
            measured = target.iloc[j]
            delta = xy[1][j] - xy[0][i]
            horizontal = float(np.linalg.norm(delta))
            dz = resolved_elevation(measured) - resolved_elevation(baseline)
            status = (
                "Outside tolerance"
                if horizontal > ht
                or (vt is not None and math.isfinite(dz) and abs(dz) > vt)
                else (
                    "Height unavailable"
                    if vt is not None and not math.isfinite(dz)
                    else "Within tolerance"
                )
            )
            ground, bearing = distances(
                baseline["Latitude"],
                baseline["Longitude"],
                np.array([measured["Latitude"]]),
                np.array([measured["Longitude"]]),
                METHODS[0],
            )
            record.update(
                {
                    "Observed row": measured["Input row"],
                    "Delta easting (m)": float(delta[0]),
                    "Delta northing (m)": float(delta[1]),
                    "Horizontal change (m)": horizontal,
                    "WGS84 ground change (m)": float(ground[0]),
                    "Delta height (m)": dz,
                    "Horizontal tolerance (m)": ht,
                    "Vertical tolerance (m)": vt,
                    "Status": status,
                }
            )
        records.append(record)
    for j, measured in target.reset_index(drop=True).iterrows():
        if j not in used:
            records.append(
                {
                    "Site ID": measured["Site ID"],
                    "Observed row": measured["Input row"],
                    "Status": "New in observed",
                    "Analysis CRS": crs.to_string(),
                }
            )
            geometries.append(Point(xy[1][j]))
    return pd.DataFrame(records), geometries, crs.to_string()


def buffers_counts(source, target, crs_text, radius, spec, cancel=None):
    from shapely import Point
    from scipy.spatial import cKDTree

    radius = positive_number(radius, "Buffer radius")
    crs, _, xy = metric_projection(crs_text, [source, target])
    tree = cKDTree(xy[1]) if len(target) else None
    records, geometries = [], []
    for i, point in enumerate(xy[0]):
        checkpoint(cancel)
        ids = np.asarray(
            tree.query_ball_point(point, radius + 1e-8, eps=0) if tree else [],
            dtype=int,
        )
        if len(ids):
            d = np.linalg.norm(xy[1][ids] - point, axis=1)
            keep = d <= radius
            if spec.ignore_same_location:
                keep &= d > 0.01
            if spec.match_attributes:
                keep &= (
                    target.iloc[ids]["Match attribute"].to_numpy()
                    == source.iloc[i]["Match attribute"]
                )
            ids, d = ids[keep], d[keep]
        else:
            d = np.array([])
        geometries.append(Point(point).buffer(radius, quad_segs=32))
        records.append(
            {
                "Source row": source.iloc[i]["Input row"],
                "Source site": source.iloc[i]["Site ID"],
                "Radius (m)": radius,
                "Nearby target count": len(ids),
                "Nearest target distance (m)": float(d.min()) if len(d) else np.nan,
                "Analysis CRS": crs.to_string(),
                "Count method": "Exact projected centre distance <= radius; display buffer approximated with 128 edges",
            }
        )
    return pd.DataFrame(records), geometries, crs.to_string()


def telecom_sectors(
    source, target, crs_text, radius, azimuth, beamwidth, spec, cancel=None
):
    from shapely import Polygon
    from pyproj import Geod

    radius = positive_number(radius, "Sector range")
    azimuth, beamwidth = float(azimuth), positive_number(beamwidth, "Beamwidth")
    if not math.isfinite(azimuth) or not 0 <= azimuth < 360 or beamwidth > 360:
        raise ValueError("Azimuth must be in [0, 360); beamwidth must be in (0, 360].")
    crs, transform, xy = metric_projection(crs_text, [source, target])
    geod = Geod(ellps="WGS84")
    query_spec = replace(
        spec,
        method=METHODS[0],
        mode=MODES[0],
        match_mode="All within radius",
        radius_m=radius,
    )
    index = SpatialIndex(target, query_spec) if not target.empty else None
    records, geometries = [], []
    angles = np.linspace(
        azimuth - beamwidth / 2,
        azimuth + beamwidth / 2,
        max(3, int(math.ceil(beamwidth)) + 1),
    )
    for i, row in source.reset_index(drop=True).iterrows():
        checkpoint(cancel)
        lon, lat, _ = geod.fwd(
            np.full(len(angles), row["Longitude"]),
            np.full(len(angles), row["Latitude"]),
            angles,
            np.full(len(angles), radius),
        )
        metric_projection(crs, [pd.DataFrame({"Longitude": lon, "Latitude": lat})])
        x, y = transform.transform(lon, lat, errcheck=True)
        vertices = list(zip(x, y))
        if beamwidth < 360:
            vertices = [tuple(xy[0][i])] + vertices + [tuple(xy[0][i])]
        geom = Polygon(vertices)
        if not geom.is_valid:
            raise ValueError(
                "Sector shape is invalid in this CRS; reduce range or choose another study area."
            )
        geometries.append(geom)
        found = index.query(row, cancel) if index else ([], [], [], [])
        bearings, ground = np.asarray(found[2]), np.asarray(found[1])
        difference = np.abs((bearings - azimuth + 180) % 360 - 180)
        count = int(((difference <= beamwidth / 2 + 1e-10) | (ground <= 0.01)).sum())
        records.append(
            {
                "Source row": row["Input row"],
                "Source site": row["Site ID"],
                "Azimuth true north (deg)": azimuth,
                "Beamwidth (deg)": beamwidth,
                "Range (m)": radius,
                "Targets in sector": count,
                "Analysis CRS": crs.to_string(),
                "Model": "Geometric planning sector; no propagation or coverage prediction",
            }
        )
    return pd.DataFrame(records), geometries, crs.to_string()


def dem_profiles(
    links,
    path,
    spacing=30,
    source_agl=30,
    target_agl=30,
    frequency_ghz=5,
    fresnel_fraction=0.6,
    k_factor=4 / 3,
    units="metres",
    band=1,
    cancel=None,
    progress=None,
):
    try:
        import rasterio
    except ImportError as error:
        raise ValueError(
            "Local DEM profiles require rasterio. Install requirements-geomatch-5.3.txt and restart."
        ) from error
    from pyproj import Geod, Transformer
    from shapely import LineString

    spacing = positive_number(spacing, "Sample spacing")
    source_agl = positive_number(source_agl, "Source AGL", True)
    target_agl = positive_number(target_agl, "Target AGL", True)
    frequency = positive_number(frequency_ghz, "Frequency") * 1e9
    k_factor = positive_number(k_factor, "Effective Earth radius factor")
    fresnel_fraction = positive_number(fresnel_fraction, "Fresnel fraction", True)
    if fresnel_fraction > 1 or units not in ("metres", "international feet"):
        raise ValueError("Choose DEM units and a Fresnel fraction from 0 to 1.")
    fingerprint = input_fingerprint(path, cancel)
    geod = Geod(ellps="WGS84")
    matched = links.loc[links["Match status"].eq("Matched")]
    if matched.empty:
        raise ValueError(
            "Calculate matched point links first; the DEM tool profiles those frozen results."
        )
    records, profiles, geometries = [], [], []
    budget = 0
    with rasterio.open(path) as dem:
        if dem.crs is None or not 1 <= int(band) <= dem.count:
            raise ValueError("The DEM requires a CRS and a valid band number.")
        transform = Transformer.from_crs(
            "EPSG:4326",
            horizontal_crs(dem.crs),
            always_xy=True,
            allow_ballpark=False,
            only_best=True,
        )
        metadata = {
            "rasterio": rasterio.__version__,
            "GDAL": rasterio.__gdal_version__,
            "path": str(path),
            "sha256": fingerprint,
            "crs": dem.crs.to_wkt(),
            "band": int(band),
            "chosen_units": units,
            "band_units": dem.units[int(band) - 1],
            "band_scale": dem.scales[int(band) - 1],
            "band_offset": dem.offsets[int(band) - 1],
            "resolution": list(dem.res),
            "sampling": "Nearest raster cell; no interpolation or vertical datum transformation",
        }
        for link_number, row in enumerate(matched.to_dict("records"), 1):
            checkpoint(cancel)
            az, _, length = geod.inv(
                row["Source longitude"],
                row["Source latitude"],
                row["Target longitude"],
                row["Target latitude"],
            )
            count = max(2, int(math.ceil(length / spacing)) + 1)
            budget += count
            if count > 10000 or budget > 200000:
                raise ValueError(
                    "DEM profile limit exceeded (10,000 samples/link or 200,000 total). Increase sample spacing or calculate fewer links."
                )
            distances_m = np.linspace(0, length, count)
            lon, lat, _ = geod.fwd(
                np.full(count, row["Source longitude"]),
                np.full(count, row["Source latitude"]),
                np.full(count, az),
                distances_m,
            )
            x, y = transform.transform(lon, lat, errcheck=True)
            elevations = []
            for n, value in enumerate(
                dem.sample(zip(x, y), indexes=int(band), masked=True)
            ):
                checkpoint(cancel)
                rr, cc = dem.index(x[n], y[n])
                missing = (
                    rr < 0
                    or cc < 0
                    or rr >= dem.height
                    or cc >= dem.width
                    or bool(np.ma.getmaskarray(value)[0])
                    or not np.isfinite(value[0])
                )
                elevations.append(
                    np.nan
                    if missing
                    else (
                        float(value[0]) * dem.scales[int(band) - 1]
                        + dem.offsets[int(band) - 1]
                    )
                    * (0.3048 if units == "international feet" else 1)
                )
            elevations = np.asarray(elevations)
            t = distances_m / length if length > 0 else np.zeros(count)
            sight = (
                elevations[0]
                + source_agl
                + t * (elevations[-1] + target_agl - elevations[0] - source_agl)
            )
            bulge = distances_m * (length - distances_m) / (2 * k_factor * 6371000)
            fresnel = (
                np.sqrt(
                    np.maximum(
                        0,
                        (299792458 / frequency)
                        * distances_m
                        * (length - distances_m)
                        / length,
                    )
                )
                if length > 0
                else np.zeros(count)
            )
            clearance = sight - elevations - bulge - fresnel_fraction * fresnel
            missing = int(np.isnan(elevations).sum())
            status = (
                "Not evaluated: zero-length link"
                if length <= 0.01
                else (
                    "Unknown: missing DEM samples"
                    if missing
                    else (
                        "Obstructed at sampled cells"
                        if np.any(clearance < 0)
                        else "Clear at sampled cells only"
                    )
                )
            )
            common = {
                "Link": link_number,
                "Source site": row["Source site"],
                "Target site": row["Target site"],
            }
            records.append(
                {
                    **common,
                    "Ground distance (m)": length,
                    "Samples": count,
                    "Missing samples": missing,
                    "Minimum clearance (m)": (
                        float(np.min(clearance)) if not missing else np.nan
                    ),
                    "Status": status,
                    "Source AGL (m)": source_agl,
                    "Target AGL (m)": target_agl,
                    "Frequency (GHz)": frequency / 1e9,
                    "Fresnel fraction": fresnel_fraction,
                    "K factor": k_factor,
                }
            )
            from shapely import MultiLineString

            parts = []
            segment = [[float(lon[0]), float(lat[0])]]
            for x1, y1, x2, y2 in zip(lon[:-1], lat[:-1], lon[1:], lat[1:]):
                if abs(x2 - x1) > 180:
                    adjusted = x2 - 360 if x2 > x1 else x2 + 360
                    edge = -180 if adjusted < x1 else 180
                    crossing = y1 + (y2 - y1) * (edge - x1) / (adjusted - x1)
                    segment.append([edge, float(crossing)])
                    parts.append(segment)
                    segment = [[-edge, float(crossing)], [float(x2), float(y2)]]
                else:
                    segment.append([float(x2), float(y2)])
            parts.append(segment)
            geometries.append(MultiLineString(parts) if length > 0.01 else None)
            for n in range(count):
                profiles.append(
                    {
                        **common,
                        "Distance along (m)": distances_m[n],
                        "Longitude": lon[n],
                        "Latitude": lat[n],
                        "DEM elevation (m)": elevations[n],
                        "Earth bulge (m)": bulge[n],
                        "Line of sight (m)": sight[n],
                        "First Fresnel radius (m)": fresnel[n],
                        "Required clearance margin (m)": clearance[n],
                        "Sample valid": bool(np.isfinite(elevations[n])),
                    }
                )
            if progress:
                progress(
                    f"DEM profile {link_number:,} / {len(matched):,}",
                    link_number / len(matched),
                )
    if input_fingerprint(path, cancel) != fingerprint:
        raise ValueError("DEM file changed during analysis.")
    return pd.DataFrame(records), pd.DataFrame(profiles), geometries, metadata


def run_geo_tool(payload, cancel=None, progress=None):
    import shapely
    import pyproj

    tool, options, spec = payload["tool"], payload["options"], payload["spec"]
    if tool not in GEO_TOOLS:
        raise ValueError("Choose a Geo Tools operation.")
    metadata = {
        "Product": TITLE,
        "Version": VERSION,
        "Tool": tool,
        "Created": datetime.now().isoformat(timespec="seconds"),
        "Runtime versions": {
            "Python": sys.version.split()[0],
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "shapely": shapely.__version__,
            "pyproj": pyproj.__version__,
        },
        "Options": options,
        "Input specification": asdict(spec),
        "Distance scope": "Geo Tools settings are independent of point-matching radius/top N/mode",
    }
    layers, tables, crs = {}, {}, "EPSG:4326"
    if tool == GEO_TOOLS[4]:
        if not options.get("dem_confirmed"):
            raise ValueError(
                "Confirm the DEM elevation units and common vertical reference before profiling."
            )
        table, profile, geometry, details = dem_profiles(
            payload["links"],
            options["dem_path"],
            options["spacing"],
            options["source_agl"],
            options["target_agl"],
            options["frequency"],
            options["fresnel"],
            options["k_factor"],
            options["dem_units"],
            options["band"],
            cancel,
            progress,
        )
        metadata["DEM"] = details
        metadata["Matched run"] = payload.get("matched_metadata", {})
        tables["Profile samples"] = profile
        layers["profile_links"] = geometry
    else:
        need_target = tool in (GEO_TOOLS[2], GEO_TOOLS[3], GEO_TOOLS[5])
        counts = tool in (GEO_TOOLS[3], GEO_TOOLS[5])
        spec = replace(
            spec,
            mode=MODES[0],
            terrain=False,
            match_attributes=spec.match_attributes if counts else False,
        )
        frames, audits = [], []
        for name, file in (
            (("Source", spec.source), ("Target", spec.target))
            if need_target
            else (("Source", spec.source),)
        ):
            checkpoint(cancel)
            if progress:
                progress(f"Validating {name.lower()} working copy for Geo Tools")
            valid, audit = validate_working(
                checked_raw(file, cancel, progress), file, spec, cancel
            )
            if payload.get("exclude_flags"):
                valid = valid.loc[valid["Flag"].eq("")].reset_index(drop=True)
                audit.loc[
                    audit["GeoMatch disposition"].eq("Flagged - retained"),
                    "GeoMatch disposition",
                ] = "Excluded by user (outlier)"
            if valid.empty:
                raise ValueError(
                    f"{name} has no eligible points. Review its validation and exclusions."
                )
            frames.append(valid.reset_index(drop=True))
            audits.append(audit)
            tables[name + " review"] = audit
        source, target = frames[0], frames[-1]
        metadata["Applied validation specification"] = asdict(spec)
        metadata["Exclude flagged outliers"] = bool(payload.get("exclude_flags"))
        metadata["Coordinate operations"] = [
            frame.attrs.get("coordinate_operation", {}) for frame in frames
        ]
        if tool in GEO_TOOLS[:2]:
            crs_obj, _, _ = metric_projection(options["crs"], [source])
            geometries, attributes, details = vector_features(
                options["vector_path"],
                options.get("layer", ""),
                crs_obj,
                "line" if tool == GEO_TOOLS[0] else "polygon",
                cancel,
            )
            metadata["Vector input"] = details
            table, geometry, crs = nearest_vector(
                source,
                geometries,
                attributes,
                options["crs"],
                tool == GEO_TOOLS[1],
                cancel,
            )
            layers["nearest_connections"] = geometry
        elif tool == GEO_TOOLS[2]:
            table, geometry, crs = survey_comparison(
                source,
                target,
                options["crs"],
                options["horizontal_tolerance"],
                options["vertical_tolerance"] if options["check_vertical"] else None,
                cancel,
            )
            layers["survey_points"] = geometry
        elif tool == GEO_TOOLS[3]:
            table, geometry, crs = buffers_counts(
                source, target, options["crs"], options["radius"], spec, cancel
            )
            layers["buffers"] = geometry
        else:
            table, geometry, crs = telecom_sectors(
                source,
                target,
                options["crs"],
                options["radius"],
                options["azimuth"],
                options["beamwidth"],
                spec,
                cancel,
            )
            layers["sectors"] = geometry
        metadata["Analysis CRS"] = crs
        metadata["Planar measurements"] = (
            "Metres in selected projected CRS; projection distortion is not corrected. No vertical datum transformation."
        )
    tables = {"Results": table, **tables}
    # WKB transports geometries across spawned workers without GUI/GDAL objects.
    return {
        "tables": tables,
        "layers": {
            name: shapely.to_wkb(geoms).tolist() for name, geoms in layers.items()
        },
        "crs": crs,
        "metadata": metadata,
    }


def json_record(record):
    def clean(value):
        if isinstance(value, dict):
            return {str(k): clean(v) for k, v in value.items()}
        if isinstance(value, (list, tuple)):
            return [clean(v) for v in value]
        if isinstance(value, np.generic):
            value = value.item()
        if pd.isna(value):
            return None
        return value

    return json.dumps(clean(record), ensure_ascii=False, default=str, allow_nan=False)


def export_geo_tool(path, analysis, cancel=None, progress=None, temporary=None):
    path = Path(path)
    temporary = (
        Path(temporary)
        if temporary
        else path.with_name(
            f".geomatch-analysis-{os.getpid()}-{time.time_ns()}{path.suffix}"
        )
    )
    if path.suffix.lower() not in (".xlsx", ".gpkg"):
        raise ValueError("Geo Tools reports support .xlsx or .gpkg.")
    try:
        if path.suffix.lower() == ".xlsx":
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell

            workbook = Workbook(write_only=True)
            tables = {
                **analysis["tables"],
                "Run settings": pd.DataFrame(
                    {
                        "Setting": list(analysis["metadata"]),
                        "Value": [
                            json.dumps(v, ensure_ascii=False, default=str)
                            for v in analysis["metadata"].values()
                        ],
                    }
                ),
            }
            if any(len(frame) > 1048575 for frame in tables.values()):
                raise ValueError(
                    "A report table exceeds Excel's row limit. Export GeoPackage instead."
                )
            try:
                for name, frame in tables.items():
                    worksheet = workbook.create_sheet(name[:31])
                    worksheet.freeze_panes = "A2"
                    headers = []
                    for name in frame.columns:
                        cell = WriteOnlyCell(worksheet, value=str(name))
                        cell.data_type = "s"
                        headers.append(cell)
                    worksheet.append(headers)
                    for n, values in enumerate(
                        frame.itertuples(index=False, name=None)
                    ):
                        if n % 1000 == 0:
                            checkpoint(cancel)
                        cells = []
                        for value in values:
                            if isinstance(value, (dict, list, tuple)):
                                value = json.dumps(
                                    value, ensure_ascii=False, default=str
                                )
                            if pd.isna(value):
                                value = None
                            elif isinstance(value, np.generic):
                                value = value.item()
                            cell = WriteOnlyCell(worksheet, value=value)
                            if isinstance(value, str):
                                cell.data_type = "s"
                            cells.append(cell)
                        worksheet.append(cells)
                workbook.save(temporary)
            finally:
                for sheet in workbook.worksheets:
                    if not sheet.closed:
                        sheet.close()
                    writer = getattr(sheet, "_writer", None)
                    if writer and os.path.exists(writer.out):
                        writer.cleanup()
                workbook.close()
        else:
            import geopandas as gpd
            import shapely
            import sqlite3
            from contextlib import closing

            for name, wkbs in analysis["layers"].items():
                checkpoint(cancel)
                attributes = analysis["tables"]["Results"].copy()
                if len({str(c).casefold() for c in attributes}) != len(
                    attributes.columns
                ):
                    attributes = pd.DataFrame(
                        {
                            "result_record": [
                                json_record(row)
                                for row in attributes.to_dict("records")
                            ]
                        }
                    )
                for col in attributes:
                    if (
                        attributes[col]
                        .map(lambda value: isinstance(value, (dict, list, tuple)))
                        .any()
                    ):
                        attributes[col] = attributes[col].map(
                            lambda value: (
                                json_record(value)
                                if isinstance(value, (dict, list, tuple))
                                else value
                            )
                        )
                frame = gpd.GeoDataFrame(
                    attributes, geometry=shapely.from_wkb(wkbs), crs=analysis["crs"]
                )
                frame.to_file(temporary, layer=name, driver="GPKG", engine="pyogrio")
            with closing(sqlite3.connect(temporary)) as database, database:
                tables = {
                    **analysis["tables"],
                    "Run settings": pd.DataFrame(
                        {
                            "Setting": list(analysis["metadata"]),
                            "Value": [
                                json.dumps(v, ensure_ascii=False, default=str)
                                for v in analysis["metadata"].values()
                            ],
                        }
                    ),
                }
                for name, frame in tables.items():
                    checkpoint(cancel)
                    # JSON preserves arbitrary input column names, casing and types.
                    table_name = name.lower().replace(" ", "_")
                    database.execute(
                        f'CREATE TABLE "{table_name}" (fid INTEGER PRIMARY KEY, record_json TEXT)'
                    )
                    for start in range(0, len(frame), 2000):
                        checkpoint(cancel)
                        records = frame.iloc[start : start + 2000].to_dict("records")
                        database.executemany(
                            f'INSERT INTO "{table_name}" (record_json) VALUES (?)',
                            [(json_record(row),) for row in records],
                        )
                    database.execute(
                        "INSERT INTO gpkg_contents(table_name,data_type,identifier,description) VALUES (?, 'attributes', ?, ?)",
                        (
                            table_name,
                            table_name,
                            "GeoMatch Geo Tools report; records stored as JSON",
                        ),
                    )
        checkpoint(cancel)
        os.replace(temporary, path)
    finally:
        for candidate in (
            temporary,
            Path(str(temporary) + "-wal"),
            Path(str(temporary) + "-shm"),
            Path(str(temporary) + "-journal"),
        ):
            candidate.unlink(missing_ok=True)
    return str(path)


def job_worker(kind, payload, events, cancel):
    def progress(message, fraction=None):
        events.put(("progress", (message, fraction)))

    try:
        if kind == "sheets":
            before = signature(payload)
            with pd.ExcelFile(
                payload,
                engine=(
                    "openpyxl" if Path(payload).suffix.lower() == ".xlsx" else "xlrd"
                ),
            ) as workbook:
                names = workbook.sheet_names
            checkpoint(cancel)
            if signature(payload) != before:
                raise ValueError("The workbook changed while reading. Open it again.")
            value = {"path": payload, "sheets": names, "signature": before}
        elif kind == "inspect":
            path = payload["path"] if isinstance(payload, dict) else payload
            sheet = payload.get("sheet") if isinstance(payload, dict) else None
            before = signature(path)
            if (
                isinstance(payload, dict)
                and payload.get("signature")
                and before != payload["signature"]
            ):
                raise ValueError(
                    "The workbook changed after selecting its worksheet. Open it again."
                )
            fingerprint = input_fingerprint(path, cancel)
            if Path(path).suffix.lower() in (".csv", ".txt"):
                sample = None
                total = 0
                with pd.read_csv(
                    path,
                    sep=None,
                    engine="python",
                    dtype=str,
                    keep_default_na=False,
                    chunksize=10000,
                ) as reader:
                    for chunk in reader:
                        checkpoint(cancel)
                        if sample is None:
                            sample = chunk.head(12).copy()
                        total += len(chunk)
                        progress(f"Inspecting {Path(path).name}: {total:,} rows")
                if not total:
                    raise ValueError("The input contains no data rows.")
                raw = sample
            else:
                raw = read_input(path, cancel, progress, sheet)
                total = len(raw)
            if input_fingerprint(path, cancel) != fingerprint:
                raise ValueError("File contents changed while reading. Import again.")
            if signature(path) != before:
                raise ValueError(
                    "The file changed while it was being read. Import it again."
                )
            value = {
                "headers": list(raw.columns),
                "sample": raw.head(12),
                "rows": total,
                "signature": before,
                "sheet": sheet,
                "sha256": fingerprint,
            }
        elif kind == "validate":
            spec = payload
            spec.validate()
            frames, audits = [], []
            for label, file in (("Source", spec.source), ("Target", spec.target)):
                checkpoint(cancel)
                if signature(file.path) != file.signature:
                    raise ValueError(
                        f"{label} file changed after mapping. Import it again."
                    )
                fingerprint = input_fingerprint(file.path, cancel)
                if file.sha256 and fingerprint != file.sha256:
                    raise ValueError(
                        f"{label} file contents changed after mapping. Import it again."
                    )
                progress(f"Validating {label.lower()} sites…")
                raw = read_input(file.path, cancel, progress, file.sheet)
                if signature(file.path) != file.signature:
                    raise ValueError(
                        f"{label} file changed while reading. Import it again."
                    )
                if input_fingerprint(file.path, cancel) != fingerprint:
                    raise ValueError(
                        f"{label} file contents changed while validating. Import again."
                    )
                valid, audit = validate_working(raw, file, spec, cancel)
                valid.attrs["sha256"] = fingerprint
                frames.append(valid)
                audits.append(audit)
            value = (frames, audits)
        elif kind == "calculate":
            source, target, spec = payload
            value = calculate(source, target, spec, cancel, progress)
        elif kind == "export":
            path, result, audits, metadata, temporary = payload
            export_workbook(path, result, audits, metadata, cancel, progress, temporary)
            value = path
        elif kind == "export_gis":
            path, result, frames, audits, metadata, temporary = payload
            export_gis(
                path, result, frames, audits, metadata, cancel, progress, temporary
            )
            value = path
        elif kind == "working":
            value = checked_raw(payload, cancel, progress)
        elif kind == "save_project":
            value = save_project(
                payload[0],
                payload[1],
                payload[2],
                cancel,
                payload[3] if len(payload) > 3 else None,
            )
        elif kind == "load_project":
            value = (
                *load_project(*payload, cancel=cancel),
                json.loads(Path(payload[0]).read_text(encoding="utf-8")).get(
                    "geo_tools", {}
                ),
            )
        elif kind == "stream_large":
            value = stream_large_run(
                payload[0],
                payload[1],
                cancel=cancel,
                progress=progress,
                temporary=payload[2],
            )
        elif kind == "geo_tool":
            value = run_geo_tool(payload, cancel, progress)
        elif kind == "export_geo_tool":
            value = export_geo_tool(
                payload[0], payload[1], cancel, progress, payload[2]
            )
        elif kind == "single":
            lat, lon, tlat, tlon = payload
            ground, bearing = distances(
                lat, lon, np.array([tlat]), np.array([tlon]), METHODS[0]
            )
            value = (float(ground[0]), float(bearing[0]))
        else:
            raise ValueError("Unknown operation")
        checkpoint(cancel)
        events.put(("done", value))
    except Cancelled:
        events.put(("cancelled", None))
    except Exception as error:
        events.put(("error", (str(error), traceback.format_exc())))


def start_gui(run=True, settings_dir=None):
    """Qt stays in the parent process; the unchanged engine runs in spawned workers."""
    from PySide6.QtCore import (
        Qt,
        QTimer,
        QAbstractTableModel,
        QModelIndex,
        QSortFilterProxyModel,
        QByteArray,
        QUrl,
        QPointF,
        QUrlQuery,
    )
    from PySide6.QtGui import (
        QFontMetrics,
        QFont,
        QIcon,
        QPixmap,
        QPainter,
        QColor,
        QPen,
        QAction,
        QKeySequence,
        QDesktopServices,
        QPalette,
    )
    from PySide6.QtWidgets import (
        QApplication,
        QMainWindow,
        QWidget,
        QVBoxLayout,
        QHBoxLayout,
        QFormLayout,
        QLabel,
        QPushButton,
        QComboBox,
        QLineEdit,
        QSpinBox,
        QCheckBox,
        QSplitter,
        QScrollArea,
        QTabWidget,
        QTableView,
        QAbstractItemView,
        QHeaderView,
        QPlainTextEdit,
        QTextBrowser,
        QProgressBar,
        QFileDialog,
        QDialog,
        QDialogButtonBox,
        QMenu,
        QInputDialog,
        QMessageBox,
        QGroupBox,
        QStyle,
    )

    application = QApplication.instance() or QApplication(sys.argv[:1])
    application.setApplicationName("GeoMatch")
    application.setStyle("Fusion")

    def label(text, name=""):
        item = QLabel(text)
        item.setWordWrap(True)
        item.setTextFormat(Qt.TextFormat.PlainText)
        if name:
            item.setObjectName(name)
        return item

    def button(text, callback, primary=False, icon=None):
        item = QPushButton(text.replace("&", "&&"))
        item.clicked.connect(callback)
        item.setCursor(Qt.CursorShape.PointingHandCursor)
        if primary:
            item.setObjectName("primary")
        if icon is not None:
            item.setIcon(application.style().standardIcon(icon))
        return item

    def combo(values):
        item = QComboBox()
        item.addItems(list(values))
        return item

    def mark():
        pix = QPixmap(64, 64)
        pix.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pix)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setBrush(QColor("#0b7e83"))
        painter.setPen(Qt.PenStyle.NoPen)
        painter.drawRoundedRect(1, 1, 62, 62, 14, 14)
        painter.setPen(QPen(QColor("white"), 4))
        painter.drawLine(17, 43, 47, 21)
        painter.setBrush(QColor("white"))
        painter.drawEllipse(11, 37, 12, 12)
        painter.drawEllipse(41, 15, 12, 12)
        painter.end()
        return QIcon(pix)

    class FrameModel(QAbstractTableModel):
        def __init__(self, parent=None):
            super().__init__(parent)
            self.frame = pd.DataFrame()

        def set_frame(self, frame):
            self.beginResetModel()
            self.frame = frame.reset_index(drop=True)
            self.endResetModel()

        def rowCount(self, parent=QModelIndex()):
            return 0 if parent.isValid() else len(self.frame)

        def columnCount(self, parent=QModelIndex()):
            return 0 if parent.isValid() else len(self.frame.columns)

        def data(self, index, role=Qt.ItemDataRole.DisplayRole):
            if not index.isValid():
                return None
            value = self.frame.iat[index.row(), index.column()]
            if role in (Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.ToolTipRole):
                if pd.isna(value):
                    return ""
                if isinstance(value, (float, np.floating)):
                    if any(
                        axis in str(self.frame.columns[index.column()]).lower()
                        for axis in ("latitude", "longitude")
                    ):
                        return f"{value:.8f}"
                    if str(self.frame.columns[index.column()]) in (
                        "Source row",
                        "Target row",
                        "Rank",
                    ):
                        return str(int(value))
                    if any(
                        term in str(self.frame.columns[index.column()]).lower()
                        for term in ("delta ", "tolerance", "change")
                    ):
                        return f"{value:,.6f}"
                    return f"{value:,.3f}"
                return str(value)
            if role == Qt.ItemDataRole.TextAlignmentRole and isinstance(
                value, (int, float, np.number)
            ):
                return int(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            return None

        def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
            if role in (Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.ToolTipRole):
                return (
                    str(self.frame.columns[section])
                    if orientation == Qt.Orientation.Horizontal
                    else str(section + 1)
                )
            return None

    class FrameProxy(QSortFilterProxyModel):
        def __init__(self, parent=None):
            super().__init__(parent)
            self.query = ""
            self.filters = {}
            self.affected = False

        def filterAcceptsRow(self, row, parent):
            frame = self.sourceModel().frame
            values = frame.iloc[row]
            if self.affected and "GeoMatch disposition" in frame:
                if values["GeoMatch disposition"] == "Retained" and not values.get(
                    "GeoMatch coordinate warning", ""
                ):
                    return False
            for name, query in self.filters.items():
                if name in frame and query not in str(values[name]).casefold():
                    return False
            return not self.query or any(
                self.query in str(v).casefold() for v in values
            )

        def lessThan(self, left, right):
            frame = self.sourceModel().frame
            a, b = (
                frame.iat[left.row(), left.column()],
                frame.iat[right.row(), right.column()],
            )
            if pd.isna(a):
                return False
            if pd.isna(b):
                return True
            if isinstance(a, (int, float, np.number)) and isinstance(
                b, (int, float, np.number)
            ):
                return float(a) < float(b)
            return str(a).casefold() < str(b).casefold()

    class DataTable(QWidget):
        def __init__(self, parent=None, preferences=None, key=""):
            super().__init__(parent)
            self.preferences, self.key = preferences, key
            self.model = FrameModel(self)
            self.proxy = FrameProxy(self)
            self.proxy.setSourceModel(self.model)
            self.view = QTableView()
            self.view.setModel(self.proxy)
            self.view.setSortingEnabled(True)
            self.view.sortByColumn(-1, Qt.SortOrder.AscendingOrder)
            self.view.setAlternatingRowColors(True)
            self.view.setSelectionMode(
                QAbstractItemView.SelectionMode.ExtendedSelection
            )
            self.view.setSelectionBehavior(
                QAbstractItemView.SelectionBehavior.SelectItems
            )
            self.view.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            self.view.setWordWrap(False)
            self.view.verticalHeader().setDefaultSectionSize(32)
            self.view.verticalHeader().hide()
            header = self.view.horizontalHeader()
            header.setSectionsMovable(True)
            header.setDefaultSectionSize(155)
            header.setMinimumSectionSize(55)
            header.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
            header.customContextMenuRequested.connect(self.header_menu)
            layout = QVBoxLayout(self)
            layout.setContentsMargins(0, 0, 0, 0)
            toolbar = QHBoxLayout()
            self.search = QLineEdit()
            self.search.setPlaceholderText("Search all columns…")
            self.search.setClearButtonEnabled(True)
            self.search.setAccessibleName("Search table")
            self.delay = QTimer(self)
            self.delay.setSingleShot(True)
            self.delay.setInterval(180)
            self.delay.timeout.connect(self.apply_search)
            self.search.textChanged.connect(lambda: self.delay.start())
            toolbar.addWidget(self.search, 1)
            toolbar.addWidget(button("Columns", self.columns_menu))
            toolbar.addWidget(button("Copy cells", self.copy_cells))
            toolbar.addWidget(button("Clear filters", self.clear_filters))
            layout.addLayout(toolbar)
            layout.addWidget(self.view, 1)
            self.count = label("No rows loaded", "muted")
            layout.addWidget(self.count)
            action = QAction(self)
            action.setShortcut(QKeySequence.StandardKey.Copy)
            action.setShortcutContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
            action.triggered.connect(self.copy_cells)
            self.addAction(action)
            self.proxy.modelReset.connect(self.update_count)
            self.proxy.rowsInserted.connect(self.update_count)
            self.proxy.rowsRemoved.connect(self.update_count)

        def save_layout(self):
            if self.preferences is not None and self.model.columnCount():
                header = self.view.horizontalHeader()
                self.preferences[self.key] = {
                    "order": [
                        str(self.model.frame.columns[header.logicalIndex(i)])
                        for i in range(header.count())
                    ],
                    "widths": {
                        str(c): self.view.columnWidth(i)
                        for i, c in enumerate(self.model.frame.columns)
                    },
                    "hidden": [
                        str(c)
                        for i, c in enumerate(self.model.frame.columns)
                        if self.view.isColumnHidden(i)
                    ],
                }

        def set_frame(self, frame):
            self.save_layout()
            self.model.set_frame(frame)
            self.proxy.invalidate()
            settings = (self.preferences or {}).get(self.key, {})
            columns = list(frame.columns)
            header = self.view.horizontalHeader()
            for dest, name in enumerate(
                c for c in settings.get("order", []) if c in columns
            ):
                header.moveSection(header.visualIndex(columns.index(name)), dest)
            for i, name in enumerate(columns):
                self.view.setColumnWidth(
                    i, max(55, min(900, settings.get("widths", {}).get(name, 155)))
                )
                self.view.setColumnHidden(i, name in settings.get("hidden", []))
            if columns and all(
                self.view.isColumnHidden(i) for i in range(len(columns))
            ):
                self.view.setColumnHidden(0, False)
            self.fit_headers()
            self.update_count()

        def fit_headers(self):
            font = QFont(application.font())
            font.setWeight(QFont.Weight.DemiBold)
            metrics = QFontMetrics(font)
            self.view.verticalHeader().setDefaultSectionSize(
                max(32, metrics.height() + 12)
            )
            for i, name in enumerate(self.model.frame.columns):
                if not self.view.isColumnHidden(i):
                    self.view.setColumnWidth(
                        i,
                        max(
                            self.view.columnWidth(i),
                            min(500, metrics.horizontalAdvance(str(name)) + 32),
                        ),
                    )

        def update_count(self, *args):
            filters = "; ".join(f"{k}: {v}" for k, v in self.proxy.filters.items())
            self.count.setText(
                f"{self.proxy.rowCount():,} shown / {self.model.rowCount():,} rows"
                + (f" · Filters: {filters}" if filters else "")
                + (" · Affected rows only" if self.proxy.affected else "")
                + " · Right-click a column heading to filter"
            )

        def apply_search(self):
            self.proxy.query = self.search.text().strip().casefold()
            self.proxy.invalidate()
            self.update_count()

        def clear_filters(self):
            self.delay.stop()
            self.search.clear()
            self.proxy.query = ""
            self.proxy.filters.clear()
            self.proxy.affected = False
            self.proxy.invalidate()
            self.update_count()

        def affected_rows(self):
            self.clear_filters()
            self.proxy.affected = True
            self.proxy.invalidate()
            self.update_count()

        def header_menu(self, point):
            column = self.view.horizontalHeader().logicalIndexAt(point)
            if column < 0:
                return
            name = str(self.model.frame.columns[column])
            text, ok = QInputDialog.getText(
                self,
                "Column filter",
                f"{name} contains (empty clears this filter):",
                text=self.proxy.filters.get(name, ""),
            )
            if ok:
                if text.strip():
                    self.proxy.filters[name] = text.strip().casefold()
                else:
                    self.proxy.filters.pop(name, None)
                self.proxy.invalidate()
                self.update_count()

        def columns_menu(self):
            menu = QMenu(self)
            for i, name in enumerate(self.model.frame.columns):
                action = menu.addAction(str(name))
                action.setCheckable(True)
                action.setChecked(not self.view.isColumnHidden(i))
                action.triggered.connect(
                    lambda checked, index=i: self.view.setColumnHidden(
                        index, not checked
                    )
                )
            menu.addSeparator()
            menu.addAction(
                "Show all columns",
                lambda: [
                    self.view.setColumnHidden(i, False)
                    for i in range(self.model.columnCount())
                ],
            )
            menu.exec(self.mapToGlobal(self.rect().center()))

        def copy_cells(self):
            selected = [
                i
                for i in self.view.selectionModel().selectedIndexes()
                if not self.view.isColumnHidden(i.column())
            ]
            if not selected:
                return
            cells = {(i.row(), i.column()) for i in selected}
            rows = sorted({i.row() for i in selected})
            columns = sorted(
                {i.column() for i in selected},
                key=self.view.horizontalHeader().visualIndex,
            )
            # Preserve raw numeric precision and ID strings in clipboard output.
            records = [
                [
                    (
                        self.model.frame.iat[
                            self.proxy.mapToSource(self.proxy.index(r, c)).row(), c
                        ]
                        if (r, c) in cells
                        else ""
                    )
                    for c in columns
                ]
                for r in rows
            ]
            application.clipboard().setText(
                pd.DataFrame(
                    records, columns=[self.model.frame.columns[c] for c in columns]
                ).to_csv(sep="\t", index=False)
            )

    class TerrainPlot(QWidget):
        def __init__(self):
            super().__init__()
            self.setMinimumHeight(200)
            self.frame = pd.DataFrame()
            self.setToolTip(
                "Sampled terrain plus Earth bulge and the required Fresnel fraction is compared with the straight antenna-to-antenna line. Gaps mean unavailable raster samples. This is not a complete obstruction or propagation survey."
            )

        def set_frame(self, frame):
            self.frame = frame
            self.update()

        def paintEvent(self, event):
            p = QPainter(self)
            p.setRenderHint(QPainter.RenderHint.Antialiasing)
            p.fillRect(self.rect(), self.palette().base())
            p.setPen(self.palette().text().color())
            if self.frame.empty:
                p.drawText(
                    12, 25, "Select a profiled link to view its terrain cross-section."
                )
                return
            x = self.frame["Distance along (m)"].to_numpy(float)
            terrain = self.frame["DEM elevation (m)"].to_numpy(float)
            sight = self.frame["Line of sight (m)"].to_numpy(float)
            required = sight - self.frame["Required clearance margin (m)"].to_numpy(
                float
            )
            finite = np.concatenate(
                [a[np.isfinite(a)] for a in (terrain, sight, required)]
            )
            p.drawText(
                12,
                20,
                "Terrain (brown) · terrain + curvature + required Fresnel (orange) · sight line (blue)",
            )
            if not len(finite):
                p.drawText(
                    12, 55, "No valid elevations. Check DEM extent, band and NoData."
                )
                return
            low, high = float(finite.min()), float(finite.max())
            padding = max(1, (high - low) * 0.1)
            low -= padding
            high += padding
            width, height = max(1, self.width() - 105), max(1, self.height() - 85)
            point = lambda xx, yy: QPointF(
                70 + xx / max(1, float(x[-1])) * width,
                35 + (high - yy) / (high - low) * height,
            )
            p.setPen(QPen(self.palette().mid().color(), 1))
            p.drawLine(point(0, low), point(x[-1], low))
            p.drawLine(point(0, low), point(0, high))
            for values, color in (
                (terrain, "#a67646"),
                (required, "#e68a25"),
                (sight, "#438fd5"),
            ):
                p.setPen(QPen(QColor(color), 2))
                previous = None
                for xx, yy in zip(x, values):
                    if not math.isfinite(yy):
                        previous = None
                        continue
                    current = point(xx, yy)
                    if previous is not None:
                        p.drawLine(previous, current)
                    previous = current
            p.setPen(self.palette().text().color())
            p.drawText(4, 43, f"{high:.1f} m")
            p.drawText(4, 35 + height, f"{low:.1f} m")
            p.drawText(
                70,
                self.height() - 15,
                f"0 m                       Distance along link → {x[-1]:,.1f} m",
            )

    class WorkingModel(FrameModel):
        def __init__(self, working, changed, parent=None):
            super().__init__(parent)
            self.working, self.changed = working, changed
            self.set_frame(working.frame)

        def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
            if (
                orientation == Qt.Orientation.Vertical
                and role == Qt.ItemDataRole.DisplayRole
            ):
                return str(section + 2)
            return super().headerData(section, orientation, role)

        def flags(self, index):
            return super().flags(index) | Qt.ItemFlag.ItemIsEditable

        def data(self, index, role=Qt.ItemDataRole.DisplayRole):
            if index.isValid() and role == Qt.ItemDataRole.EditRole:
                return str(self.working.frame.iat[index.row(), index.column()])
            if index.isValid() and role == Qt.ItemDataRole.BackgroundRole:
                if index.row() in self.working.excluded:
                    return QColor("#f3c4b6")
                if (
                    index.row(),
                    self.frame.columns[index.column()],
                ) in self.working.edits:
                    return QColor("#b8e3d9")
            if role == Qt.ItemDataRole.ForegroundRole:
                return (
                    QColor("#172b38")
                    if index.row() in self.working.excluded
                    or (index.row(), self.frame.columns[index.column()])
                    in self.working.edits
                    else None
                )
            return super().data(index, role)

        def setData(self, index, value, role=Qt.ItemDataRole.EditRole):
            if role != Qt.ItemDataRole.EditRole or not index.isValid():
                return False
            self.working.edit(
                index.row(), str(self.frame.columns[index.column()]), str(value)
            )
            self.changed()
            return True

    class WorkingDialog(QDialog):
        def __init__(self, parent, role, raw, file):
            super().__init__(parent)
            self.setWindowTitle(
                f"{role.title()} working copy — originals stay unchanged"
            )
            self.resize(1000, 650)
            self.working = WorkingCopy(raw, file)
            layout = QVBoxLayout(self)
            layout.addWidget(
                label(
                    "Double-click a cell to edit. Green = edited; orange = excluded. Row numbers correspond to input data rows (header is row 1). Apply changes, then validate again. Cancel discards this dialog's changes."
                )
            )
            self.table = DataTable()
            self.table.model = WorkingModel(self.working, self.refresh, self.table)
            self.table.proxy.setSourceModel(self.table.model)
            self.table.view.setEditTriggers(
                QAbstractItemView.EditTrigger.DoubleClicked
                | QAbstractItemView.EditTrigger.EditKeyPressed
            )
            self.table.view.verticalHeader().show()
            actions = QHBoxLayout()
            self.undo_button = button("Undo", lambda: self.history(False))
            self.redo_button = button("Redo", lambda: self.history(True))
            for widget in (
                self.undo_button,
                self.redo_button,
                button("Exclude selected rows", lambda: self.exclude(True)),
                button("Include selected rows", lambda: self.exclude(False)),
            ):
                actions.addWidget(widget)
            layout.addLayout(actions)
            layout.addWidget(self.table, 1)
            self.summary = label("")
            layout.addWidget(self.summary)
            buttons = QDialogButtonBox(
                QDialogButtonBox.StandardButton.Apply
                | QDialogButtonBox.StandardButton.Cancel
            )
            buttons.button(QDialogButtonBox.StandardButton.Apply).clicked.connect(
                self.accept
            )
            buttons.rejected.connect(self.reject)
            layout.addWidget(buttons)
            for key, callback in (
                ("Ctrl+Z", lambda: self.history(False)),
                ("Ctrl+Y", lambda: self.history(True)),
            ):
                action = QAction(self)
                action.setShortcut(QKeySequence(key))
                action.triggered.connect(callback)
                self.addAction(action)
            self.refresh()

        def refresh(self):
            self.table.model.set_frame(self.working.frame)
            self.table.proxy.invalidate()
            self.table.update_count()
            self.undo_button.setEnabled(bool(self.working.undo_stack))
            self.redo_button.setEnabled(bool(self.working.redo_stack))
            self.summary.setText(
                f"{len(self.working.edits):,} edited cells · {len(self.working.excluded):,} excluded rows · Changes are saved with a project"
            )

        def history(self, redo):
            (self.working.redo if redo else self.working.undo)()
            self.refresh()

        def exclude(self, excluded):
            rows = {
                self.table.proxy.mapToSource(i).row()
                for i in self.table.view.selectionModel().selectedIndexes()
            }
            self.working.exclude(rows, excluded)
            self.refresh()

    class SpatialPreview(QWidget):
        """Offline equirectangular display, sampled for responsive interaction."""

        def __init__(self, table):
            super().__init__()
            self.table = table
            self.setMinimumHeight(180)
            self.setToolTip(
                "Offline coordinate diagram, not a basemap or distance scale. Up to 2,000 evenly sampled result rows plus 2,000 selected rows are drawn. Table filters do not filter the diagram; filtered-out rows cannot be selected from the diagram."
            )
            self.setAccessibleName(
                "Offline spatial preview. Select result rows to highlight links. Wheel to zoom; drag to pan."
            )
            self.frame = pd.DataFrame()
            self.selected = set()
            self.zoom = 1.0
            self.pan = QPointF(0, 0)
            self.drag = None
            self.hit = []
            table.view.selectionModel().selectionChanged.connect(self.selection_changed)

        def set_frame(self, frame):
            self.frame = frame
            self.selected.clear()
            self.extent = None
            if not frame.empty:
                self.centre = float(frame.iloc[0]["Source longitude"])
                xs = np.concatenate(
                    [
                        frame["Source longitude"].to_numpy(float),
                        frame["Target longitude"].dropna().to_numpy(float),
                    ]
                )
                ys = np.concatenate(
                    [
                        frame["Source latitude"].to_numpy(float),
                        frame["Target latitude"].dropna().to_numpy(float),
                    ]
                )
                xs = (xs - self.centre + 180) % 360 - 180
                self.extent = (
                    float(xs.min()),
                    float(xs.max()),
                    float(-ys.max()),
                    float(-ys.min()),
                )
            self.fit()

        def fit(self):
            self.zoom = 1.0
            self.pan = QPointF(0, 0)
            self.update()

        def selection_changed(self, *args):
            self.selected = {
                self.table.proxy.mapToSource(i).row()
                for i in self.table.view.selectionModel().selectedIndexes()
            }
            self.update()

        def paintEvent(self, event):
            p = QPainter(self)
            p.setRenderHint(QPainter.RenderHint.Antialiasing)
            p.fillRect(self.rect(), self.palette().base())
            p.setPen(self.palette().text().color())
            p.drawText(
                12, 20, "WGS84 diagram · blue source / teal target / orange selection"
            )
            self.hit = []
            if self.frame.empty:
                p.drawText(12, 50, "Run matching to preview points and connections.")
                return
            sampled = np.linspace(
                0, len(self.frame) - 1, min(2000, len(self.frame)), dtype=int
            )
            ids = sorted(set(sampled) - self.selected) + sorted(self.selected)[:2000]
            data = self.frame.iloc[ids]
            # Centre wrapping on the first source so short date-line links stay short.
            centre = float(self.frame.iloc[0]["Source longitude"])
            wrap = lambda x: (float(x) - centre + 180) % 360 - 180
            xmin, xmax, ymin, ymax = self.extent
            scale = (
                min(
                    max(1, self.width() - 60) / max(xmax - xmin, 0.001),
                    max(1, self.height() - 80) / max(ymax - ymin, 0.001),
                )
                * self.zoom
            )

            def point(x, y):
                return QPointF(
                    self.width() / 2
                    + (wrap(x) - (xmin + xmax) / 2) * scale
                    + self.pan.x(),
                    (self.height() + 30) / 2
                    + (-y - (ymin + ymax) / 2) * scale
                    + self.pan.y(),
                )

            for idx, values in zip(ids, data.itertuples(index=False, name=None)):
                r = dict(zip(data.columns, values))
                a = point(r["Source longitude"], r["Source latitude"])
                chosen = idx in self.selected
                color = QColor("#ed8c28" if chosen else "#78949b")
                if pd.notna(r["Target longitude"]):
                    b = point(r["Target longitude"], r["Target latitude"])
                    # Do not draw a connector across the preview's wrapping seam.
                    if (
                        abs(wrap(r["Target longitude"]) - wrap(r["Source longitude"]))
                        <= 180
                    ):
                        p.setPen(QPen(color, 2.5 if chosen else 1))
                        p.drawLine(a, b)
                    p.setPen(Qt.PenStyle.NoPen)
                    p.setBrush(QColor("#ed8c28" if chosen else "#168a82"))
                    p.drawEllipse(b, 5 if chosen else 3, 5 if chosen else 3)
                    self.hit.append((b, idx))
                p.setPen(Qt.PenStyle.NoPen)
                p.setBrush(QColor("#ed8c28" if chosen else "#387ad0"))
                p.drawEllipse(a, 5 if chosen else 3, 5 if chosen else 3)
                self.hit.append((a, idx))
            p.setPen(self.palette().text().color())
            p.drawText(
                12,
                self.height() - 10,
                f"{len(ids):,} / {len(self.frame):,} rows drawn · Wheel: zoom · Drag: pan · Click point: select row",
            )

        def wheelEvent(self, event):
            self.zoom = max(
                0.2,
                min(200, self.zoom * (1.2 if event.angleDelta().y() > 0 else 1 / 1.2)),
            )
            self.update()
            event.accept()

        def mousePressEvent(self, event):
            if event.button() == Qt.MouseButton.LeftButton:
                self.drag = self.press = event.position()

        def mouseMoveEvent(self, event):
            if self.drag is not None:
                self.pan += event.position() - self.drag
                self.drag = event.position()
                self.update()

        def mouseReleaseEvent(self, event):
            if (
                self.drag is not None
                and (event.position() - self.press).manhattanLength() < 5
                and self.hit
            ):
                pos, row = min(
                    self.hit, key=lambda v: (v[0] - event.position()).manhattanLength()
                )
                if (pos - event.position()).manhattanLength() < 16:
                    index = self.table.proxy.mapFromSource(
                        self.table.model.index(row, 0)
                    )
                    if index.isValid():
                        self.table.view.selectRow(index.row())
                        self.table.view.scrollTo(index)
            self.drag = None

    class MappingDialog(QDialog):
        def __init__(self, parent, role, path, info):
            super().__init__(parent)
            self.setWindowTitle(f"Map {role} columns — {Path(path).name}")
            area = application.primaryScreen().availableGeometry()
            self.resize(
                min(960, int(area.width() * 0.92)), min(820, int(area.height() * 0.9))
            )
            self.info = info
            self.path = path
            self.preview_signature = None
            shell = QVBoxLayout(self)
            scroll = QScrollArea()
            scroll.setWidgetResizable(True)
            content = QWidget()
            scroll.setWidget(content)
            shell.addWidget(scroll, 1)
            outer = QVBoxLayout(content)
            outer.addWidget(
                label(
                    f"{info['rows']:,} rows · {info.get('sheet') or 'File preview'}",
                    "heading",
                )
            )
            outer.addWidget(
                label(
                    "Choose each coordinate/height field once. Match attribute may reuse the ID column. Heights are metres; generic height columns need your explicit choice."
                )
            )
            crs_row = QHBoxLayout()
            self.crs = QLineEdit("EPSG:4326")
            self.crs.setPlaceholderText("EPSG code, e.g. EPSG:32636")
            self.crs.setAccessibleName("Input coordinate reference system")
            crs_row.addWidget(label("Input CRS"))
            crs_row.addWidget(self.crs, 1)
            self.crs_search = button("Search EPSG", self.find_crs)
            crs_row.addWidget(self.crs_search)
            outer.addLayout(crs_row)
            outer.addWidget(
                label(
                    "Map X to easting/longitude and Y to northing/latitude. Geographic fields accept decimal degrees or DMS. Projected fields use their CRS units; heights are never transformed.",
                    "muted",
                )
            )
            if Path(path).suffix.lower() == ".geojson":
                self.crs.setReadOnly(True)
                self.crs_search.setEnabled(False)
            form = QFormLayout()
            self.choices = {}
            suggestions, conflicts = suggest_mapping(info["headers"])
            for field in FIELDS:
                widget = combo(["(not mapped)"] + info["headers"])
                widget.setCurrentText(suggestions[field] or "(not mapped)")
                self.choices[field] = widget
                display = {
                    "Latitude": "Y / latitude / northing",
                    "Longitude": "X / longitude / easting",
                }.get(field, field)
                form.addRow(display + (" *" if field in FIELDS[:3] else ""), widget)
            outer.addLayout(form)
            if conflicts:
                outer.addWidget(
                    label("Ambiguous columns: " + "; ".join(conflicts), "error")
                )
            self.preview = DataTable(self)
            self.preview.setMinimumHeight(210)
            self.preview.set_frame(info["sample"])
            outer.addWidget(self.preview, 1)
            outer.addWidget(
                button("Preview converted coordinates", self.preview_coordinates)
            )
            self.confirm_crs = QCheckBox(
                "I confirm the CRS, X/Y mapping and coordinate preview"
            )
            self.preview_note = label("", "muted")
            shell.addWidget(self.preview_note)
            shell.addWidget(self.confirm_crs)
            self.error_label = label("", "error")
            shell.addWidget(self.error_label)
            buttons = QDialogButtonBox(
                QDialogButtonBox.StandardButton.Ok
                | QDialogButtonBox.StandardButton.Cancel
            )
            buttons.button(QDialogButtonBox.StandardButton.Ok).setText("Use mapping")
            buttons.accepted.connect(self.accept_mapping)
            buttons.rejected.connect(self.reject)
            shell.addWidget(buttons)

        def find_crs(self):
            query, ok = QInputDialog.getText(
                self,
                "Find coordinate system",
                "EPSG code or name (for example WGS 84 UTM 36N)",
            )
            if not ok:
                return
            try:
                entries = search_crs(query)
                if not entries:
                    raise ValueError("No matching CRS. Refine the code or name.")
                labels = [f"{code} — {name}" for code, name in entries]
                chosen, ok = QInputDialog.getItem(
                    self,
                    "Choose CRS",
                    "Matches (first 100). Confirm the correct zone and datum:",
                    labels,
                    0,
                    False,
                )
                if ok:
                    self.crs.setText(entries[labels.index(chosen)][0])
            except ValueError as error:
                self.error_label.setText(str(error))

        def preview_coordinates(self):
            mapping = {
                k: ("" if v.currentIndex() == 0 else v.currentText())
                for k, v in self.choices.items()
            }
            try:
                validate_mapping(self.info["headers"], mapping)
                crs = horizontal_crs(self.crs.text())
                file = FileSpec(
                    self.path,
                    tuple(mapping.items()),
                    tuple(self.info["signature"]),
                    crs=crs.to_string(),
                )
                lat, lon, reasons, warnings, operation = normalize_coordinates(
                    self.info["sample"], mapping, file
                )
                preview = self.info["sample"].copy()
                preview["Preview WGS84 latitude"] = lat
                preview["Preview WGS84 longitude"] = lon
                preview["Coordinate check"] = [
                    reason or warning or "Valid"
                    for reason, warning in zip(reasons, warnings)
                ]
                shown = [
                    "Preview WGS84 latitude",
                    "Preview WGS84 longitude",
                    "Coordinate check",
                ]
                self.preview.set_frame(
                    preview[shown + [c for c in preview if c not in shown]]
                )
                self.preview_signature = (self.crs.text(), tuple(mapping.items()))
                self.confirm_crs.setChecked(False)
                self.error_label.clear()
                self.preview_note.setText(
                    f"{crs.name} → WGS84. Preview checked {len(preview)} rows; full validation follows. Review converted coordinates and warnings before confirming."
                )
            except ValueError as error:
                self.preview_signature = None
                self.preview_note.clear()
                self.error_label.setText(str(error))

        def accept_mapping(self):
            self.mapping = {
                k: ("" if v.currentIndex() == 0 else v.currentText())
                for k, v in self.choices.items()
            }
            try:
                validate_mapping(self.info["headers"], self.mapping)
                self.crs_value = horizontal_crs(self.crs.text()).to_string()
                if (
                    self.preview_signature
                    != (self.crs.text(), tuple(self.mapping.items()))
                    or not self.confirm_crs.isChecked()
                ):
                    raise ValueError(
                        "Preview the current coordinates and confirm the CRS/mapping before continuing."
                    )
            except ValueError as error:
                self.error_label.setText(str(error))
                return
            self.accept()

    class App(QMainWindow):
        def __init__(self):
            super().__init__()
            self.setWindowTitle(TITLE)
            self.setWindowIcon(mark())
            self.settings_dir = (
                Path(settings_dir)
                if settings_dir is not None
                else Path(os.environ.get("LOCALAPPDATA", str(Path.home() / ".config")))
                / "SiteSpan"
            )
            self.settings_dir.mkdir(parents=True, exist_ok=True)
            self.settings_path = self.settings_dir / "settings.json"
            try:
                self.settings = json.loads(
                    self.settings_path.read_text(encoding="utf-8")
                )
                if not isinstance(self.settings, dict):
                    self.settings = {}
            except (OSError, ValueError):
                self.settings = {}
            # Migrate saved audit-column layouts to the new display headings.
            self.layouts = self.settings.setdefault("qt_tables", {})
            review_layout = self.layouts.get("review", {})
            for key in ("order", "hidden"):
                if key in review_layout:
                    review_layout[key] = [
                        name.replace("SiteSpan ", "GeoMatch ", 1)
                        for name in review_layout[key]
                    ]
            if "widths" in review_layout:
                review_layout["widths"] = {
                    name.replace("SiteSpan ", "GeoMatch ", 1): width
                    for name, width in review_layout["widths"].items()
                }
            self.layouts.setdefault(
                "results",
                {
                    "order": [
                        "Source site",
                        "Target site",
                        "Distance (m)",
                        "Bearing (deg)",
                        "Match status",
                        "Rank",
                        "Terrain status",
                        "Ground distance (m)",
                        "Distance mode",
                        "Source row",
                        "Target row",
                    ],
                    "hidden": [
                        "Source row",
                        "Target row",
                        "Ground distance (m)",
                        "Distance mode",
                        "Terrain status",
                    ],
                    "widths": {
                        "Source site": 105,
                        "Target site": 105,
                        "Distance (m)": 120,
                        "Bearing (deg)": 110,
                        "Match status": 210,
                        "Rank": 65,
                    },
                },
            )
            self.layouts.setdefault(
                "review",
                {
                    "order": [
                        "GeoMatch input row",
                        "GeoMatch disposition",
                        "GeoMatch reason",
                    ],
                    "widths": {
                        "GeoMatch input row": 125,
                        "GeoMatch disposition": 155,
                        "GeoMatch reason": 310,
                    },
                    "hidden": [],
                },
            )
            self.context = mp.get_context("spawn")
            self.process = self.events = self.cancel = self.pending = self.callback = (
                None
            )
            self.stopping = self.closing = False
            self.started = 0
            self.kind = ""
            self.export_temporary = self.last_export = None
            self.files = {"source": None, "target": None}
            self.frames = self.audits = self.run_spec = None
            self.result = self.result_spec = self.result_audits = None
            self.result_frames = None
            self.project_dirty = False
            self.project_active = False
            self.project_exclude_flags = None
            self.single_value = None
            self.inputs = []
            self.clipboard_store = tempfile.TemporaryDirectory(prefix="geomatch-paste-")
            self.setAcceptDrops(True)
            self.build()
            self.restore_layout()
            self.apply_theme()
            self.timer = QTimer(self)
            self.timer.timeout.connect(self.poll)
            self.timer.start(80)
            self.refresh_summary()
            self.set_busy(False)
            self.update_steps()
            self.log(
                f"GeoMatch {VERSION} ready. Confirm input CRS and X/Y mapping; outlier exclusion off."
            )

        def build(self):
            container = QWidget()
            self.setCentralWidget(container)
            outer = QVBoxLayout(container)
            outer.setContentsMargins(18, 12, 18, 12)
            header = QHBoxLayout()
            logo = QLabel()
            logo.setPixmap(self.windowIcon().pixmap(38, 38))
            header.addWidget(logo)
            titles = QVBoxLayout()
            titles.addWidget(label("GeoMatch", "brand"))
            titles.addWidget(label("Spatial Matching & Proximity Analysis", "muted"))
            header.addLayout(titles, 1)
            self.theme = combo(["Light", "Dark", "System"])
            self.theme.setCurrentText(self.settings.get("theme", "Light"))
            self.theme.setAccessibleName("Theme")
            self.scale = combo(["100%", "125%", "150%"])
            self.scale.setCurrentText(self.settings.get("scale", "100%"))
            self.scale.setAccessibleName("Text size")
            header.addWidget(self.theme)
            header.addWidget(self.scale)
            self.details_toggle = QCheckBox("Record details")
            self.details_toggle.setChecked(self.settings.get("qt_details", False))
            header.addWidget(self.details_toggle)
            self.help_button = QPushButton("Help & Support")
            self.help_button.setAccessibleName("Help and Support menu")
            self.help_button.setCursor(Qt.CursorShape.PointingHandCursor)
            self.help_menu = QMenu(self.help_button)
            for text, callback in (
                ("About", self.about),
                ("Quick Guide", self.quick_guide),
                ("Contact us / Support", self.contact_support),
            ):
                self.help_menu.addAction(text, callback)
            self.help_button.setMenu(self.help_menu)
            header.addWidget(self.help_button)
            outer.addLayout(header)
            steps = QHBoxLayout()
            self.steps = []
            for text, callback in (
                ("1  Load", lambda: self.select_tab(0)),
                ("2  Review", lambda: self.select_tab(1)),
                ("3  Match", self.primary_action),
                ("4  Export", self.export),
            ):
                b = button(text, callback)
                steps.addWidget(b)
                self.steps.append(b)
            outer.addLayout(steps)
            self.splitter = QSplitter(Qt.Orientation.Horizontal)
            self.side = QScrollArea()
            self.side.setWidgetResizable(True)
            self.side.setMinimumWidth(285)
            settings = QWidget()
            self.side.setWidget(settings)
            left = QVBoxLayout(settings)
            self.file_labels = {}
            for role, title in (
                ("source", "Source · sites to connect"),
                ("target", "Target · candidate sites"),
            ):
                group = QGroupBox(title)
                layout = QVBoxLayout(group)
                self.file_labels[role] = label("No file selected", "muted")
                layout.addWidget(self.file_labels[role])
                actions = QHBoxLayout()
                actions.addWidget(
                    button(
                        "Browse & map",
                        lambda checked=False, r=role: self.browse(r),
                        icon=QStyle.StandardPixmap.SP_DirOpenIcon,
                    )
                )
                actions.addWidget(
                    button("Recent", lambda checked=False, r=role: self.recent(r))
                )
                layout.addLayout(actions)
                layout.addWidget(
                    button(
                        "Paste table from Excel",
                        lambda checked=False, r=role: self.paste_table(r),
                    )
                )
                layout.addWidget(
                    button(
                        "Edit working copy",
                        lambda checked=False, r=role: self.edit_working(r),
                    )
                )
                left.addWidget(group)
            left.addWidget(button("Swap source / target", self.swap_files))
            group = QGroupBox("Matching settings")
            form = QFormLayout(group)
            form.setRowWrapPolicy(QFormLayout.RowWrapPolicy.WrapAllRows)
            self.match_mode = combo(MATCH_MODES)
            self.match_attributes = QCheckBox("Require equal match attributes")
            self.match_attributes.setToolTip(
                "Map Match attribute in both files. Exact, case-sensitive text after trimming; blanks are rejected."
            )
            self.method = combo(METHODS)
            self.mode = combo(MODES)
            self.top_n = QSpinBox()
            self.top_n.setRange(1, 10)
            self.radius = QLineEdit("5")
            self.radius_unit = combo(["km", "m"])
            self.unlimited = QCheckBox("Unlimited distance")
            self.unlimited.setChecked(True)
            self.ignore = QCheckBox("Ignore Identical locations (Ignore 0 distance) ")
            self.flag = QCheckBox("Flag outliers for review")
            self.terrain = QCheckBox("Terrain check (online)")
            for title, widget in (
                ("Matching mode", self.match_mode),
                ("Method", self.method),
                ("Distance mode", self.mode),
                ("Targets per source", self.top_n),
            ):
                form.addRow(title, widget)
            distance_row = QHBoxLayout()
            distance_row.addWidget(self.radius, 1)
            distance_row.addWidget(self.radius_unit)
            form.addRow("Maximum distance", distance_row)
            for widget in (
                self.unlimited,
                self.ignore,
                self.flag,
                self.terrain,
                self.match_attributes,
            ):
                form.addRow(widget)
            left.addWidget(group)
            self.settings_summary = label("", "muted")
            left.addWidget(self.settings_summary)
            left.addWidget(
                label(
                    "Input CRS is confirmed during mapping. Calculations use WGS84. Heights in metres with a common vertical reference. 3D requires absolute elevation, or terrain elevation + antenna AGL.",
                    "muted",
                )
            )
            presets = QHBoxLayout()
            self.preset = combo(list(self.settings.get("presets", {})) or ["Default"])
            presets.addWidget(self.preset, 1)
            presets.addWidget(button("Load", self.load_preset))
            presets.addWidget(button("Save", self.save_preset))
            left.addLayout(presets)
            left.addWidget(button("Save sample files", self.samples))
            left.addStretch()
            self.splitter.addWidget(self.side)
            self.tabs = QTabWidget()
            self.splitter.addWidget(self.tabs)
            load = QWidget()
            ll = QVBoxLayout(load)
            ll.addWidget(label("Load and map your site files", "heading"))
            project_actions = QHBoxLayout()
            for text, callback in (
                ("Open project", self.open_project_dialog),
                ("Save project", self.save_project_dialog),
                ("Large-file run to ZIP", self.large_run_dialog),
            ):
                project_actions.addWidget(button(text, callback))
            project_actions.addStretch()
            ll.addLayout(project_actions)
            ll.addWidget(
                label(
                    "Choose source and target on the left. Preview the columns, then validate the records before matching. Your original files are not changed."
                )
            )
            ll.addWidget(
                label(
                    "Drop files here or paste a table with headers from Excel. Supports CSV, TXT, XLSX, XLS and WGS84 point GeoJSON; select a CRS for each tabular input.",
                    "muted",
                )
            )
            self.load_summary = label("No files mapped yet. Start with Browse & map.")
            ll.addWidget(self.load_summary)
            self.preview_table = DataTable(preferences=self.layouts, key="preview")
            ll.addWidget(self.preview_table, 1)
            self.tabs.addTab(load, "Load")
            review = QWidget()
            rl = QVBoxLayout(review)
            self.review_summary = label(
                "Validate both files to review their records.", "heading"
            )
            rl.addWidget(self.review_summary)
            controls = QHBoxLayout()
            self.review_role = combo(["Source", "Target"])
            controls.addWidget(self.review_role)
            controls.addWidget(button("Show affected rows", self.show_affected))
            controls.addStretch()
            rl.addLayout(controls)
            self.exclude_flags = QCheckBox("Exclude flagged outliers in this run")
            rl.addWidget(self.exclude_flags)
            self.review_table = DataTable(preferences=self.layouts, key="review")
            rl.addWidget(self.review_table, 1)
            self.tabs.addTab(review, "Review")
            results = QWidget()
            res = QVBoxLayout(results)
            self.result_summary = label(
                "Results will appear after matching reviewed data.", "heading"
            )
            res.addWidget(self.result_summary)
            self.stale = label("", "error")
            self.stale.hide()
            res.addWidget(self.stale)
            actions = QHBoxLayout()
            self.export_button = button(
                "Export report / GIS",
                self.export,
                True,
                QStyle.StandardPixmap.SP_DialogSaveButton,
            )
            self.folder_button = button("Open export folder", self.open_folder)
            self.export_button.setEnabled(False)
            self.folder_button.setEnabled(False)
            actions.addWidget(self.export_button)
            actions.addWidget(self.folder_button)
            actions.addStretch()
            res.addLayout(actions)
            self.result_table = DataTable(preferences=self.layouts, key="results")
            preview_controls = QHBoxLayout()
            self.spatial_toggle = QCheckBox("Offline spatial preview")
            self.spatial_toggle.setChecked(True)
            preview_controls.addWidget(self.spatial_toggle)
            self.spatial = SpatialPreview(self.result_table)
            preview_controls.addWidget(button("Fit preview", self.spatial.fit))
            preview_controls.addStretch()
            res.addLayout(preview_controls)
            self.result_splitter = QSplitter(Qt.Orientation.Vertical)
            self.result_splitter.addWidget(self.result_table)
            self.result_splitter.addWidget(self.spatial)
            self.result_splitter.setSizes([400, 240])
            self.spatial_toggle.toggled.connect(self.spatial.setVisible)
            res.addWidget(self.result_splitter, 1)
            self.tabs.addTab(results, "Results")
            self.build_single()
            self.logs = QPlainTextEdit()
            self.logs.setReadOnly(True)
            self.logs.setMaximumBlockCount(3000)
            self.tabs.addTab(self.logs, "Activity")
            self.build_geo_tools()
            self.details = QPlainTextEdit()
            self.details.setReadOnly(True)
            self.details.setPlaceholderText(
                "Select a cell to preview its full record here."
            )
            self.details.setMinimumWidth(190)
            self.splitter.addWidget(self.details)
            self.splitter.setStretchFactor(1, 1)
            self.splitter.setChildrenCollapsible(False)
            self.details.setVisible(self.details_toggle.isChecked())
            outer.addWidget(self.splitter, 1)
            self.inline_error = label("", "error")
            self.inline_error.hide()
            outer.addWidget(self.inline_error)
            footer = QHBoxLayout()
            self.status = label(
                "Load or paste source and target files, then confirm their coordinate systems."
            )
            footer.addWidget(self.status, 1)
            self.elapsed = label("", "muted")
            footer.addWidget(self.elapsed)
            self.primary = button(
                "Validate & review",
                self.primary_action,
                True,
                QStyle.StandardPixmap.SP_DialogApplyButton,
            )
            footer.addWidget(self.primary)
            self.stop_button = button("Stop · Esc", self.stop)
            self.stop_button.setEnabled(False)
            footer.addWidget(self.stop_button)
            outer.addLayout(footer)
            self.progress = QProgressBar()
            self.progress.setRange(0, 100)
            self.progress.setValue(0)
            self.progress.setTextVisible(False)
            self.progress.setFixedHeight(5)
            outer.addWidget(self.progress)
            self.theme.currentTextChanged.connect(self.apply_theme)
            self.scale.currentTextChanged.connect(self.apply_theme)
            self.details_toggle.toggled.connect(self.details.setVisible)
            self.review_role.currentTextChanged.connect(self.show_audit)
            self.exclude_flags.toggled.connect(self.exclusion_changed)
            for widget in (self.method, self.mode, self.radius_unit, self.match_mode):
                widget.currentTextChanged.connect(self.settings_changed)
            self.top_n.valueChanged.connect(self.settings_changed)
            self.radius.textChanged.connect(self.settings_changed)
            for widget in (
                self.unlimited,
                self.ignore,
                self.flag,
                self.terrain,
                self.match_attributes,
            ):
                widget.toggled.connect(self.settings_changed)
            for table in (self.preview_table, self.review_table, self.result_table):
                table.view.selectionModel().currentChanged.connect(
                    lambda index, old, t=table: self.show_details(t, index)
                )
            self.tabs.currentChanged.connect(lambda: self.details.clear())
            self.tabs.currentChanged.connect(self.update_steps)
            self.shortcuts = []
            for key, callback in (
                ("Ctrl+O", lambda: self.browse("source")),
                ("Ctrl+Shift+O", lambda: self.browse("target")),
                (
                    "Ctrl+Return",
                    lambda: (
                        self.single()
                        if self.tabs.currentIndex() == 3
                        else self.primary_action()
                    ),
                ),
                ("Ctrl+E", self.export),
                ("Ctrl+Shift+S", self.save_project_dialog),
                ("Ctrl+Shift+P", self.open_project_dialog),
                ("Escape", self.stop),
                ("F1", self.about),
                ("Shift+F1", self.quick_guide),
            ):
                self.shortcut(key, callback)
            for i in range(6):
                self.shortcut(
                    f"Ctrl+{i + 1}",
                    lambda checked=False, index=i: self.select_tab(index),
                )

        def shortcut(self, key, callback):
            action = QAction(self)
            action.setShortcut(QKeySequence(key))
            action.triggered.connect(callback)
            self.addAction(action)
            self.shortcuts.append(action)

        def build_geo_tools(self):
            page = QWidget()
            layout = QVBoxLayout(page)
            layout.addWidget(
                label("Geo Tools — spatial analysis and engineering checks", "heading")
            )
            layout.addWidget(
                label(
                    "Use the mapped working copies from Load. Geo Tools validates them again and keeps its own results. DEM profiles use the last calculated point links. Settings below are separate from point matching.",
                    "muted",
                )
            )
            split = QSplitter(Qt.Orientation.Horizontal)
            self.geo_controls = QScrollArea()
            self.geo_controls.setWidgetResizable(True)
            self.geo_controls.setMinimumWidth(310)
            controls = QWidget()
            self.geo_controls.setWidget(controls)
            form = QFormLayout(controls)
            form.setRowWrapPolicy(QFormLayout.RowWrapPolicy.WrapAllRows)
            self.geo_operation = combo(GEO_TOOLS)
            form.addRow("Operation", self.geo_operation)
            self.geo_description = label("", "muted")
            form.addRow(self.geo_description)
            self.geo_fields = {}
            self.geo_field_rows = {}

            def add(key, title, widget):
                self.geo_fields[key] = widget
                caption = label(title)
                form.addRow(caption, widget)
                self.geo_field_rows[key] = (caption, widget)
                return widget

            add(
                "crs",
                "Projected analysis CRS in metres (choose for your area)",
                QLineEdit(""),
            )
            self.geo_fields["crs"].setPlaceholderText(
                "e.g. EPSG:32636 — only for UTM zone 36N"
            )
            for key, title, filter_text in (
                (
                    "vector_path",
                    "Line / polygon file",
                    "Vector layers (*.gpkg *.geojson *.json)",
                ),
                (
                    "dem_path",
                    "Local elevation raster",
                    "Elevation raster (*.tif *.tiff)",
                ),
            ):
                row = QWidget()
                actions = QHBoxLayout(row)
                actions.setContentsMargins(0, 0, 0, 0)
                field = QLineEdit()
                actions.addWidget(field, 1)
                actions.addWidget(
                    button(
                        "Browse",
                        lambda checked=False, k=key, f=filter_text: self.geo_browse(
                            k, f
                        ),
                    )
                )
                add(key, title, row)
                self.geo_fields[key] = field
            add("layer", "GeoPackage layer (required if more than one)", QLineEdit())
            for key, title, default in (
                ("horizontal_tolerance", "Horizontal tolerance (m)", "0.1"),
                ("vertical_tolerance", "Vertical tolerance (m)", "0.1"),
                ("radius", "Buffer radius / sector range (m)", "1000"),
                ("azimuth", "Sector azimuth from true north (degrees)", "0"),
                ("beamwidth", "Sector beamwidth (degrees)", "65"),
                ("spacing", "DEM sample spacing (m)", "30"),
                ("source_agl", "Source antenna above sampled DEM (m)", "30"),
                ("target_agl", "Target antenna above sampled DEM (m)", "30"),
                ("frequency", "Frequency (GHz)", "5"),
                ("fresnel", "Required first Fresnel fraction (0–1)", "0.6"),
                ("k_factor", "Effective Earth-radius factor K", "1.3333333333333333"),
            ):
                add(key, title, QLineEdit(default))
            add("check_vertical", "", QCheckBox("Check vertical survey tolerance too"))
            add(
                "dem_units",
                "DEM elevation units after band scale/offset",
                combo(["metres", "international feet"]),
            )
            band = QSpinBox()
            band.setRange(1, 999)
            add("band", "DEM elevation band", band)
            add(
                "dem_confirmed",
                "",
                QCheckBox("I confirm DEM units and the vertical reference"),
            )
            self.geo_run_button = button("Run analysis", self.run_geo_analysis, True)
            form.addRow(self.geo_run_button)
            split.addWidget(self.geo_controls)
            right = QWidget()
            results = QVBoxLayout(right)
            results.setContentsMargins(0, 0, 0, 0)
            self.geo_summary = label("Choose an operation and run analysis.", "heading")
            results.addWidget(self.geo_summary)
            self.geo_stale = label("", "error")
            self.geo_stale.hide()
            results.addWidget(self.geo_stale)
            actions = QHBoxLayout()
            self.geo_table_choice = combo(["Results"])
            self.geo_export_button = button("Export analysis", self.export_geo_dialog)
            self.geo_export_button.setEnabled(False)
            actions.addWidget(self.geo_table_choice, 1)
            actions.addWidget(self.geo_export_button)
            results.addLayout(actions)
            self.geo_table = DataTable(preferences=self.layouts, key="geo_tools")
            self.geo_table.view.selectionModel().currentChanged.connect(
                lambda index, old: self.show_details(self.geo_table, index)
            )
            results.addWidget(self.geo_table, 1)
            self.geo_profile_choice = combo([])
            self.geo_profile_choice.setAccessibleName("Terrain profile link")
            self.geo_profile_choice.currentIndexChanged.connect(self.show_geo_profile)
            results.addWidget(self.geo_profile_choice)
            self.geo_profile_plot = TerrainPlot()
            results.addWidget(self.geo_profile_plot)
            self.geo_profile_choice.hide()
            self.geo_profile_plot.hide()
            split.addWidget(right)
            split.setStretchFactor(1, 1)
            split.setSizes([330, 750])
            layout.addWidget(split, 1)
            self.geo_analysis = None
            self.geo_tab_index = self.tabs.addTab(page, "Geo Tools")
            self.geo_operation.currentTextChanged.connect(self.geo_operation_changed)
            self.geo_table_choice.currentTextChanged.connect(self.show_geo_table)
            for key, widget in self.geo_fields.items():
                if isinstance(widget, QLineEdit):
                    widget.textChanged.connect(self.geo_settings_changed)
                elif isinstance(widget, QComboBox):
                    widget.currentTextChanged.connect(self.geo_settings_changed)
                elif isinstance(widget, QCheckBox):
                    widget.toggled.connect(self.geo_settings_changed)
                elif isinstance(widget, QSpinBox):
                    widget.valueChanged.connect(self.geo_settings_changed)
            self.geo_operation_changed()

        def geo_operation_changed(self, *args):
            index = self.geo_operation.currentIndex()
            common = {"crs"} if index != 4 else set()
            specific = [
                {"vector_path", "layer"},
                {"vector_path", "layer"},
                {"horizontal_tolerance", "vertical_tolerance", "check_vertical"},
                {"radius"},
                {
                    "dem_path",
                    "spacing",
                    "source_agl",
                    "target_agl",
                    "frequency",
                    "fresnel",
                    "k_factor",
                    "dem_units",
                    "band",
                    "dem_confirmed",
                },
                {"radius", "azimuth", "beamwidth"},
            ][index]
            for key, widgets in self.geo_field_rows.items():
                for widget in widgets:
                    widget.setVisible(key in common | specific)
            descriptions = [
                "Source points → nearest line geometry. Requires a local projected CRS in metres. Equidistant features keep file order. Source input only; target point file is unused.",
                "Source points → covering polygons or nearest polygon. Includes boundary distance and overlap count. Holes count as outside. Source input only.",
                "Source = baseline; target = observed survey. Match exact trimmed IDs. Deltas are observed minus baseline, in the chosen projected CRS. Heights need the same vertical reference.",
                "Creates a projected buffer per source and counts eligible targets by exact projected distance. Boundary included. The point-matching identical-location and attribute filters apply.",
                "Offline profiles of ALL matched rows from the last point run. Uses chosen constant antenna heights above the DEM for every endpoint, not input height fields. Nearest-cell samples; no-data is Unknown. Up to 10,000 samples/link and 200,000 total.",
                "One sector per source using a shared true-north azimuth, beamwidth and WGS84 range. Counts targets geometrically; no propagation, terrain loss or coverage prediction. Identical-location and attribute filters apply.",
            ]
            self.geo_description.setText(descriptions[index])
            self.geo_settings_changed()

        def geo_settings_changed(self, *args):
            if self.sender() in (
                self.geo_fields.get("dem_path"),
                self.geo_fields.get("dem_units"),
                self.geo_fields.get("band"),
            ):
                self.geo_fields["dem_confirmed"].setChecked(False)
            if self.project_active:
                self.project_dirty = True
            self.mark_geo_stale()

        def mark_geo_stale(self, *args):
            if getattr(self, "geo_analysis", None) is not None:
                self.geo_stale.setText(
                    "Previous analysis — exports retain its original inputs and settings. Run analysis to update."
                )
                self.geo_stale.show()

        def geo_browse(self, key, filter_text):
            path, _ = QFileDialog.getOpenFileName(
                self, "Open analysis input", "", filter_text
            )
            if path:
                self.geo_fields[key].setText(path)
                if key == "dem_path":
                    self.geo_fields["dem_confirmed"].setChecked(False)

        def geo_options(self):
            options = {}
            for key, widget in self.geo_fields.items():
                options[key] = (
                    widget.text().strip()
                    if isinstance(widget, QLineEdit)
                    else (
                        widget.currentText()
                        if isinstance(widget, QComboBox)
                        else (
                            widget.isChecked()
                            if isinstance(widget, QCheckBox)
                            else widget.value()
                        )
                    )
                )
            return options

        def run_geo_analysis(self):
            if self.process is not None:
                return
            tool = self.geo_operation.currentText()
            try:
                if tool == GEO_TOOLS[4]:
                    if self.result is None or self.result_spec is None:
                        raise ValueError(
                            "Calculate point matches first. The DEM tool profiles the frozen matched results."
                        )
                    spec = self.result_spec
                else:
                    if not self.files["source"]:
                        raise ValueError("Load and map source points first.")
                    if tool in GEO_TOOLS[2:4] or tool == GEO_TOOLS[5]:
                        if not self.files["target"]:
                            raise ValueError(
                                "This operation needs mapped source and target points."
                            )
                    spec = RunSpec(
                        self.files["source"],
                        self.files["target"] or self.files["source"],
                        ignore_same_location=self.ignore.isChecked(),
                        flag_outliers=self.flag.isChecked(),
                        match_attributes=self.match_attributes.isChecked(),
                    )
                options = self.geo_options()
                if tool == GEO_TOOLS[4] and not options["dem_confirmed"]:
                    raise ValueError(
                        "Confirm DEM units and vertical reference before profiling."
                    )
                if tool != GEO_TOOLS[4] and not options["crs"]:
                    raise ValueError(
                        "Enter a projected analysis CRS with metre units, suitable for your study area."
                    )
            except ValueError as error:
                self.error(error)
                return
            payload = {
                "tool": tool,
                "options": options,
                "spec": spec,
                "exclude_flags": self.exclude_flags.isChecked(),
            }
            if tool == GEO_TOOLS[4]:
                payload.update(
                    links=self.result,
                    matched_metadata={
                        "Run specification": asdict(self.result_spec),
                        "Calculated at": self.result_created,
                    },
                )

            def complete(analysis):
                self.geo_analysis = analysis
                self.geo_profile_choice.clear()
                profile = analysis["tables"].get("Profile samples")
                self.geo_profile_choice.setVisible(profile is not None)
                self.geo_profile_plot.setVisible(profile is not None)
                if profile is not None:
                    for row in analysis["tables"]["Results"].to_dict("records"):
                        self.geo_profile_choice.addItem(
                            f"Link {row['Link']}: {row['Source site']} → {row['Target site']} · {row['Status']}",
                            int(row["Link"]),
                        )
                    self.show_geo_profile()
                self.geo_table_choice.clear()
                self.geo_table_choice.addItems(list(analysis["tables"]))
                self.geo_table_choice.setCurrentText("Results")
                self.show_geo_table()
                self.geo_summary.setText(
                    f"{tool}\n{len(analysis['tables']['Results']):,} result rows · {analysis['metadata'].get('Analysis CRS','WGS84 geodesic profiles')}"
                )
                self.geo_stale.hide()
                self.geo_export_button.setEnabled(True)
                self.update_steps()
                self.status.setText(
                    "Analysis complete. Inspect results and audit tables; export this analysis separately."
                )

            self.launch("geo_tool", payload, complete)

        def show_geo_profile(self, *args):
            analysis = getattr(self, "geo_analysis", None)
            if analysis is None:
                return
            profile = analysis["tables"].get("Profile samples")
            if profile is not None:
                self.geo_profile_plot.set_frame(
                    profile.loc[
                        profile["Link"].eq(self.geo_profile_choice.currentData())
                    ]
                )

        def show_geo_table(self, *args):
            if self.geo_analysis is not None:
                self.geo_table.clear_filters()
                self.geo_table.set_frame(
                    self.geo_analysis["tables"].get(
                        self.geo_table_choice.currentText(), pd.DataFrame()
                    )
                )

        def export_geo_dialog(self):
            if self.process is not None or self.geo_analysis is None:
                return
            path, selected = QFileDialog.getSaveFileName(
                self,
                "Export Geo Tools analysis",
                "GeoMatch_analysis",
                "Excel workbook (*.xlsx);;GeoPackage (*.gpkg)",
            )
            if path:
                if not Path(path).suffix:
                    path += ".gpkg" if "GeoPackage" in selected else ".xlsx"
                    if (
                        Path(path).exists()
                        and QMessageBox.question(
                            self, "Replace report?", f"Replace {path}?"
                        )
                        != QMessageBox.StandardButton.Yes
                    ):
                        return
                self.export_geo_to(path)

        def export_geo_to(self, path):
            if self.process is not None or self.geo_analysis is None:
                return
            try:
                self.protect_inputs(path)
                frozen = self.geo_analysis["metadata"]
                inputs = [
                    frozen.get("Vector input", {}).get("path"),
                    frozen.get("DEM", {}).get("path"),
                ]
                inputs += [
                    frozen["Input specification"][r]["path"]
                    for r in ("source", "target")
                ]
                if any(
                    value and Path(path).resolve() == Path(value).resolve()
                    for value in inputs
                ):
                    raise ValueError(
                        "Choose a separate output; analysis inputs cannot be overwritten."
                    )
                if Path(path).suffix.lower() not in (".xlsx", ".gpkg"):
                    raise ValueError("Choose an .xlsx or .gpkg analysis output.")
            except ValueError as error:
                self.error(error)
                return
            self.export_temporary = str(
                Path(path).with_name(
                    f".geomatch-analysis-{os.getpid()}-{time.time_ns()}{Path(path).suffix}"
                )
            )

            def complete(saved):
                self.last_export = saved
                self.folder_button.setEnabled(True)
                self.status.setText(
                    f"Saved analysis: {saved}. Full tables exported, regardless of table filters."
                )

            self.launch(
                "export_geo_tool",
                (path, self.geo_analysis, self.export_temporary),
                complete,
            )

        def build_single(self):
            self.single_panel = QWidget()
            layout = QVBoxLayout(self.single_panel)
            layout.addWidget(label("Measure a single link", "heading"))
            layout.addWidget(
                label(
                    "WGS84 geodesic ground distance and forward bearing. Paste latitude first, longitude second. Decimal degrees or DMS; separate DMS pairs with a comma or semicolon."
                )
            )
            self.coordinates = []
            for role in ("Source", "Target"):
                group = QGroupBox(role)
                form = QFormLayout(group)
                for name in ("Latitude", "Longitude"):
                    entry = QLineEdit()
                    entry.setPlaceholderText("Decimal degrees or DMS")
                    entry.setAccessibleName(f"{role} {name}")
                    entry.returnPressed.connect(self.single)
                    entry.textChanged.connect(self.invalidate_single)
                    self.coordinates.append(entry)
                    form.addRow(name, entry)
                form.addRow(
                    button(
                        "Paste coordinate pair",
                        lambda checked=False, offset=(
                            0 if role == "Source" else 2
                        ): self.paste_pair(offset),
                    )
                )
                layout.addWidget(group)
            actions = QHBoxLayout()
            actions.addWidget(button("Swap endpoints", self.swap_endpoints))
            actions.addWidget(
                button("Clear", lambda: [c.clear() for c in self.coordinates])
            )
            actions.addWidget(button("Calculate", self.single, True))
            layout.addLayout(actions)
            self.single_result = label("Enter both coordinate pairs.", "heading")
            layout.addWidget(self.single_result)
            self.copy_single = button(
                "Copy result",
                lambda: application.clipboard().setText(self.single_value or ""),
            )
            self.copy_single.setEnabled(False)
            layout.addWidget(self.copy_single)
            layout.addStretch()
            scroll = QScrollArea()
            scroll.setWidgetResizable(True)
            scroll.setWidget(self.single_panel)
            self.tabs.addTab(scroll, "Single link")

        def update_steps(self, *args):
            active = self.tabs.currentIndex()
            if active == self.geo_tab_index:
                self.primary.setText("Run analysis")
                self.steps[3].setEnabled(
                    self.process is None and self.geo_analysis is not None
                )
            else:
                self.primary.setText(
                    "Validate & review"
                    if self.frames is None
                    else "Match reviewed data"
                )
                self.steps[3].setEnabled(
                    self.process is None and self.result is not None
                )
            for index, control in enumerate(self.steps):
                control.setProperty("activeStep", index == active and active < 3)
                control.style().unpolish(control)
                control.style().polish(control)

        def select_tab(self, index):
            self.tabs.setCurrentIndex(index)
            if index == 0:
                self.side.setFocus()

        def show_details(self, table, index):
            if index.isValid():
                row = table.proxy.mapToSource(index).row()
                record = table.model.frame.iloc[row]
                self.details.setPlainText(
                    "RECORD DETAILS\n\n"
                    + "\n\n".join(
                        f"{key}\n{'' if pd.isna(value) else value}"
                        for key, value in record.items()
                    )
                )

        def apply_theme(self, *args):
            dark = self.theme.currentText() == "Dark" or (
                self.theme.currentText() == "System"
                and application.styleHints().colorScheme() == Qt.ColorScheme.Dark
            )
            bg, card, text, muted, border, alternate = (
                ("#111c2a", "#1c2b3d", "#edf4fa", "#a8bccd", "#3c5064", "#223347")
                if dark
                else ("#f3f6fa", "#ffffff", "#182d43", "#52677d", "#ccd7e3", "#f5f9fc")
            )
            palette = QPalette()
            for role, color in (
                (QPalette.ColorRole.Window, bg),
                (QPalette.ColorRole.WindowText, text),
                (QPalette.ColorRole.Base, card),
                (QPalette.ColorRole.AlternateBase, alternate),
                (QPalette.ColorRole.Text, text),
                (QPalette.ColorRole.Button, card),
                (QPalette.ColorRole.ButtonText, text),
                (QPalette.ColorRole.Highlight, "#0b737c"),
                (QPalette.ColorRole.HighlightedText, "#ffffff"),
                (QPalette.ColorRole.PlaceholderText, muted),
            ):
                palette.setColor(role, QColor(color))
            application.setPalette(palette)
            point_size = {"100%": 10, "125%": 12, "150%": 14}[self.scale.currentText()]
            application.setFont(QFont("Segoe UI", point_size))
            self.setStyleSheet(f"""
                QMainWindow, QWidget {{ color: {text}; }}
                QMainWindow {{ background: {bg}; }}
                QLabel#brand {{ font-size: 25px; font-weight: 700; }}
                QLabel#heading {{ font-weight: 600; font-size: {point_size + 3}pt; }}
                QLabel#muted {{ color: {muted}; }}
                QLabel#error {{ color: {'#ffb4a9' if dark else '#ad2938'}; padding: 5px; }}
                QGroupBox {{ border: 1px solid {border}; border-radius: 8px; margin-top: 13px; padding: 12px 8px 8px; font-weight: 600; }}
                QGroupBox::title {{ subcontrol-origin: margin; left: 10px; padding: 0 4px; }}
                QPushButton {{ background: {card}; border: 1px solid {border}; border-radius: 6px; padding: 7px 11px; }}
                QPushButton:hover {{ border-color: #128d94; }}
                QPushButton[activeStep="true"] {{ border: 2px solid #128d94; font-weight: 600; }}
                QPushButton:focus, QLineEdit:focus, QComboBox:focus, QSpinBox:focus {{ border: 2px solid #128d94; }}
                QPushButton#primary {{ background: #086e76; color: white; border-color: #086e76; font-weight: 600; }}
                QPushButton:disabled {{ color: {muted}; background: {bg}; }}
                QLineEdit, QComboBox, QSpinBox {{ background: {card}; border: 1px solid {border}; border-radius: 5px; padding: 6px; }}
                QTabWidget::pane {{ border: 1px solid {border}; border-radius: 6px; }}
                QTabBar::tab {{ padding: 9px 14px; border: 0; background: transparent; }}
                QTabBar::tab:selected {{ color: {'#6ce0de' if dark else '#086e76'}; border-bottom: 3px solid #128d94; font-weight: 600; }}
                QTableView, QPlainTextEdit {{ background: {card}; alternate-background-color: {alternate}; border: 1px solid {border}; border-radius: 4px; gridline-color: {border}; selection-background-color: #0b737c; selection-color: white; }}
                QHeaderView::section {{ background: {bg}; color: {text}; padding: 8px; border: 0; border-bottom: 1px solid {border}; font-weight: 600; }}
                QSplitter::handle {{ background: {border}; }}
                QProgressBar {{ border: 0; background: {border}; }} QProgressBar::chunk {{ background: #128d94; }}
            """)
            for table in (self.preview_table, self.review_table, self.result_table):
                table.fit_headers()
            for dialog in getattr(self, "help_dialogs", {}).values():
                if hasattr(dialog, "content"):
                    self.style_help_browser(dialog.content)
            self.settings["theme"] = self.theme.currentText()
            self.settings["scale"] = self.scale.currentText()

        def restore_layout(self):
            area = application.primaryScreen().availableGeometry()
            self.resize(
                min(1280, int(area.width() * 0.92)), min(840, int(area.height() * 0.90))
            )
            self.setMinimumSize(min(900, area.width()), min(600, area.height()))
            saved = self.settings.get("qt_geometry")
            if saved:
                self.restoreGeometry(QByteArray.fromBase64(saved.encode()))
                # Fit saved geometry to the available screen after monitor changes.
                self.resize(
                    min(self.width(), area.width()), min(self.height(), area.height())
                )
                self.move(
                    max(area.left(), min(self.x(), area.right() - self.width())),
                    max(area.top(), min(self.y(), area.bottom() - self.height())),
                )
            self.splitter.setSizes(self.settings.get("qt_splitter", [330, 750, 230]))

        def persist_settings(self):
            try:
                temp = self.settings_path.with_suffix(".tmp")
                temp.write_text(json.dumps(self.settings, indent=2), encoding="utf-8")
                os.replace(temp, self.settings_path)
            except OSError as error:
                self.log(f"Could not save preferences: {error}")

        def log(self, message):
            self.logs.appendPlainText(f"{datetime.now():%H:%M:%S}  {message}")

        def error(self, message):
            self.inline_error.setText(str(message))
            self.inline_error.show()
            self.status.setText(
                "Correct the issue and try again. Details are in Activity."
            )
            self.log(str(message))

        def clear_error(self):
            self.inline_error.clear()
            self.inline_error.hide()

        def edit_working(self, role):
            if self.process is not None:
                return
            file = self.files[role]
            if file is None:
                self.error(f"Map the {role} file first.")
                return

            def opened(raw):
                self.working_dialog = WorkingDialog(self, role, raw, file)
                if self.working_dialog.exec() == QDialog.DialogCode.Accepted:
                    revised = self.working_dialog.working.file(file)
                    if revised != file:
                        self.files[role] = revised
                        self.project_dirty = True
                        self.project_active = True
                        self.invalidate_review()
                        self.update_file_labels()
                        self.status.setText(
                            "Working copy updated. Validate again; save a project to keep these changes."
                        )

            self.launch("working", file, opened)

        def save_project_dialog(self):
            if self.process is not None:
                return
            path, _ = QFileDialog.getSaveFileName(
                self,
                "Save GeoMatch project",
                "GeoMatch.geomatch",
                "GeoMatch project (*.geomatch)",
            )
            if path:
                if not Path(path).suffix:
                    path += ".geomatch"
                    if (
                        Path(path).exists()
                        and QMessageBox.question(
                            self, "Replace project?", f"Replace {path}?"
                        )
                        != QMessageBox.StandardButton.Yes
                    ):
                        return
                self.save_project_to(path)

        def save_project_to(self, path):
            if self.process is not None:
                return
            try:
                spec = self.current_spec()
                self.protect_inputs(path)
            except ValueError as error:
                self.error(error)
                return

            def complete(saved):
                self.project_dirty = False
                self.project_active = True
                self.status.setText(
                    f"Saved project: {saved}. Reopen it to restore inputs, settings and working edits."
                )

            self.launch(
                "save_project",
                (
                    path,
                    spec,
                    self.exclude_flags.isChecked(),
                    {
                        "tool": self.geo_operation.currentText(),
                        "options": self.geo_options(),
                    },
                ),
                complete,
            )

        def open_project_dialog(self):
            if self.process is not None:
                return
            path, _ = QFileDialog.getOpenFileName(
                self, "Open GeoMatch project", "", "GeoMatch project (*.geomatch)"
            )
            if path:
                self.open_project_from(path)

        def open_project_from(self, path):
            if self.process is not None:
                return
            if (
                self.project_dirty
                and QMessageBox.question(
                    self,
                    "Replace unsaved working changes?",
                    "Open this project and discard unsaved working-copy changes? Save a project first if you want to keep them.",
                )
                != QMessageBox.StandardButton.Yes
            ):
                return

            def complete(value):
                spec, exclude_flags, geo_saved = value
                if isinstance(geo_saved, dict):
                    self.geo_operation.setCurrentText(
                        geo_saved.get("tool", GEO_TOOLS[0])
                    )
                    for key, saved in geo_saved.get("options", {}).items():
                        widget = self.geo_fields.get(key)
                        if isinstance(widget, QLineEdit):
                            widget.setText(str(saved))
                        elif isinstance(widget, QComboBox):
                            widget.setCurrentText(str(saved))
                        elif isinstance(widget, QCheckBox):
                            widget.setChecked(
                                bool(saved) if key != "dem_confirmed" else False
                            )
                        elif isinstance(widget, QSpinBox):
                            widget.setValue(int(saved))
                self.files = {"source": spec.source, "target": spec.target}
                self.method.setCurrentText(spec.method)
                self.mode.setCurrentText(spec.mode)
                self.match_mode.setCurrentText(spec.match_mode)
                self.top_n.setValue(spec.top_n)
                self.radius_unit.setCurrentText("m")
                self.radius.setText(str(spec.radius_m or 5000))
                self.unlimited.setChecked(spec.radius_m == 0)
                for widget, value in (
                    (self.ignore, spec.ignore_same_location),
                    (self.flag, spec.flag_outliers),
                    (self.terrain, spec.terrain),
                    (self.match_attributes, spec.match_attributes),
                ):
                    widget.setChecked(value)
                self.exclude_flags.setChecked(exclude_flags)
                self.project_exclude_flags = exclude_flags
                self.update_file_labels()
                self.invalidate_review()
                self.preview_table.set_frame(pd.DataFrame())
                self.project_dirty = False
                self.project_active = True
                self.tabs.setCurrentIndex(0)
                self.load_summary.setText(
                    f"Project restored: {Path(path).name}. Input fingerprints checked; validate before matching."
                )
                self.status.setText(
                    "Project opened. Working-copy edits and row exclusions restored."
                )

            self.launch("load_project", (path, self.clipboard_store.name), complete)

        def protect_inputs(self, path):
            resolved = Path(path).resolve()
            if any(
                file and resolved == Path(file.path).resolve()
                for file in list(self.files.values())
                + (
                    [self.result_spec.source, self.result_spec.target]
                    if self.result_spec
                    else []
                )
            ):
                raise ValueError(
                    "Choose a separate output path; an input file cannot be overwritten."
                )

        def large_run_dialog(self):
            if self.process is not None:
                return
            try:
                spec = self.current_spec()
                if (
                    spec.flag_outliers
                    or spec.terrain
                    or spec.match_mode == "Mutual nearest"
                    or Path(spec.source.path).suffix.lower() not in (".csv", ".txt")
                ):
                    raise ValueError(
                        "Large-file mode needs a CSV/TXT source, with outlier and terrain checks off and a mode other than Mutual nearest. Target stays in memory. Working edits and explicit exclusions are applied."
                    )
            except ValueError as error:
                self.error(error)
                return
            path, _ = QFileDialog.getSaveFileName(
                self,
                "Stream validated results to ZIP (no results-table preview)",
                "GeoMatch_large_report.zip",
                "ZIP report (*.zip)",
            )
            if path:
                if not Path(path).suffix:
                    path += ".zip"
                    if (
                        Path(path).exists()
                        and QMessageBox.question(
                            self, "Replace report?", f"Replace {path}?"
                        )
                        != QMessageBox.StandardButton.Yes
                    ):
                        return
                self.large_run_to(path)

        def large_run_to(self, path):
            if self.process is not None:
                return
            try:
                spec = self.current_spec()
                self.protect_inputs(path)
            except ValueError as error:
                self.error(error)
                return

            def complete(info):
                self.last_export = info["path"]
                self.folder_button.setEnabled(True)
                self.status.setText(
                    f"Saved ZIP: {info['source_rows']:,} source rows, {info['eligible']:,} eligible, {info['results']:,} results. Open its review tables to inspect rejected/excluded rows."
                )
                self.log(self.status.text())

            self.export_temporary = str(
                Path(path).with_name(
                    f".geomatch-stream-{os.getpid()}-{time.time_ns()}.zip"
                )
            )
            self.launch("stream_large", (path, spec, self.export_temporary), complete)

        def refresh_summary(self):
            self.top_n.setEnabled(self.match_mode.currentText() == "Nearest N")
            self.radius.setEnabled(not self.unlimited.isChecked())
            self.radius_unit.setEnabled(not self.unlimited.isChecked())
            try:
                radius_in_metres(
                    self.radius.text(),
                    self.radius_unit.currentText(),
                    self.unlimited.isChecked(),
                )
                reach = (
                    "no distance limit"
                    if self.unlimited.isChecked()
                    else f"within {self.radius.text()} {self.radius_unit.currentText()}"
                )
                self.settings_summary.setText(
                    f"{self.match_mode.currentText()} · {reach}\n{self.method.currentText()} · {self.mode.currentText()}"
                    + (
                        "\nExact attribute match required"
                        if self.match_attributes.isChecked()
                        else ""
                    )
                )
            except ValueError as error:
                self.settings_summary.setText(str(error))

        def settings_changed(self, *args):
            if self.project_active:
                self.project_dirty = True
            self.invalidate_review()
            self.refresh_summary()

        def invalidate_review(self):
            self.mark_geo_stale()
            self.frames = self.audits = self.run_spec = None
            self.primary.setText("Validate & review")
            self.review_summary.setText(
                "Files or settings changed. Validate again before matching."
            )
            self.review_table.set_frame(pd.DataFrame())
            if self.result is not None:
                self.stale.setText(
                    "Previous run · These results keep their original settings. Validate and match again to update them."
                )
                self.stale.show()

        def exclusion_changed(self, *args):
            self.mark_geo_stale()
            if self.project_active:
                self.project_dirty = True
            self.show_audit()
            if self.frames is not None:
                self.primary.setText("Match reviewed data")
            if self.result is not None:
                self.stale.setText(
                    "Review choice changed · Existing results and exports still represent the previous run."
                )
                self.stale.show()

        def current_spec(self):
            if not all(self.files.values()):
                raise ValueError("Load and map both source and target files first.")
            spec = RunSpec(
                self.files["source"],
                self.files["target"],
                self.method.currentText(),
                self.mode.currentText(),
                self.top_n.value(),
                radius_in_metres(
                    self.radius.text(),
                    self.radius_unit.currentText(),
                    self.unlimited.isChecked(),
                ),
                self.ignore.isChecked(),
                self.flag.isChecked(),
                self.terrain.isChecked(),
                self.match_mode.currentText(),
                self.match_attributes.isChecked(),
            )
            spec.validate()
            return spec

        def set_busy(self, busy):
            self.geo_controls.setEnabled(not busy)
            self.geo_export_button.setEnabled(
                not busy and self.geo_analysis is not None
            )
            self.side.setEnabled(not busy)
            self.single_panel.setEnabled(not busy)
            self.exclude_flags.setEnabled(not busy)
            self.primary.setEnabled(not busy)
            self.steps[2].setEnabled(not busy)
            self.steps[3].setEnabled(not busy and self.result is not None)
            self.export_button.setEnabled(not busy and self.result is not None)
            self.stop_button.setEnabled(busy)
            self.update_steps()

        def launch(self, kind, payload, callback):
            if self.process is not None or self.closing:
                return
            self.clear_error()
            self.kind, self.callback = kind, callback
            self.pending = None
            self.stopping = False
            self.started = time.monotonic()
            self.cancel = self.context.Event()
            self.events = self.context.Queue()
            self.process = self.context.Process(
                target=job_worker, args=(kind, payload, self.events, self.cancel)
            )
            self.set_busy(True)
            self.progress.setRange(0, 0)
            self.status.setText(f"{kind.capitalize()} in progress…")
            try:
                self.process.start()
                self.log(f"Started {kind}.")
            except Exception as error:
                self.process = None
                self.events.close()
                self.events = None
                self.set_busy(False)
                self.progress.setRange(0, 100)
                self.progress.setValue(0)
                self.error(error)

        def poll(self):
            if self.process is None:
                return
            self.elapsed.setText(f"{time.monotonic() - self.started:.1f} s")
            try:
                for _ in range(0 if self.stopping else 100):
                    tag, payload = self.events.get_nowait()
                    if tag == "progress":
                        message, fraction = payload
                        self.status.setText(message)
                        if fraction is not None:
                            self.progress.setRange(0, 100)
                            self.progress.setValue(int(fraction * 100))
                    elif tag in ("done", "error", "cancelled"):
                        self.pending = (tag, payload)
            except queue.Empty:
                pass
            if (
                self.stopping
                and self.process.is_alive()
                and time.monotonic() - self.stop_time > 1.5
            ):
                self.process.terminate()
            if self.process.is_alive():
                return
            try:
                while not self.stopping:
                    tag, payload = self.events.get_nowait()
                    if tag in ("done", "error", "cancelled"):
                        self.pending = (tag, payload)
            except queue.Empty:
                pass
            exitcode = self.process.exitcode
            self.process.join(timeout=0)
            self.process.close()
            self.process = None
            self.events.close()
            self.events = None
            if self.export_temporary:
                try:
                    Path(self.export_temporary).unlink(missing_ok=True)
                except OSError as error:
                    self.log(f"Temporary export cleanup failed: {error}")
                for suffix in (
                    "-wal",
                    "-shm",
                    "-journal",
                    ".results.csv",
                    ".audit.csv",
                ):
                    Path(self.export_temporary + suffix).unlink(missing_ok=True)
                self.export_temporary = None
            self.progress.setRange(0, 100)
            self.set_busy(False)
            if self.closing:
                self.close()
                return
            if self.stopping or (self.pending and self.pending[0] == "cancelled"):
                self.progress.setValue(0)
                self.status.setText(
                    "Stopped. Worker has exited; you can start another operation."
                )
                self.log(f"Stopped {self.kind}.")
            elif self.pending and self.pending[0] == "done":
                self.progress.setValue(100)
                self.status.setText("Ready.")
                try:
                    self.callback(self.pending[1])
                except Exception as error:
                    self.log(traceback.format_exc())
                    self.error(error)
            else:
                message, detail = (
                    self.pending[1]
                    if self.pending and self.pending[0] == "error"
                    else (f"Worker exited unexpectedly (code {exitcode}).", "")
                )
                self.log(detail)
                self.progress.setValue(0)
                self.error(message)

        def stop(self):
            if self.process is not None and not self.stopping:
                self.stopping = True
                self.stop_time = time.monotonic()
                self.cancel.set()
                self.stop_button.setEnabled(False)
                self.status.setText("Stopping… waiting for the worker to exit.")

        def closeEvent(self, event):
            if self.process is not None:
                self.closing = True
                self.stop()
                self.status.setText("Stopping active work before closing…")
                event.ignore()
                return
            if (
                self.project_dirty
                and QMessageBox.question(
                    self,
                    "Unsaved project changes",
                    "Close without saving these project or working-copy changes? Choose No, then Save project to keep them.",
                )
                != QMessageBox.StandardButton.Yes
            ):
                self.closing = False
                event.ignore()
                return
            for table in (self.preview_table, self.review_table, self.result_table):
                table.save_layout()
            self.settings["qt_geometry"] = bytes(
                self.saveGeometry().toBase64()
            ).decode()
            self.settings["qt_splitter"] = self.splitter.sizes()
            self.settings["qt_details"] = self.details_toggle.isChecked()
            self.persist_settings()
            self.timer.stop()
            self.clipboard_store.cleanup()
            event.accept()

        def dragEnterEvent(self, event):
            if (
                self.process is None
                and event.mimeData().hasUrls()
                and all(url.isLocalFile() for url in event.mimeData().urls())
            ):
                event.acceptProposedAction()

        def dropEvent(self, event):
            paths = [
                url.toLocalFile()
                for url in event.mimeData().urls()
                if url.isLocalFile()
            ]
            if self.process is not None or not paths:
                return
            if len(paths) != 1:
                self.error(
                    "Drop one file at a time, then choose whether it is source or target."
                )
                return
            role, ok = QInputDialog.getItem(
                self,
                "Open dropped file",
                "Use this file as",
                ["source", "target"],
                0,
                False,
            )
            if ok:
                self.open_input(role, paths[0])
                event.acceptProposedAction()

        def paste_table(self, role):
            if self.process is not None:
                return
            try:
                frame = parse_pasted_table(application.clipboard().text())
                path = (
                    Path(self.clipboard_store.name)
                    / f"Pasted_{role}_{time.time_ns()}.csv"
                )
                frame.to_csv(path, index=False)
                self.open_input(role, str(path))
            except (ValueError, OSError) as error:
                self.error(error)

        def browse(self, role):
            if self.process is not None:
                return
            path, _ = QFileDialog.getOpenFileName(
                self,
                f"Open {role} sites",
                "",
                "Site files (*.csv *.txt *.xlsx *.xls *.geojson);;All files (*)",
            )
            if path:
                self.open_input(role, path)

        def open_input(self, role, path, preferred_sheet=None):
            if self.process is not None:
                return
            if not Path(path).is_file():
                self.error(f"File not found: {path}. Browse to its new location.")
                return
            if Path(path).suffix.lower() in (".xlsx", ".xls"):
                self.launch(
                    "sheets",
                    path,
                    lambda info: self.choose_sheet(role, info, preferred_sheet),
                )
            else:
                self.launch(
                    "inspect", path, lambda info: self.mapping_dialog(role, path, info)
                )

        def choose_sheet(self, role, info, preferred_sheet=None):
            names = info["sheets"]
            if not names:
                self.error("Workbook contains no worksheets.")
                return
            if len(names) > 1:
                sheet, ok = QInputDialog.getItem(
                    self,
                    "Choose worksheet",
                    Path(info["path"]).name,
                    names,
                    names.index(preferred_sheet) if preferred_sheet in names else 0,
                    False,
                )
                if not ok:
                    return
            else:
                sheet = names[0]
            self.launch(
                "inspect",
                {"path": info["path"], "sheet": sheet, "signature": info["signature"]},
                lambda preview: self.mapping_dialog(role, info["path"], preview),
            )

        def mapping_dialog(self, role, path, info):
            self.mapping_window = MappingDialog(self, role, path, info)
            if self.mapping_window.exec() == QDialog.DialogCode.Accepted:
                self.accept_file(
                    role,
                    path,
                    info,
                    self.mapping_window.mapping,
                    self.mapping_window.crs_value,
                )

        def accept_file(self, role, path, info, mapping, crs="EPSG:4326"):
            if self.files[role] and (
                self.files[role].edits or self.files[role].excluded
            ):
                if (
                    QMessageBox.question(
                        self,
                        "Replace working copy?",
                        "Replace this input and its working-copy changes? Save a project first if you want to keep them.",
                    )
                    != QMessageBox.StandardButton.Yes
                ):
                    return
            if self.project_active:
                self.project_dirty = True
            validate_mapping(info["headers"], mapping)
            self.files[role] = FileSpec(
                path,
                tuple(mapping.items()),
                tuple(info["signature"]),
                info.get("sheet"),
                info["rows"],
                horizontal_crs(crs).to_string(),
                (info.get("sha256") or input_fingerprint(path)),
                (
                    "clipboard"
                    if Path(path).parent == Path(self.clipboard_store.name)
                    else "file"
                ),
            )
            self.preview_table.set_frame(info["sample"])
            self.load_summary.setText(
                f"Latest preview: {role} · {Path(path).name} · {info['rows']:,} rows (first 12 shown)"
            )
            self.update_file_labels()
            self.invalidate_review()
            record = {"path": str(Path(path).resolve()), "sheet": info.get("sheet")}
            recent = [
                e
                for e in self.settings.get("recent_files", [])
                if isinstance(e, dict)
                and (os.path.normcase(e.get("path", "")), e.get("sheet"))
                != (os.path.normcase(record["path"]), record["sheet"])
            ]
            self.settings["recent_files"] = (
                ([record] + recent[:7])
                if self.files[role].origin == "file"
                else recent[:8]
            )
            self.persist_settings()
            self.tabs.setCurrentIndex(0)
            self.status.setText(
                f"Mapped {role}. "
                + (
                    "Both files ready to validate."
                    if all(self.files.values())
                    else "Load the other file next."
                )
            )
            self.log(f"Mapped {role}: {path}")

        def recent(self, role):
            entries = [
                e
                for e in self.settings.get("recent_files", [])
                if isinstance(e, dict) and e.get("path")
            ]
            if not entries:
                self.error("No recent files yet. Browse and confirm a mapping first.")
                return
            names = [
                f"{i + 1}. {e['path']}" + (f" · {e['sheet']}" if e.get("sheet") else "")
                for i, e in enumerate(entries)
            ]
            choice, ok = QInputDialog.getItem(
                self, "Recent files", f"Open as {role}", names, 0, False
            )
            if ok:
                record = entries[names.index(choice)]
                self.open_input(role, record["path"], record.get("sheet"))

        def update_file_labels(self):
            for role, file in self.files.items():
                self.file_labels[role].setText(
                    "No file selected"
                    if file is None
                    else f"{Path(file.path).name}\n{file.rows:,} rows · {file.sheet or 'Columns mapped'}\n{file.crs}"
                )
                self.file_labels[role].setToolTip(file.path if file else "")

        def swap_files(self):
            if self.process is not None:
                return
            if not all(self.files.values()):
                self.error("Map both files before swapping.")
                return
            if self.project_active:
                self.project_dirty = True
            self.files["source"], self.files["target"] = (
                self.files["target"],
                self.files["source"],
            )
            self.update_file_labels()
            self.invalidate_review()
            self.preview_table.set_frame(pd.DataFrame())
            self.load_summary.setText(
                "Source and target swapped. Validate the new direction."
            )

        def primary_action(self):
            if self.tabs.currentIndex() == self.geo_tab_index:
                self.run_geo_analysis()
                return
            if self.process is not None:
                return
            if self.frames is None:
                self.validate_run()
            else:
                self.calculate_reviewed()

        def validate_run(self):
            if self.process is not None:
                return
            try:
                spec = self.current_spec()
            except ValueError as error:
                self.error(error)
                return
            self.invalidate_review()
            self.run_spec = spec
            self.exclude_flags.setChecked(
                self.exclude_flags.isChecked()
                if self.project_exclude_flags is None
                else self.project_exclude_flags
            )
            self.project_exclude_flags = None
            self.launch("validate", spec, self.validated)

        def validated(self, value):
            self.frames, self.audits = value
            self.review_table.clear_filters()
            self.show_audit()
            self.tabs.setCurrentIndex(1)
            self.primary.setText("Match reviewed data")
            self.status.setText(
                "Review affected rows, then match."
                if all(not f.empty for f in self.frames)
                else "One or both files have no eligible sites. Review affected rows and correct the input."
            )

        def show_audit(self, *args):
            if self.audits is None:
                return
            lines = []
            for role, audit in zip(("Source", "Target"), self.audits):
                rejected = int(
                    audit["GeoMatch disposition"]
                    .isin(["Rejected", "Excluded by user"])
                    .sum()
                )
                flagged = int(
                    audit["GeoMatch disposition"].eq("Flagged — retained").sum()
                )
                eligible = (
                    len(audit)
                    - rejected
                    - (flagged if self.exclude_flags.isChecked() else 0)
                )
                warned = int(
                    audit.get("GeoMatch coordinate warning", pd.Series(dtype=str))
                    .fillna("")
                    .ne("")
                    .sum()
                )
                lines.append(
                    f"{role}: {eligible:,} eligible · {rejected:,} rejected · {flagged:,} outliers · {warned:,} coordinate warnings"
                )
            self.review_summary.setText("\n".join(lines))
            audit = self.audits[self.review_role.currentIndex()].copy()
            if self.exclude_flags.isChecked():
                audit.loc[
                    audit["GeoMatch disposition"].eq("Flagged — retained"),
                    "GeoMatch disposition",
                ] = "Excluded by user (outlier)"
            self.review_table.set_frame(audit)

        def show_affected(self):
            self.review_table.affected_rows()

        def calculate_reviewed(self):
            if self.process is not None or self.frames is None:
                return
            try:
                current = self.current_spec()
                if current != self.run_spec:
                    self.invalidate_review()
                    self.error("Settings changed. Validate again before matching.")
                    return
                frames = [
                    (
                        f.loc[f["Flag"].eq("")].copy()
                        if self.exclude_flags.isChecked()
                        else f.copy()
                    )
                    for f in self.frames
                ]
                if any(f.empty for f in frames):
                    raise ValueError(
                        "Both files need eligible sites. Review rejected records or retain flagged rows."
                    )
            except ValueError as error:
                self.error(error)
                return
            if (
                current.terrain
                and QMessageBox.question(
                    self,
                    "Terrain check",
                    "Send path coordinates to Open-Elevation?\n\nNearest target only; sampled terrain, 5 GHz, 60% Fresnel clearance, K=4/3. Internet required.",
                )
                != QMessageBox.StandardButton.Yes
            ):
                return
            audits = [a.copy() for a in self.audits]
            excluded = self.exclude_flags.isChecked()
            if excluded:
                for audit in audits:
                    audit.loc[
                        audit["GeoMatch disposition"].eq("Flagged — retained"),
                        "GeoMatch disposition",
                    ] = "Excluded by user (outlier)"
            started = time.monotonic()

            def complete(result):
                self.mark_geo_stale()
                self.result, self.result_spec, self.result_audits = (
                    result,
                    current,
                    audits,
                )
                self.result_frames = frames
                self.result_excluded = excluded
                self.result_created = datetime.now().isoformat(timespec="seconds")
                self.result_table.clear_filters()
                self.result_table.set_frame(result)
                self.spatial.set_frame(result)
                matched = result.loc[
                    result["Match status"].eq("Matched"), "Source row"
                ].nunique()
                self.result_summary.setText(
                    f"{matched:,} matched · {len(frames[0]) - matched:,} unmatched · {len(result):,} result rows\n{current.method} · {current.mode} · {time.monotonic() - started:.1f} s"
                )
                self.stale.hide()
                self.tabs.setCurrentIndex(2)
                self.export_button.setEnabled(True)
                self.steps[3].setEnabled(True)
                self.primary.setText("Match again")
                self.status.setText(
                    "Matching complete. Inspect results or export the full report."
                )
                self.log(f"Matched data: {len(result):,} result rows.")

            self.launch("calculate", (frames[0], frames[1], current), complete)

        def export(self):
            if self.tabs.currentIndex() == self.geo_tab_index:
                self.export_geo_dialog()
                return
            if self.process is not None or self.result is None:
                return
            path, selected = QFileDialog.getSaveFileName(
                self,
                "Export full report or GIS data",
                str(
                    Path(self.settings.get("export_dir") or ".")
                    / f"GeoMatch_{datetime.now():%Y%m%d_%H%M%S}"
                ),
                "Excel workbook (*.xlsx);;GeoPackage (*.gpkg);;GeoJSON (*.geojson)",
            )
            if path:
                extension = (
                    ".gpkg"
                    if "GeoPackage" in selected
                    else ".geojson" if "GeoJSON" in selected else ".xlsx"
                )
                if not Path(path).suffix:
                    path += extension
                    if (
                        Path(path).exists()
                        and QMessageBox.question(
                            self, "Replace existing file?", f"Replace {path}?"
                        )
                        != QMessageBox.StandardButton.Yes
                    ):
                        return
                self.export_to(path)

        def export_to(self, path):
            if self.process is not None or self.result is None:
                return
            try:
                self.protect_inputs(path)
            except ValueError as error:
                self.error(error)
                return
            suffix = Path(path).suffix.lower()
            if suffix not in (".xlsx", ".gpkg", ".geojson"):
                self.error("Choose an .xlsx, .gpkg or .geojson output filename.")
                return
            import pyproj

            metadata = {
                "Product": TITLE,
                "Version": VERSION,
                "Calculated at": self.result_created,
                "Runtime versions": json.dumps(
                    {
                        "Python": sys.version.split()[0],
                        "numpy": np.__version__,
                        "pandas": pd.__version__,
                        "pyproj": pyproj.__version__,
                        "PROJ": pyproj.proj_version_str,
                    }
                ),
                "Run specification": json.dumps(
                    asdict(self.result_spec), ensure_ascii=False
                ),
                "User excluded flagged outliers": self.result_excluded,
                "Coordinate operations": json.dumps(
                    [
                        f.attrs.get("coordinate_operation", {})
                        for f in self.result_frames
                    ],
                    ensure_ascii=False,
                ),
                "Input SHA256": json.dumps(
                    [f.attrs.get("sha256", "") for f in self.result_frames]
                ),
                "Calculation CRS": "WGS84 EPSG:4326; longitude/latitude",
                "Distance units": "metres",
                "Height units": "metres; common vertical reference; no vertical transformation",
                "Tie rule": "Equal distances keep input order; duplicate IDs keep first valid row; mutual nearest uses this rule in both directions",
                "Attribute rule": "Exact case-sensitive text after trimming; blanks rejected when enabled",
                "GIS connectors": "Endpoint connectors, split at antimeridian; not routed paths or sampled geodesics",
                "GIS point scope": "All eligible source and target points used in this run; unmatched sources retained",
                "Terrain assumptions": "5 GHz; 60% Fresnel; K=4/3; nearest target only; sampled terrain",
            }

            def complete(saved):
                self.last_export = saved
                self.settings["export_dir"] = str(Path(saved).parent)
                self.persist_settings()
                self.folder_button.setEnabled(True)
                self.status.setText(
                    f"Saved {Path(saved).name} · Complete run, regardless of table filters."
                )
                self.log(f"Exported: {saved}")

            self.export_temporary = str(
                Path(path).with_name(
                    f".geomatch-{os.getpid()}-{time.time_ns()}{suffix}"
                )
            )
            if suffix == ".xlsx":
                self.launch(
                    "export",
                    (
                        path,
                        self.result,
                        self.result_audits,
                        metadata,
                        self.export_temporary,
                    ),
                    complete,
                )
            else:
                self.launch(
                    "export_gis",
                    (
                        path,
                        self.result,
                        self.result_frames,
                        self.result_audits,
                        metadata,
                        self.export_temporary,
                    ),
                    complete,
                )

        def open_folder(self):
            if self.last_export:
                QDesktopServices.openUrl(
                    QUrl.fromLocalFile(str(Path(self.last_export).parent))
                )

        def invalidate_single(self, *args):
            self.single_value = None
            if hasattr(self, "copy_single"):
                self.copy_single.setEnabled(False)
                self.single_result.setText(
                    "Coordinates changed. Calculate to update the result."
                )

        def paste_pair(self, offset):
            try:
                pair = parse_coordinate_pair(application.clipboard().text())
                for entry, value in zip(self.coordinates[offset : offset + 2], pair):
                    entry.setText(f"{value:.12g}")
            except ValueError as error:
                self.single_result.setText(str(error))

        def swap_endpoints(self):
            values = [e.text() for e in self.coordinates]
            for entry, value in zip(self.coordinates, values[2:] + values[:2]):
                entry.setText(value)

        def single(self):
            if self.process is not None:
                return
            try:
                values = [
                    parse_angle(e.text(), "latitude" if i % 2 == 0 else "longitude")
                    for i, e in enumerate(self.coordinates)
                ]
                if any(
                    abs(v) > (90 if i % 2 == 0 else 180) for i, v in enumerate(values)
                ):
                    raise ValueError(
                        "Latitude must be within −90…90 and longitude within −180…180."
                    )
            except ValueError as error:
                self.single_result.setText(str(error))
                return

            def complete(value):
                distance, bearing = value
                self.single_value = (
                    f"Distance: {distance:,.2f} m ({distance / 1000:,.4f} km)\nBearing: "
                    + (
                        f"{bearing:.2f}°"
                        if distance > 0.01
                        else "undefined at the same location"
                    )
                )
                self.single_result.setText(self.single_value)
                self.copy_single.setEnabled(True)

            self.launch("single", values, complete)

        def save_preset(self):
            try:
                radius_in_metres(
                    self.radius.text(),
                    self.radius_unit.currentText(),
                    self.unlimited.isChecked(),
                )
            except ValueError as error:
                self.error(error)
                return
            name, ok = QInputDialog.getText(self, "Save preset", "Preset name")
            if not ok or not name.strip():
                return
            name = name.strip()
            self.settings.setdefault("presets", {})[name] = {
                "method": self.method.currentText(),
                "mode": self.mode.currentText(),
                "top_n": str(self.top_n.value()),
                "radius": self.radius.text(),
                "radius_unit": self.radius_unit.currentText(),
                "unlimited": self.unlimited.isChecked(),
                "ignore": self.ignore.isChecked(),
                "flag": self.flag.isChecked(),
                "terrain": self.terrain.isChecked(),
                "match_mode": self.match_mode.currentText(),
                "match_attributes": self.match_attributes.isChecked(),
            }
            self.preset.clear()
            self.preset.addItems(list(self.settings["presets"]))
            self.preset.setCurrentText(name)
            self.persist_settings()

        def load_preset(self):
            preset = self.settings.get("presets", {}).get(self.preset.currentText())
            if not isinstance(preset, dict):
                self.error("Select a saved preset first.")
                return
            try:
                count = int(preset.get("top_n", 1))
                radius = str(preset.get("radius", 0))
                unlimited = preset.get("unlimited", float(radius) == 0)
                unit = preset.get("radius_unit", "m")
                if (
                    not 1 <= count <= 10
                    or preset.get("method", METHODS[0]) not in METHODS
                    or preset.get("mode", MODES[0]) not in MODES
                ):
                    raise ValueError("Preset contains unsupported settings.")
                radius_in_metres(radius, unit, unlimited)
            except (ValueError, TypeError) as error:
                self.error(error)
                return
            self.match_mode.setCurrentText(preset.get("match_mode", "Nearest N"))
            self.match_attributes.setChecked(
                bool(preset.get("match_attributes", False))
            )
            self.method.setCurrentText(preset.get("method", METHODS[0]))
            self.mode.setCurrentText(preset.get("mode", MODES[0]))
            self.top_n.setValue(count)
            self.radius.setText(radius)
            self.radius_unit.setCurrentText(unit)
            self.unlimited.setChecked(unlimited)
            for key in ("ignore", "flag", "terrain"):
                getattr(self, key).setChecked(bool(preset.get(key, False)))
            self.invalidate_review()

        def samples(self):
            folder = QFileDialog.getExistingDirectory(self, "Save sample files")
            if not folder:
                return
            try:
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
                for role, rows in (
                    (
                        "Source",
                        [("S001", 30, 31, 25, 40), ("S002", 30.01, 31.01, 30, 45)],
                    ),
                    (
                        "Target",
                        [
                            ("T001", 30.005, 31.004, 35, 42),
                            ("T002", 30.02, 31.02, 40, 50),
                        ],
                    ),
                ):
                    pd.DataFrame(rows, columns=FIELDS[:5]).to_csv(
                        Path(folder) / f"GeoMatch_{role}_{stamp}.csv", index=False
                    )
                self.status.setText("Saved source and target sample files.")
            except OSError as error:
                self.error(error)

        def help_window(self, key, title, subtitle, width=720, height=580):
            dialogs = getattr(self, "help_dialogs", {})
            self.help_dialogs = dialogs
            if key in dialogs:
                dialog = dialogs[key]
                dialog.show()
                dialog.raise_()
                dialog.activateWindow()
                return dialog, None
            dialog = QDialog(self)
            dialog.setWindowTitle(f"{title} — GeoMatch")
            dialog.setWindowIcon(self.windowIcon())
            dialog.setSizeGripEnabled(True)
            area = self.screen().availableGeometry()
            dialog.resize(
                min(width, area.width() - 60), min(height, area.height() - 80)
            )
            layout = QVBoxLayout(dialog)
            layout.setContentsMargins(24, 20, 24, 20)
            layout.setSpacing(14)
            heading = QHBoxLayout()
            logo = QLabel()
            logo.setPixmap(self.windowIcon().pixmap(52, 52))
            heading.addWidget(logo)
            text = QVBoxLayout()
            text.addWidget(label(title, "heading"))
            text.addWidget(label(subtitle, "muted"))
            heading.addLayout(text, 1)
            layout.addLayout(heading)
            dialogs[key] = dialog
            return dialog, layout

        def help_browser(self, html):
            browser = QTextBrowser()
            browser.setOpenExternalLinks(False)
            browser.setOpenLinks(False)
            browser.help_html = html
            self.style_help_browser(browser)
            browser.setAccessibleName("Help content")
            browser.anchorClicked.connect(
                lambda url: (
                    browser.scrollToAnchor(url.fragment()) if not url.scheme() else None
                )
            )
            return browser

        def style_help_browser(self, browser):
            position = browser.verticalScrollBar().value()
            dark = (
                application.palette().color(QPalette.ColorRole.Base).lightness() < 128
            )
            link = "#78caff" if dark else "#075fa8"
            browser.document().setDefaultStyleSheet(
                f"a {{ color: {link}; text-decoration: underline; }} p, li {{ line-height: 130%; }}"
            )
            browser.setHtml(browser.help_html)
            browser.verticalScrollBar().setValue(position)

        def help_finish(self, dialog, layout, buttons=()):
            footer = QHBoxLayout()
            for title, callback in buttons:
                footer.addWidget(button(title, callback))
            footer.addStretch()
            close = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
            close.rejected.connect(dialog.close)
            footer.addWidget(close)
            layout.addLayout(footer)
            dialog.show()
            dialog.raise_()
            dialog.activateWindow()

        def about(self):
            dialog, layout = self.help_window(
                "about",
                "About GeoMatch",
                "Spatial Matching & Proximity Analysis",
                680,
                540,
            )
            if layout is None:
                return
            browser = self.help_browser(f"""
                <h2>GeoMatch</h2>
                <p>A desktop workspace for spatial matching, proximity analysis and practical geospatial engineering checks.</p>
                <table cellpadding="9" cellspacing="0">
                <tr><td><b>Version</b></td><td>{VERSION}</td></tr>
                <tr><td><b>Build</b></td><td>{BUILD_NUMBER}</td></tr>
                <tr><td><b>Release date</b></td><td>{BUILD_DATE}</td></tr>
                <tr><td><b>Author</b></td><td>{AUTHOR}</td></tr>
                </table>
                <h3>Tools in this release</h3>
                <p>Point matching · CRS conversion · Survey comparisons · Road and polygon proximity · Buffers and nearby counts · Local DEM profiles · Telecom planning sectors.</p>
                <p>Use <b>Quick Guide</b> for the complete workflow, or <b>Contact us / Support</b> to reach the author.</p>
                """)
            dialog.content = browser
            layout.addWidget(browser, 1)
            dialog.copy_button = button(
                "Copy version details",
                lambda: application.clipboard().setText(
                    f"{TITLE}\nVersion: {VERSION}\nBuild: {BUILD_NUMBER}\nRelease date: {BUILD_DATE}\nAuthor: {AUTHOR}"
                ),
            )
            layout.addWidget(dialog.copy_button)
            self.help_finish(
                dialog,
                layout,
                (
                    ("Quick Guide", self.quick_guide),
                    ("Contact / Support", self.contact_support),
                ),
            )

        def quick_guide(self):
            dialog, layout = self.help_window(
                "guide",
                "Quick Guide",
                "From loading your data to reviewing and exporting an analysis",
                940,
                760,
            )
            if layout is None:
                return
            sections = [
                (
                    "start",
                    "1. Start and choose a workflow",
                    """
                <ol><li>For site-to-site distances, follow <b>Load → Review → Match → Export</b>.</li>
                <li>For roads, polygons, surveys, buffers or sectors, map your points on <b>Load</b>, then switch to <b>Geo Tools</b>.</li>
                <li>For a local DEM profile, calculate point matches first. The DEM tool uses those completed links.</li>
                <li>Use Light/Dark/System and the text-size selector to adjust the interface. Most tables support sorting, searching, column filters and copying.</li></ol>""",
                ),
                (
                    "load",
                    "2. Load source and target data",
                    """
                <ol><li><b>Source</b> is the set of points you want to analyse; <b>Target</b> contains candidate points. Click <b>Browse &amp; map</b> under each.</li>
                <li>Choose CSV, TXT, XLSX, XLS or WGS84 point GeoJSON. Select the worksheet for Excel inputs. Alternatively, drop one file and choose its role.</li>
                <li>To paste from Excel, copy a table <b>including its column headers</b>, then use <b>Paste table from Excel</b>.</li>
                <li>Map a unique <b>Site ID</b>, X/longitude/easting and Y/latitude/northing. Map optional height and match-attribute fields when needed.</li></ol>
                <p>Identifiers retain leading zeros. Original input files are not edited by GeoMatch.</p>""",
                ),
                (
                    "crs",
                    "3. Confirm coordinates and CRS",
                    """
                <ol><li>Choose the input CRS from the dataset's specification. Search EPSG if necessary. Do not infer it solely from how numbers look.</li>
                <li>Check X/Y mapping: <b>X = longitude/easting; Y = latitude/northing</b>. Each input may have a different CRS.</li>
                <li>Click <b>Preview converted coordinates</b>. Inspect the normalized WGS84 coordinates and warnings.</li>
                <li>Tick the CRS/mapping confirmation, then <b>Use mapping</b>. A mapping or CRS change requires confirmation again.</li></ol>
                <p>Geographic inputs accept decimal degrees or DMS. Separate pasted DMS pairs with a comma or semicolon. Heights are in metres with a common vertical reference; no vertical datum transformation is performed.</p>""",
                ),
                (
                    "settings",
                    "4. Choose point-matching settings",
                    """
                <ol><li>Select <b>Nearest N</b>, <b>Nearest eligible</b>, <b>All within radius</b>, or <b>Mutual nearest</b>. Nearest N allows 1–10 targets per source.</li>
                <li>Choose WGS84 geodesic or spherical Haversine distance, and 2D ground or approximate 3D slant distance.</li>
                <li>Set a distance limit in metres/kilometres, or use Unlimited. <b>All within radius requires a positive limit</b>.</li>
                <li>For 3D, provide absolute antenna elevation, or terrain elevation plus nonnegative antenna AGL.</li>
                <li>Optionally ignore identical locations, flag outliers, or require equal match attributes.</li></ol>
                <p>Identical-location exclusion uses ground distance ≤ 0.01 m. Attribute equality is case-sensitive after trimming; blank attributes are rejected when enabled. Mutual nearest is not a global assignment. The existing online terrain option asks before sending coordinates.</p>""",
                ),
                (
                    "review",
                    "5. Validate, correct and review",
                    """
                <ol><li>Click <b>Validate &amp; review</b>. Select Source or Target and use <b>Show affected rows</b>.</li>
                <li>Inspect rejection reasons, duplicate IDs, outliers and coordinate warnings. Duplicate IDs keep the first eligible valid occurrence.</li>
                <li>To fix data without altering the source file, choose <b>Edit working copy</b> beneath that input. Double-click a cell or press F2.</li>
                <li>Use <b>Exclude selected rows</b> / <b>Include selected rows</b>, and Undo/Redo. Green means edited; orange means excluded.</li>
                <li>Apply changes, then validate again. Choose whether to exclude flagged outliers before matching.</li></ol>
                <p>Undo/Redo history lasts for the current editor dialog. Save a project to retain applied edits and exclusions between sessions.</p>""",
                ),
                (
                    "results",
                    "6. Match and inspect results",
                    """
                <ol><li>Click <b>Match reviewed data</b>. Use <b>Stop</b> or Esc to cancel an active job.</li>
                <li>Review matched and unmatched records in Results. Click a column to sort; right-click its heading to filter.</li>
                <li>Use Search, Columns and Copy cells. Copied/exported numeric values keep their full precision.</li>
                <li>Select rows to highlight their points and connections in the offline spatial preview. Click a point to select a corresponding visible row.</li>
                <li>Use the mouse wheel to zoom, drag to pan and <b>Fit preview</b> to reset.</li></ol>
                <p>The preview is a sampled coordinate diagram, not a basemap or distance scale. Table filters do not filter the diagram. Changed settings mark previous results; rerun to update them.</p>""",
                ),
                (
                    "export",
                    "7. Export the complete point run",
                    """
                <ol><li>Choose <b>Export report / GIS</b> on Results.</li>
                <li><b>Excel</b>: results, source/target review sheets and settings. <b>GeoPackage</b>: point/link layers plus report tables. <b>GeoJSON</b>: eligible point/link features and metadata.</li>
                <li>Choose a separate output filename and wait for completion. Use <b>Open export folder</b> to find it.</li></ol>
                <p>Exports contain the full calculated run regardless of table filters. Excel/GeoPackage include rejected-row audits; point-run GeoJSON does not. Connection geometry represents endpoints, not a routed path.</p>""",
                ),
                (
                    "projects",
                    "8. Save and resume a project",
                    """
                <ol><li>With both point inputs mapped, open Load and choose <b>Save project</b>.</li>
                <li>The .geomatch file saves paths, mappings, CRS, settings, working edits, exclusions and Geo Tools controls. Pasted clipboard inputs are embedded.</li>
                <li>Use <b>Open project</b> to resume. Ordinary disk inputs must remain at their saved paths with matching fingerprints.</li>
                <li>Validate and run again. Calculated results and Undo history are not stored. Reconfirm DEM units before a new profile run.</li></ol>""",
                ),
                (
                    "large",
                    "9. Process large point datasets",
                    """
                <ol><li>Use a CSV/TXT source and map both inputs. Select Nearest N, Nearest eligible, or All within radius.</li>
                <li>Turn outlier and online terrain checks off. Mutual nearest uses the regular reviewed workflow.</li>
                <li>On Load, choose <b>Large-file run to ZIP</b>. It validates source chunks and writes results and audit CSVs incrementally.</li>
                <li>Inspect the ZIP reports after completion. Import ID columns as Text when opening CSVs in Excel.</li></ol>
                <p>There is no full Results-table preview for this route. Targets, their index, unique source IDs and edits still use memory; allow disk space for temporary report files. Geo Tools operations are not streamed by this mode.</p>""",
                ),
                (
                    "geo",
                    "10. Roads, polygons, surveys and buffers",
                    """
                <ol><li>Open <b>Geo Tools</b> and select an operation. Enter a suitable <b>projected analysis CRS in metres</b>; geographic coordinates, Web Mercator and foot-based analysis CRSs are rejected.</li>
                <li><b>Nearest line / road</b>: map source points, select a line GeoPackage/GeoJSON and enter its layer name if needed. The result gives nearest feature, foot coordinates and planar distance.</li>
                <li><b>Polygon membership / proximity</b>: use a polygon layer. Inspect Inside/Boundary/Outside, covering polygons and boundary distance. Holes count as outside.</li>
                <li><b>Survey comparison by ID</b>: Source = baseline; Target = observed. Set horizontal and optional vertical tolerances. Deltas are observed minus baseline; IDs are compared exactly after trimming.</li>
                <li><b>Buffers and nearby counts</b>: load source and target, set radius, then run. Counts use projected centre distance; buffer polygons are approximations.</li>
                <li>Review Results and audit tables. Use <b>Export analysis</b> for Excel or GeoPackage. Export before another tool replaces the current analysis.</li></ol>
                <p>Line/polygon tools need source points only. Other tools here need source and target. Planar distances include the chosen projection's distortion; choose a CRS appropriate for the full study area.</p>""",
                ),
                (
                    "dem",
                    "11. Profile links using a local DEM",
                    """
                <ol><li>Calculate point matches first. The DEM tool uses <b>all matched result rows</b>, including additional ranks; filters do not restrict them.</li>
                <li>In Geo Tools choose <b>DEM profiles of matched links</b>, then browse to a local GeoTIFF with a CRS and elevation band.</li>
                <li>Choose band, units after band scale/offset, sample spacing, source/target antenna AGL, frequency, Fresnel fraction and Earth-radius factor K.</li>
                <li>Confirm units and vertical reference, then run. Panel AGL values are constant heights above the DEM for every endpoint; mapped height fields are not used here.</li>
                <li>Select a link beneath the table to view the chart, or choose Profile samples for detailed elevations and clearance margins.</li>
                <li>Export analysis to Excel or GeoPackage.</li></ol>
                <p>Sampling uses nearest raster cells. Missing data gives Unknown; zero-length links are not evaluated. Limits: 10,000 samples/link and 200,000 total. “Clear at sampled cells only” is not a complete obstruction or RF propagation survey.</p>""",
                ),
                (
                    "sectors",
                    "12. Create telecom planning sectors",
                    """
                <ol><li>Map source sites and target candidates. In Geo Tools choose <b>Telecom sectors</b> and enter the projected CRS.</li>
                <li>Set a shared azimuth clockwise from true north (0–&lt;360°), beamwidth (&gt;0–360°) and range in metres.</li>
                <li>Run to create one sector per source and count targets within the WGS84 bearing/range limits.</li>
                <li>Export GeoPackage polygons, or the Excel report. Use separate runs for sites needing different sector parameters.</li></ol>
                <p>These are geometric planning shapes, not predicted radio coverage. They do not model propagation, antenna patterns, tilt, interference or terrain loss.</p>""",
                ),
                (
                    "shortcuts",
                    "13. Shortcuts and troubleshooting",
                    """
                <p><b>F1</b>: product Help · <b>Shift+F1</b>: Quick Guide<br>
                <b>Ctrl+O / Ctrl+Shift+O</b>: source / target<br>
                <b>Ctrl+Enter</b>: active matching, Single Link or Geo Tools action<br>
                <b>Ctrl+E</b>: export current Results or Geo Tools analysis<br>
                <b>Ctrl+Shift+S / Ctrl+Shift+P</b>: save / open project<br>
                <b>Ctrl+1…6</b>: tabs · <b>Esc</b>: stop active work (or close a help popup)<br>
                <b>Ctrl+C</b>: copy table cells · <b>Ctrl+Z / Ctrl+Y</b>: Undo/Redo in working editor</p>
                <ul><li><b>No eligible points:</b> inspect review reasons, mappings, height requirements and exclusions.</li>
                <li><b>Input changed:</b> import/remap it again, or restore the exact file referenced by a project.</li>
                <li><b>CRS/transform error:</b> verify CRS, axis mapping and area of use. Required local transformation resources must be available.</li>
                <li><b>Missing DEM dependency:</b> install the supplied requirements, then restart.</li>
                <li><b>Need assistance:</b> open Contact us / Support and include version/build, reproduction steps, expected result and error text from Activity.</li></ul>""",
                ),
            ]
            chapters = combo([title for _, title, _ in sections])
            chapters.setAccessibleName("Quick Guide section")
            layout.addWidget(chapters)
            html = (
                "<p><b>Choose a section above or scroll through the complete guide.</b></p>"
                + "".join(
                    f'<h2 id="{key}"><a name="{key}"></a>{title}</h2>{body}<hr>'
                    for key, title, body in sections
                )
            )
            browser = self.help_browser(html)
            chapters.currentIndexChanged.connect(
                lambda index: browser.scrollToAnchor(sections[index][0])
            )
            dialog.content = browser
            dialog.chapters = chapters
            layout.addWidget(browser, 1)
            self.help_finish(
                dialog, layout, (("Contact / Support", self.contact_support),)
            )

        def contact_support(self):
            dialog, layout = self.help_window(
                "contact",
                "Contact us / Support",
                "Connect with Mohamed Ashraf, author of GeoMatch",
                740,
                580,
            )
            if layout is None:
                return
            browser = self.help_browser(
                f"""
                <h3>Questions, feedback or a reproducible issue?</h3>
                <p>Use email for support enquiries and feedback, or visit the author's LinkedIn profile.</p>
                <p><b>Email</b><br><a href="mailto:{SUPPORT_EMAIL}">{SUPPORT_EMAIL}</a></p>
                <p><b>LinkedIn</b><br><a href="{SUPPORT_LINKEDIN}">{SUPPORT_LINKEDIN}</a></p>
                <h3>What to include</h3>
                <ol><li>Your GeoMatch version/build and operation.</li><li>Steps to reproduce the issue.</li>
                <li>Expected result, actual result and relevant Activity error text.</li></ol>
                <p><b>This installation:</b> Version {VERSION} · Build {BUILD_NUMBER}</p>
                <p>Email opens a draft in your default mail application. Nothing is sent automatically. Copy the address if no mail application is configured.</p>"""
            )

            def open_contact(url):
                value = url.toString()
                if value not in (f"mailto:{SUPPORT_EMAIL}", SUPPORT_LINKEDIN):
                    return
                if not QDesktopServices.openUrl(url):
                    dialog.feedback.setText(
                        "Could not open the application. Copy the contact details and open them manually."
                    )

            browser.anchorClicked.connect(open_contact)
            dialog.content = browser
            layout.addWidget(browser, 1)
            actions = QHBoxLayout()
            email = QUrl(f"mailto:{SUPPORT_EMAIL}")
            query = QUrlQuery()
            query.addQueryItem(
                "subject", f"GeoMatch support — v{VERSION} / build {BUILD_NUMBER}"
            )
            email.setQuery(query)

            def compose():
                if not QDesktopServices.openUrl(email):
                    dialog.feedback.setText(
                        "No email application opened. Use Copy email, then compose a message in your mail service."
                    )

            dialog.email_button = button("Compose email", compose)
            dialog.linkedin_button = button(
                "Open LinkedIn", lambda: open_contact(QUrl(SUPPORT_LINKEDIN))
            )
            dialog.copy_button = button(
                "Copy email",
                lambda: (
                    application.clipboard().setText(SUPPORT_EMAIL),
                    dialog.feedback.setText("Email address copied."),
                ),
            )
            for item in (
                dialog.email_button,
                dialog.linkedin_button,
                dialog.copy_button,
            ):
                actions.addWidget(item)
            layout.addLayout(actions)
            dialog.feedback = label("", "muted")
            layout.addWidget(dialog.feedback)
            self.help_finish(
                dialog,
                layout,
                (
                    (
                        "Copy support details",
                        lambda: (
                            application.clipboard().setText(
                                f"{TITLE}\nVersion: {VERSION}\nBuild: {BUILD_NUMBER}\nAuthor: {AUTHOR}\nEmail: {SUPPORT_EMAIL}\nLinkedIn: {SUPPORT_LINKEDIN}"
                            ),
                            dialog.feedback.setText("Support details copied."),
                        ),
                    ),
                ),
            )

    window = App()
    window._application = application
    window.show()
    if run:
        application.exec()
    return window


def main():
    import argparse

    parser = argparse.ArgumentParser(description=TITLE)
    parser.add_argument(
        "--smoke",
        action="store_true",
        help="Open and close a Qt window using temporary preferences",
    )
    args = parser.parse_args()
    try:
        if args.smoke:
            with tempfile.TemporaryDirectory(prefix="sitespan-qt-smoke-") as directory:
                window = start_gui(run=False, settings_dir=directory)
                window._application.processEvents()
                window.close()
                window._application.processEvents()
            print("GeoMatch Qt startup self-check passed.")
        else:
            start_gui()
        return 0
    except Exception:
        detail = traceback.format_exc()
        folder = (
            Path(os.environ.get("LOCALAPPDATA", str(Path.home() / ".config")))
            / "SiteSpan"
        )
        try:
            folder.mkdir(parents=True, exist_ok=True)
            (folder / "startup-error.txt").write_text(detail, encoding="utf-8")
        except OSError:
            pass
        if sys.stderr is not None:
            print(detail, file=sys.stderr)
        return 1


if __name__ == "__main__":
    mp.freeze_support()
    raise SystemExit(main())
