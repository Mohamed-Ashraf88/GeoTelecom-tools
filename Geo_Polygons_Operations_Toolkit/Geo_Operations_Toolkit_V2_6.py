"""Geo Operations Toolkit 2.6 — single-file desktop application.

Requires Python 3.11+. Install the UI and processing dependencies once:
    python -m pip install PySide6-Essentials shapely pyproj pandas openpyxl defusedxml
Optional Shapefile/GeoPackage input:
    python -m pip install geopandas pyogrio
Run:
    python Geo_Operations_Toolkit_V2_6.py

Author: Mohamed Ashraf
Help menu: About, Quick Guide (F1), and Contact & Support.

The icon is drawn from hardcoded shapes using Qt. All
application code and icon artwork are in this file; no sibling modules,
image files, or resources are needed. Third-party processing libraries are
imported only when a function needs them. Settings/logs use the user's app
data directory. The original input files are never modified by processing.
"""

from dataclasses import dataclass, field
from typing import Any
import math
import re
from collections import Counter
from typing import Callable
import csv
import io
import json
import os
from pathlib import Path
import tempfile
import zipfile
from xml.etree import ElementTree as ET
import logging
from logging.handlers import RotatingFileHandler
import sys
import threading
import time

__version__ = "2.6.0"
BUILD_NUMBER = "20261001.9"
RELEASE_DATE = "1 October 2026"
AUTHOR_NAME = "Mohamed Ashraf"
SUPPORT_EMAIL = "Mohamed--Ashraf@outlook.com"
LINKEDIN_URL = "https://www.linkedin.com/in/mohamed---ashraf/"

import struct
import zlib


def create_app_icon_png(size: int = 64) -> bytes:
    """Draw overlapping map polygons and a surveyed point as an RGBA PNG.

    Artwork uses a fixed 64-unit coordinate system and 2x supersampling.
    Encoding uses only the standard library; neither Pillow nor an asset file
    is required to launch the application.
    """
    if not 8 <= size <= 256:
        raise ValueError("Icon size must be between 8 and 256 pixels")
    polygon_a = ((12, 17), (34, 11), (46, 28), (35, 48), (13, 42))
    polygon_b = ((29, 24), (50, 18), (55, 43), (37, 54), (22, 42))

    def inside(x, y, vertices):
        hit = False
        previous = vertices[-1]
        for current in vertices:
            ax, ay = previous
            bx, by = current
            if (ay > y) != (by > y) and x < (bx - ax) * (y - ay) / (by - ay) + ax:
                hit = not hit
            previous = current
        return hit

    def near_edge(x, y, vertices):
        previous = vertices[-1]
        for current in vertices:
            ax, ay = previous
            bx, by = current
            dx, dy = bx - ax, by - ay
            t = max(0, min(1, ((x - ax) * dx + (y - ay) * dy) / (dx * dx + dy * dy)))
            if (x - ax - t * dx) ** 2 + (y - ay - t * dy) ** 2 <= 0.9**2:
                return True
            previous = current
        return False

    def sample(x, y):
        # Rounded navy tile, transparent outside the corner radius.
        cx, cy = min(53, max(11, x)), min(53, max(11, y))
        if (x - cx) ** 2 + (y - cy) ** 2 > 9**2:
            return (0, 0, 0, 0)
        color = (25, 49, 65, 255)
        a, b = inside(x, y, polygon_a), inside(x, y, polygon_b)
        if a:
            color = (63, 129, 177, 255)
        if b:
            color = (40, 159, 152, 255)
        if a and b:
            color = (111, 208, 190, 255)
        if near_edge(x, y, polygon_a) or near_edge(x, y, polygon_b):
            color = (206, 236, 239, 255)
        # An amber point with a white halo lies within the shared footprint.
        distance = (x - 34) ** 2 + (y - 33) ** 2
        if distance <= 6**2:
            color = (255, 255, 255, 255)
        if distance <= 4.1**2:
            color = (246, 176, 67, 255)
        return color

    scanlines = bytearray()
    for row in range(size):
        scanlines.append(0)  # PNG filter: none
        for column in range(size):
            samples = [
                sample((column + sx) * 64 / size, (row + sy) * 64 / size)
                for sy in (0.25, 0.75)
                for sx in (0.25, 0.75)
            ]
            alpha = sum(pixel[3] for pixel in samples)
            rgb = [
                round(sum(pixel[channel] * pixel[3] for pixel in samples) / alpha)
                if alpha
                else 0
                for channel in range(3)
            ]
            scanlines.extend((*rgb, round(alpha / 4)))

    def chunk(kind, payload):
        return (
            struct.pack(">I", len(payload))
            + kind
            + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 6, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(bytes(scanlines), 9))
        + chunk(b"IEND", b"")
    )


# ==========================================================================
# MODELS
# ==========================================================================

OPERATIONS = (
    "Points inside polygon",
    "Combine",
    "Intersection",
    "Difference",
    "SymDifference",
    "Buffer",
    "Simplify",
    "Area",
    "Perimeter",
    "Distance Calculator",
    "Split polygons",
    "Fill holes",
    "Geometry validation",
    "Boundary / vertices",
    "Coordinate Converter",
)

DESCRIPTIONS = {
    "Coordinate Converter": "Convert single points or tables from DD, DDM, or DMS to decimal degrees and optional UTM.",
    "Split polygons": "Divide Layer A using cutting lines or a local projected grid. Retain source attributes.",
    "Fill holes": "Fill all interior holes, or only those within a chosen area limit.",
    "Geometry validation": "Inspect original imported geometries and optionally produce valid polygon copies.",
    "Boundary / vertices": "Export polygon rings as lines, or vertices as numbered points and coordinate tables.",
    "Distance Calculator": "Find the nearest N reference sites for each query location, with WGS84 distance and bearing.",
    "Points inside polygon": "Match every point row to polygon names in Layer A. Holes are respected.",
    "Combine": "Merge Layer A into one footprint. Include Layer B when selected below.",
    "Intersection": "Keep the geographic area shared by the footprints of A and B.",
    "Difference": "Subtract the footprint of B from A. The order is A minus B.",
    "SymDifference": "Keep areas belonging to A or B, excluding their shared area.",
    "Buffer": "Expand or shrink each feature using a distance in meters or kilometers.",
    "Simplify": "Reduce vertices with a tolerance in meters, preserving each feature's topology.",
    "Area": "Measure on the WGS84 ellipsoid. Choose merged coverage or a sum of features.",
    "Perimeter": "Measure boundary length on WGS84, including the boundaries of holes.",
}


@dataclass(frozen=True)
class Feature:
    name: str
    geometry: Any
    properties: dict = field(default_factory=dict)
    source: str = ""


@dataclass
class Layer:
    features: list[Feature]
    sources: tuple[str, ...] = ()
    warnings: list[str] = field(default_factory=list)
    raw_features: list[Feature] | None = None


@dataclass(frozen=True)
class Options:
    operation: str
    distance: float = 50.0
    unit: str = "Meters"
    include_boundary: bool = True
    swap_coordinates: bool = False
    name_column: str = ""
    latitude_column: str = ""
    longitude_column: str = ""
    measurement_scope: str = "Merged coverage"
    split_mode: str = "Grid"
    grid_rows: int = 2
    grid_columns: int = 2
    fill_all_holes: bool = True
    hole_area_limit: float = 1000.0
    repair_geometry: bool = False
    boundary_output: str = "Boundaries"
    include_holes: bool = True
    include_closing_vertex: bool = False


@dataclass
class Result:
    operation: str
    features: list[Feature] = field(default_factory=list)
    tables: dict[str, list[dict]] = field(default_factory=dict)
    summary: dict[str, Any] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


class Cancelled(Exception):
    """A user requested cancellation at a safe checkpoint."""


# ==========================================================================
# ENGINE
# ==========================================================================

Progress = Callable[[float, str], None]


def checkpoint(cancel) -> None:
    if cancel is not None and cancel.is_set():
        raise Cancelled("Cancelled. Your previous result is still available.")


def polygon_parts(geometry) -> list:
    """Extract polygonal components, including those returned by make_valid."""
    if geometry is None or geometry.is_empty:
        return []
    if geometry.geom_type == "Polygon":
        return [geometry]
    if hasattr(geometry, "geoms"):
        return [part for child in geometry.geoms for part in polygon_parts(child)]
    return []


def clean_polygon(geometry):
    from shapely import force_2d
    from shapely import get_coordinates
    from shapely import make_valid
    import numpy as np
    from shapely import union_all

    geometry = force_2d(geometry)
    if not np.isfinite(get_coordinates(geometry)).all():
        raise ValueError("Geometry contains non-finite coordinates.")
    repaired = not geometry.is_valid
    if repaired:
        geometry = make_valid(geometry)
    parts = polygon_parts(geometry)
    if not parts:
        return None, repaired
    return (parts[0] if len(parts) == 1 else union_all(parts)), repaired


def validate_extent(features: list[Feature]) -> tuple:
    """Reject domains that cannot safely use ordinary longitude/latitude topology."""
    bounds = [f.geometry.bounds for f in features if not f.geometry.is_empty]
    if not bounds:
        raise ValueError("The layer contains no usable polygon area.")
    xmin, ymin = min(b[0] for b in bounds), min(b[1] for b in bounds)
    xmax, ymax = max(b[2] for b in bounds), max(b[3] for b in bounds)
    if not all(math.isfinite(v) for b in bounds for v in b):
        raise ValueError("Polygon coordinates must be finite numbers.")
    if xmin < -180 or xmax > 180 or ymin < -90 or ymax > 90:
        raise ValueError(
            "Polygon coordinates are outside WGS84 limits. Check the input CRS."
        )
    if xmax - xmin > 180:
        raise ValueError(
            "Dateline-crossing or greater-than-hemisphere data must be split before use."
        )
    return xmin, ymin, xmax, ymax


def metric_transformers(features: list[Feature]):
    from pyproj import CRS
    from pyproj import Transformer

    xmin, ymin, xmax, ymax = validate_extent(features)
    if xmax - xmin > 6 or ymax - ymin > 8 or ymin < -80 or ymax > 84:
        raise ValueError(
            "Buffer, Simplify, and grid splitting require a local extent (at most 6° wide and 8° high, within UTM latitudes). Split this layer into smaller regions."
        )
    lon, lat = (xmin + xmax) / 2, (ymin + ymax) / 2
    zone = min(60, max(1, int((lon + 180) // 6) + 1))
    epsg = (32600 if lat >= 0 else 32700) + zone
    crs = CRS.from_epsg(epsg)
    return (
        Transformer.from_crs(4326, crs, always_xy=True).transform,
        Transformer.from_crs(crs, 4326, always_xy=True).transform,
        crs.to_string(),
    )


def geodesic_metrics(geometry) -> tuple[float, float]:
    """Area excludes holes; perimeter includes exterior and interior rings."""
    from pyproj import Geod
    from shapely.geometry.polygon import orient

    geod = Geod(ellps="WGS84")
    area = perimeter = 0.0
    for part in polygon_parts(geometry):
        polygon = orient(part, sign=1.0)
        signed_area, _ = geod.geometry_area_perimeter(polygon)
        area += abs(signed_area)
        perimeter += geod.geometry_length(polygon.exterior)
        perimeter += sum(geod.geometry_length(ring) for ring in polygon.interiors)
    return area, perimeter


def parse_coordinate(value) -> float:
    """Accept decimal, scientific notation, decimal comma, DM, and DMS.

    Unrecognized text is rejected instead of stripping arbitrary characters.
    Decimal comma is supported only when there is no decimal point.
    """
    text = str(value).strip().upper().replace("−", "-")
    if not text or text in {"NAN", "NONE", "NAT", "<NA>"}:
        raise ValueError("Missing coordinate")
    direction = ""
    if text[:1] in "NSEW":
        direction, text = text[0], text[1:].strip()
    if text[-1:] and text[-1:] in "NSEW":
        if direction:
            raise ValueError("Repeated hemisphere direction")
        direction, text = text[-1], text[:-1].strip()
    if "," in text:
        if text.count(",") != 1 or "." in text:
            raise ValueError("Ambiguous decimal separator")
        text = text.replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        cleaned = re.sub(r"[°º'\"′″:_]", " ", text)
        tokens = cleaned.split()
        if len(tokens) not in (2, 3) or not all(
            re.fullmatch(r"[+-]?\d+(?:\.\d+)?", t) for t in tokens
        ):
            raise ValueError(f"Unsupported coordinate: {value}") from None
        degrees, minutes = float(tokens[0]), float(tokens[1])
        seconds = float(tokens[2]) if len(tokens) == 3 else 0.0
        if not 0 <= minutes < 60 or not 0 <= seconds < 60:
            raise ValueError("Minutes and seconds must be between 0 and 60")
        number = abs(degrees) + minutes / 60 + seconds / 3600
        if tokens[0].startswith("-"):
            number = -number
    if not math.isfinite(number):
        raise ValueError("Coordinate must be finite")
    if direction in ("N", "E") and number < 0:
        raise ValueError("Negative coordinate conflicts with N/E hemisphere")
    if direction in ("S", "W"):
        number = -abs(number)
    return number


def match_points(
    layer: Layer, table, options: Options, cancel, progress: Progress
) -> Result:
    from shapely.geometry import Point
    from shapely import STRtree

    import pandas as pd

    columns = (options.name_column, options.latitude_column, options.longitude_column)
    if len(set(columns)) != 3 or any(c not in table.columns for c in columns):
        raise ValueError(
            "Choose three different columns for point name, latitude, and longitude."
        )
    # Index individual pieces for compact bounding boxes. Keep names and holes.
    geometries, names = [], []
    for feature in layer.features:
        for part in polygon_parts(feature.geometry):
            geometries.append(part)
            names.append(feature.name)
    tree = STRtree(geometries)
    predicate = "covered_by" if options.include_boundary else "within"
    cache: dict[tuple, tuple[str, ...]] = {}
    rows, rejected = [], []
    counts = Counter({feature.name: 0 for feature in layer.features})
    selected = table.loc[:, list(columns)]
    total = len(selected)
    for index, (name, latitude, longitude) in enumerate(
        selected.itertuples(index=False, name=None), 1
    ):
        if index % 250 == 1:
            checkpoint(cancel)
            progress(
                index / max(total, 1) * 95, f"Matching point {index:,} of {total:,}"
            )
        try:
            if pd.isna(name) or not str(name).strip():
                raise ValueError("Missing point name")
            lat, lon = parse_coordinate(latitude), parse_coordinate(longitude)
            if options.swap_coordinates:
                lat, lon = lon, lat
            if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                raise ValueError(
                    "Coordinate outside latitude/longitude limits; check mapping or Swap"
                )
        except (ValueError, TypeError) as exc:
            rejected.append(
                {
                    "Input row": index + 1,
                    "Point name": str(name),
                    "Latitude input": str(latitude),
                    "Longitude input": str(longitude),
                    "Reason": str(exc),
                }
            )
            continue
        key = (lon, lat)
        if key not in cache:
            matches = sorted(tree.query(Point(lon, lat), predicate=predicate).tolist())
            cache[key] = tuple(dict.fromkeys(names[i] for i in matches))
        matched = cache[key]
        counts.update(matched)
        rows.append(
            {
                "Input row": index + 1,
                "Point name": str(name).strip(),
                "Latitude": lat,
                "Longitude": lon,
                "Polygon Name": "; ".join(matched) if matched else "Not in any polygon",
                "Match count": len(matched),
            }
        )
    checkpoint(cancel)
    inside = sum(row["Match count"] > 0 for row in rows)
    result = Result(
        options.operation,
        tables={
            "Matches": rows,
            "Empty polygons": [
                {"Polygon Name": name} for name in sorted(counts) if counts[name] == 0
            ],
            "Polygon counts": [
                {"Polygon Name": name, "Point rows": count}
                for name, count in sorted(counts.items())
            ],
            "Rejected rows": rejected,
        },
        summary={
            "Input point rows": total,
            "Valid point rows": len(rows),
            "Matched rows": inside,
            "Outside rows": len(rows) - inside,
            "Rejected rows": len(rejected),
            "Unique tested coordinates": len(cache),
            "Boundary": "Included" if options.include_boundary else "Excluded",
            "Coordinates swapped": options.swap_coordinates,
        },
        warnings=list(layer.warnings),
    )
    if not rows:
        result.warnings.append(
            "No valid point rows. Review the Rejected rows sheet and column mapping."
        )
    return result


def parse_converter_coordinate(value, axis, correct_typos=False):
    """Validate one angular coordinate without discarding unrecognized text."""
    original = str(value).strip()
    text = (
        original.upper()
        .replace("−", "-")
        .replace("’", "'")
        .replace("‘", "'")
        .replace("”", '"')
    )
    corrected = False
    if correct_typos:
        cleaned = text.translate(str.maketrans({"O": "0", "I": "1", "L": "1"}))
        corrected = cleaned != text
        text = cleaned
    directions = [text[0]] if text and text[0] in "NSEW" else []
    if text and text[-1] in "NSEW":
        directions.append(text[-1])
    allowed = "NS" if axis == "Latitude" else "EW"
    if any(d not in allowed for d in directions):
        raise ValueError(f"{axis} has the wrong hemisphere letter")
    core = text.strip("NSEW").strip()
    if not re.fullmatch(r"[+-]?(?:\d+(?:[.,]\d*)?|[.,]\d+)(?:E[+-]?\d+)?", core):
        tokens = re.sub(r"[°º'\"′″:_]", " ", core.replace(",", ".")).split()
        if len(tokens) >= 2 and "." in tokens[0]:
            raise ValueError("Degrees must be whole numbers when minutes are supplied")
        if len(tokens) == 3 and "." in tokens[1]:
            raise ValueError("Minutes must be whole numbers when seconds are supplied")
    number = parse_coordinate(text)
    limit = 90 if axis == "Latitude" else 180
    if not -limit <= number <= limit:
        raise ValueError(f"{axis} must be between {-limit} and {limit}")
    return number, corrected


def convert_coordinates(
    table,
    latitude_column,
    longitude_column,
    name_column=None,
    precision=6,
    include_utm=True,
    correct_typos=False,
    cancel=None,
    progress=lambda value, message: None,
):
    """Convert WGS84 angular input to decimal degrees and optional UTM."""
    import pandas as pd
    from shapely.geometry import Point
    from pyproj import Transformer

    if precision not in (4, 6, 8):
        raise ValueError("Choose 4, 6, or 8 decimal places")
    columns = [latitude_column, longitude_column] + (
        [name_column] if name_column else []
    )
    if len(set(columns)) != len(columns) or any(
        c not in table.columns for c in columns
    ):
        raise ValueError(
            "Map distinct latitude, longitude, and optional point-name columns"
        )
    converted, rejected, corrections, features = [], [], [], []
    transformers = {}
    missing = utm_unavailable = 0

    def blank(value):
        return pd.isna(value) or not str(value).strip()

    for index, values in enumerate(
        table.loc[:, columns].itertuples(index=False, name=None)
    ):
        checkpoint(cancel)
        lat_raw, lon_raw = values[:2]
        name = (
            str(values[2]).strip()
            if name_column and not blank(values[2])
            else f"Point_{index + 1}"
        )
        is_missing = blank(lat_raw) or blank(lon_raw)
        try:
            if is_missing:
                raise ValueError("Missing latitude or longitude")
            lat, lat_fixed = parse_converter_coordinate(
                lat_raw, "Latitude", correct_typos
            )
            lon, lon_fixed = parse_converter_coordinate(
                lon_raw, "Longitude", correct_typos
            )
            record = {
                "Point name": name,
                "Latitude": round(lat, precision),
                "Longitude": round(lon, precision),
                "Input row": index + 2,
            }
            if include_utm:
                record.update(
                    {
                        "UTM Zone": "",
                        "UTM EPSG": "",
                        "UTM Easting (m)": "",
                        "UTM Northing (m)": "",
                        "UTM note": "",
                    }
                )
                if -80 <= lat <= 84:
                    zone = min(60, int((lon + 180) // 6) + 1)
                    if 56 <= lat < 64 and 3 <= lon < 12:
                        zone = 32
                    elif 72 <= lat <= 84 and 0 <= lon < 42:
                        zone = (
                            31
                            if lon < 9
                            else 33
                            if lon < 21
                            else 35
                            if lon < 33
                            else 37
                        )
                    epsg = (32600 if lat >= 0 else 32700) + zone
                    if epsg not in transformers:
                        transformers[epsg] = Transformer.from_crs(
                            4326, epsg, always_xy=True
                        )
                    east, north = transformers[epsg].transform(lon, lat, errcheck=True)
                    if not math.isfinite(east) or not math.isfinite(north):
                        raise ValueError(
                            "UTM projection returned a non-finite coordinate"
                        )
                    record.update(
                        {
                            "UTM Zone": f"{zone}{'N' if lat >= 0 else 'S'}",
                            "UTM EPSG": epsg,
                            "UTM Easting (m)": round(east, 2),
                            "UTM Northing (m)": round(north, 2),
                        }
                    )
                else:
                    record["UTM note"] = (
                        "Outside UTM latitude range (80°S to 84°N); decimal coordinates retained"
                    )
                    utm_unavailable += 1
            converted.append(record)
            features.append(Feature(name, Point(lon, lat), dict(record)))
            if lat_fixed or lon_fixed:
                corrections.append(
                    {
                        "Input row": index + 2,
                        "Point name": name,
                        "Latitude input": str(lat_raw),
                        "Longitude input": str(lon_raw),
                        "Latitude": lat,
                        "Longitude": lon,
                    }
                )
        except (ValueError, TypeError) as exc:
            missing += int(is_missing)
            rejected.append(
                {
                    "Input row": index + 2,
                    "Point name": name,
                    "Latitude input": str(lat_raw),
                    "Longitude input": str(lon_raw),
                    "Reason": str(exc),
                }
            )
        if index % 100 == 0 or index + 1 == len(table):
            progress(
                100 * (index + 1) / max(1, len(table)),
                f"Converting coordinates: {index + 1:,} / {len(table):,}",
            )
    checkpoint(cancel)
    return Result(
        "Coordinate Converter",
        features=features,
        tables={
            "Converted coordinates": converted,
            "Rejected rows": rejected,
            "Corrections": corrections,
        },
        summary={
            "Input rows": len(table),
            "Converted rows": len(converted),
            "Rejected rows": len(rejected),
            "Missing coordinate rows": missing,
            "Corrected rows": len(corrections),
            "UTM unavailable rows": utm_unavailable,
            "Decimal places": precision,
            "Include UTM": include_utm,
            "Typo correction": correct_typos,
            "Source CRS": "WGS84 (EPSG:4326)",
        },
        warnings=["UTM is unavailable for some polar points; see UTM note."]
        if utm_unavailable
        else [],
    )


def nearest_sites(
    queries,
    sites,
    query_columns,
    site_columns,
    count=3,
    ignore_zero_distance=False,
    exclude_same_id=False,
    swap_queries=False,
    cancel=None,
    progress=lambda value, message: None,
):
    """Exact ellipsoidal ranking, with bounded distance batches and stable ties.

    Each reference row is a candidate, including repeated IDs. Coincident
    points can be excluded before ranking using unrounded geodesic distances.
    Distances are ground distances, and bearings are forward from query to site.
    """
    import numpy as np
    import pandas as pd
    from pyproj import Geod

    if isinstance(count, bool) or not isinstance(count, int) or not 1 <= count <= 1000:
        raise ValueError("Nearest N must be a whole number from 1 to 1000.")
    for table, columns in ((queries, query_columns), (sites, site_columns)):
        if (
            len(columns) != 3
            or len(set(columns)) != 3
            or any(c not in table.columns for c in columns)
        ):
            raise ValueError(
                "Map three distinct ID, latitude, and longitude columns in each table."
            )

    def normalize(table, columns, swap=False):
        valid, rejected = [], []
        for position, values in enumerate(
            table.loc[:, list(columns)].itertuples(index=False, name=None)
        ):
            if position % 256 == 0:
                checkpoint(cancel)
            name, latitude, longitude = values
            try:
                if pd.isna(name) or not str(name).strip():
                    raise ValueError("Missing ID")
                lat, lon = parse_coordinate(latitude), parse_coordinate(longitude)
                if swap:
                    lat, lon = lon, lat
                if not (-90 <= lat <= 90 and -180 <= lon <= 180):
                    raise ValueError(
                        "Coordinate outside WGS84 latitude/longitude limits"
                    )
                valid.append((position, str(name).strip(), lat, lon))
            except (ValueError, TypeError) as exc:
                rejected.append(
                    {
                        "Input row": position + 2,
                        "ID": str(name),
                        "Latitude input": str(latitude),
                        "Longitude input": str(longitude),
                        "Reason": str(exc),
                    }
                )
        return valid, rejected

    checkpoint(cancel)
    qrows, bad_queries = normalize(queries, query_columns, swap_queries)
    srows, bad_sites = normalize(sites, site_columns)
    candidates = srows
    lats = np.array([row[2] for row in candidates], dtype=float)
    lons = np.array([row[3] for row in candidates], dtype=float)
    ids = np.array([row[1] for row in candidates], dtype=object)
    geod = Geod(ellps="WGS84")
    matches, unmatched = [], []
    for qi, (position, name, lat, lon) in enumerate(qrows):
        checkpoint(cancel)
        best = []
        for start in range(0, len(candidates), 8192):
            checkpoint(cancel)
            end = min(start + 8192, len(candidates))
            azimuth, _, distances = geod.inv(
                np.full(end - start, lon),
                np.full(end - start, lat),
                lons[start:end],
                lats[start:end],
            )
            indices = np.arange(start, end)
            valid = np.isfinite(distances)
            if ignore_zero_distance:
                valid &= distances > 0.0
            if exclude_same_id:
                valid &= ids[start:end] != name
            local = indices[valid]
            order = np.lexsort((local, distances[valid]))[:count]
            best.extend(
                (float(distances[j - start]), int(j), float(azimuth[j - start]))
                for j in local[order]
            )
            best = sorted(best, key=lambda item: (item[0], item[1]))[:count]
        if not best:
            unmatched.append(
                {
                    "Input row": position + 2,
                    "Query ID": name,
                    "Latitude": lat,
                    "Longitude": lon,
                    "Reason": "No eligible reference sites after validation and the selected ID/zero-distance exclusions",
                }
            )
        for rank, (distance, index, azimuth) in enumerate(best, 1):
            site_row, site_id, site_lat, site_lon = candidates[index]
            matches.append(
                {
                    "Input row": position + 2,
                    "Query ID": name,
                    "Query latitude": lat,
                    "Query longitude": lon,
                    "Rank": rank,
                    "Reference row": site_row + 2,
                    "Site ID": site_id,
                    "Site latitude": site_lat,
                    "Site longitude": site_lon,
                    "Distance (m)": distance,
                    "Distance (km)": distance / 1000,
                    "Bearing (degrees)": azimuth % 360 if distance > 1e-6 else "",
                    "Bearing note": "Coincident: bearing undefined"
                    if distance <= 1e-6
                    else "",
                }
            )
        if qi % 25 == 0 or qi + 1 == len(qrows):
            progress(
                (qi + 1) / max(len(qrows), 1) * 100,
                f"Finding nearest sites: {qi + 1:,} / {len(qrows):,} locations",
            )
    checkpoint(cancel)
    first_columns = ("Query ID", "Rank", "Site ID", "Distance (m)", "Bearing (degrees)")
    matches = [{**{key: row[key] for key in first_columns}, **row} for row in matches]
    return Result(
        "Distance Calculator",
        tables={
            "Distance Calculator": matches,
            "Unmatched queries": unmatched,
            "Rejected queries": bad_queries,
            "Rejected sites": bad_sites,
        },
        summary={
            "Input queries": len(queries),
            "Matched queries": len(qrows) - len(unmatched),
            "Unmatched queries": len(unmatched),
            "Rejected queries": len(bad_queries),
            "Input reference rows": len(sites),
            "Eligible reference rows": len(candidates),
            "Rejected reference rows": len(bad_sites),
            "Requested nearest N": count,
            "Result pairs": len(matches),
            "Ignore 0 distance": ignore_zero_distance,
            "Exclude matching IDs": exclude_same_id,
            "Query coordinates swapped": swap_queries,
            "Distance model": "WGS84 ellipsoid",
            "Bearing": "Clockwise from true north, query to site",
        },
        warnings=[
            "Each reference row is ranked separately; repeated site IDs are not merged. Ties follow reference input order."
        ],
    )


def geometry_leaves(geometry):
    """Yield individual geometries without changing their topology."""
    if geometry is None or geometry.is_empty:
        return
    if hasattr(geometry, "geoms"):
        for child in geometry.geoms:
            yield from geometry_leaves(child)
    else:
        yield geometry


def with_metadata(feature, **metadata):
    properties = dict(feature.properties)
    key = "_geo_toolkit"
    while key in properties:
        key = "_" + key
    properties[key] = metadata
    return properties


def load_cutting_lines(paths, cancel=None, progress=lambda value, text: None):
    from shapely import force_2d, get_coordinates
    import numpy as np

    features, warnings = [], []
    for index, path in enumerate(paths):
        checkpoint(cancel)
        progress(
            index / max(1, len(paths)) * 100,
            f"Reading cutting lines: {Path(path).name}",
        )
        for feature in read_polygon_file(path, warnings, include_lines=True):
            for part in geometry_leaves(feature.geometry):
                checkpoint(cancel)
                if part.geom_type not in {"LineString", "LinearRing"}:
                    warnings.append(f"Skipped non-line geometry: {feature.name}")
                    continue
                line = force_2d(part)
                if (
                    not line.is_valid
                    or not np.isfinite(get_coordinates(line)).all()
                    or line.length == 0
                ):
                    warnings.append(
                        f"Skipped invalid or zero-length line: {feature.name}"
                    )
                    continue
                features.append(
                    Feature(
                        feature.name, line, dict(feature.properties), feature.source
                    )
                )
    if not features:
        raise ValueError(
            "No usable cutting lines found. Choose a file containing LineString geometries."
        )
    validate_extent(features)
    return Layer(features, tuple(paths), warnings)


def split_polygons(layer, cutters, options, cancel, progress):
    from shapely.geometry import box
    from shapely.ops import transform, split
    from shapely import STRtree

    mode = options.split_mode
    if mode not in {"Grid", "Cutting lines"}:
        raise ValueError("Choose Grid or Cutting lines.")
    result = Result("Split polygons", warnings=list(layer.warnings))
    rows = []
    working_crs = "WGS84 longitude/latitude"
    if mode == "Grid":
        nr, nc = options.grid_rows, options.grid_columns
        if (
            any(
                isinstance(n, bool) or not isinstance(n, int) or not 1 <= n <= 1000
                for n in (nr, nc)
            )
            or nr * nc > 10000
        ):
            raise ValueError(
                "Grid rows and columns must be 1–1000, with at most 10,000 cells."
            )
        forward, inverse, working_crs = metric_transformers(layer.features)
        projected = [transform(forward, feature.geometry) for feature in layer.features]
        xmin = min(g.bounds[0] for g in projected)
        ymin = min(g.bounds[1] for g in projected)
        xmax = max(g.bounds[2] for g in projected)
        ymax = max(g.bounds[3] for g in projected)
        dx, dy = (xmax - xmin) / nc, (ymax - ymin) / nr
        if dx <= 0 or dy <= 0:
            raise ValueError("Grid needs a nonzero polygon extent.")
    else:
        if cutters is None or not cutters.features:
            raise ValueError("Load cutting lines first.")
        validate_extent(layer.features + cutters.features)
        lines = [
            part
            for f in cutters.features
            for part in geometry_leaves(f.geometry)
            if part.geom_type in {"LineString", "LinearRing"}
        ]
        if not lines:
            raise ValueError("The cutter layer contains no lines.")
        tree = STRtree(lines)
        result.warnings.extend(cutters.warnings)

    for index, feature in enumerate(layer.features, 1):
        checkpoint(cancel)
        pieces = []
        if mode == "Grid":
            geometry = projected[index - 1]
            left, bottom, right, top = geometry.bounds
            c0, c1 = (
                max(0, int((left - xmin) / dx)),
                min(nc - 1, int((right - xmin) / dx)),
            )
            r0, r1 = (
                max(0, int((bottom - ymin) / dy)),
                min(nr - 1, int((top - ymin) / dy)),
            )
            for row in range(r0, r1 + 1):
                for col in range(c0, c1 + 1):
                    checkpoint(cancel)
                    cell = box(
                        xmin + col * dx,
                        ymin + row * dy,
                        xmax if col == nc - 1 else xmin + (col + 1) * dx,
                        ymax if row == nr - 1 else ymin + (row + 1) * dy,
                    )
                    clipped = geometry.intersection(cell)
                    for part in polygon_parts(clipped):
                        if part.area > 0:
                            pieces.append((transform(inverse, part), row + 1, col + 1))
                            if len(pieces) + len(result.features) > 100000:
                                raise ValueError(
                                    "Split exceeds 100,000 pieces. Use a coarser grid or a smaller input."
                                )
        else:
            current = polygon_parts(feature.geometry)
            for line_index in sorted(
                tree.query(feature.geometry, predicate="intersects").tolist()
            ):
                checkpoint(cancel)
                next_parts = []
                for part in current:
                    checkpoint(cancel)
                    next_parts.extend(polygon_parts(split(part, lines[line_index])))
                    if len(next_parts) + len(result.features) > 100000:
                        raise ValueError(
                            "Split exceeds 100,000 pieces. Use fewer cutters or a smaller input."
                        )
                current = next_parts
            pieces = [(part, "", "") for part in current]
        if len(result.features) + len(pieces) > 100000:
            raise ValueError(
                "Split exceeds 100,000 pieces. Use a coarser grid or a smaller input."
            )
        for piece_number, (geometry, row, col) in enumerate(pieces, 1):
            checkpoint(cancel)
            name = f"{feature.name} / piece {piece_number}"
            result.features.append(
                Feature(
                    name,
                    geometry,
                    with_metadata(
                        feature,
                        source_feature=index,
                        piece=piece_number,
                        grid_row=row,
                        grid_column=col,
                    ),
                    feature.source,
                )
            )
            rows.append(
                {
                    "Input feature": index,
                    "Name": feature.name,
                    "Piece": piece_number,
                    "Grid row": row,
                    "Grid column": col,
                    "Area (m²)": geodesic_metrics(geometry)[0],
                    "Source": feature.source,
                }
            )
        progress(
            index / len(layer.features) * 100,
            f"Splitting polygon {index:,} / {len(layer.features):,}",
        )
    checkpoint(cancel)
    result.tables["Pieces"] = rows
    result.summary = {
        "Input A features": len(layer.features),
        "Result features": len(result.features),
        "Mode": mode,
        "Working CRS": working_crs,
        "Grid rows": options.grid_rows if mode == "Grid" else "",
        "Grid columns": options.grid_columns if mode == "Grid" else "",
    }
    result.warnings.append(
        "Grid rows are numbered south to north, columns west to east. Cutting lines must cross an interior to divide it; lines that only touch a boundary leave it unchanged."
        if mode == "Grid"
        else "Cutters are used as supplied, without snapping or extension. Lines that stop inside or only touch a boundary may not divide a polygon."
    )
    return result


def fill_holes(layer, options, cancel, progress):
    from shapely.geometry import Polygon
    from shapely import union_all

    limit = options.hole_area_limit
    if not options.fill_all_holes and (not math.isfinite(limit) or limit < 0):
        raise ValueError(
            "Hole area limit must be a finite, nonnegative number in square meters."
        )
    result = Result("Fill holes", warnings=list(layer.warnings))
    audit, changes = [], []
    filled = 0
    for index, feature in enumerate(layer.features, 1):
        checkpoint(cancel)
        parts = []
        feature_filled = 0
        for part_number, part in enumerate(polygon_parts(feature.geometry), 1):
            retained = []
            for ring_number, ring in enumerate(part.interiors, 1):
                checkpoint(cancel)
                area = geodesic_metrics(Polygon(ring))[0]
                remove = options.fill_all_holes or area <= limit
                if remove:
                    filled += 1
                    feature_filled += 1
                else:
                    retained.append(ring.coords)
                audit.append(
                    {
                        "Input feature": index,
                        "Name": feature.name,
                        "Part": part_number,
                        "Hole": ring_number,
                        "Hole area (m²)": area,
                        "Action": "Filled" if remove else "Retained",
                    }
                )
            parts.append(Polygon(part.exterior.coords, retained))
        # Islands within a newly filled hole can overlap; dissolve within this
        # source feature only, preserving a valid polygon result.
        changed = union_all(parts)
        result.features.append(
            Feature(feature.name, changed, dict(feature.properties), feature.source)
        )
        before, after = (
            geodesic_metrics(feature.geometry)[0],
            geodesic_metrics(changed)[0],
        )
        changes.append(
            {
                "Input feature": index,
                "Name": feature.name,
                "Holes filled": feature_filled,
                "Area before (m²)": before,
                "Area after (m²)": after,
                "Area gained (m²)": after - before,
            }
        )
        progress(index / len(layer.features) * 100, f"Filling holes: polygon {index:,}")
    checkpoint(cancel)
    result.tables = {"Changes": changes, "Holes": audit}
    result.summary = {
        "Input A features": len(layer.features),
        "Result features": len(result.features),
        "Holes inspected": len(audit),
        "Holes filled": filled,
        "Mode": "All holes" if options.fill_all_holes else "Area limit",
        "Maximum hole area (m²)": "All" if options.fill_all_holes else limit,
    }
    return result


def validation_report(
    layer, repair=False, cancel=None, progress=lambda value, text: None
):
    from shapely import force_2d, get_coordinates, make_valid, union_all
    from shapely.validation import explain_validity
    import numpy as np

    originals = layer.raw_features if layer.raw_features is not None else layer.features
    result = Result("Geometry validation", warnings=list(layer.warnings))
    records = []
    counts = Counter()
    for index, feature in enumerate(originals, 1):
        checkpoint(cancel)
        original = feature.geometry
        record = {
            "Input feature": index,
            "Name": feature.name,
            "Status": "",
            "Reason": "",
            "Source": feature.source,
            "Original type": original.geom_type if original is not None else "Missing",
            "Original parts": len(polygon_parts(original)),
            "Original area (m²)": "",
            "Repair action": "Not requested",
            "Output type": "",
            "Output parts": 0,
            "Output area (m²)": "",
            "Discarded non-polygon pieces": 0,
            "Had Z coordinates": bool(original is not None and original.has_z),
        }
        valid = None
        try:
            if original is None:
                record.update(Status="Missing", Reason="Null geometry")
            elif original.is_empty:
                record.update(Status="Empty", Reason="Empty geometry")
            else:
                coordinates = get_coordinates(original)
                if not np.isfinite(coordinates).all():
                    record.update(
                        Status="Invalid coordinates", Reason="Non-finite coordinate"
                    )
                elif np.any(np.abs(coordinates[:, 0]) > 180) or np.any(
                    np.abs(coordinates[:, 1]) > 90
                ):
                    record.update(
                        Status="Invalid coordinates",
                        Reason="Outside WGS84 longitude/latitude limits; define or reproject the source CRS",
                    )
                elif original.bounds[2] - original.bounds[0] > 180:
                    record.update(
                        Status="Unsupported extent",
                        Reason="Date-line or greater-than-hemisphere shape requires preprocessing",
                    )
                elif original.geom_type not in {"Polygon", "MultiPolygon"}:
                    record.update(
                        Status="Non-polygon",
                        Reason="A polygon or multipolygon is required",
                    )
                else:
                    geometry = force_2d(original)
                    record["Status"] = "Valid" if geometry.is_valid else "Invalid"
                    record["Reason"] = explain_validity(geometry)
                    if geometry.is_valid:
                        valid = geometry
                        record["Original area (m²)"] = geodesic_metrics(geometry)[0]
                        if repair:
                            record["Repair action"] = "Retained valid polygon (2D)"
                    elif repair:
                        fixed = make_valid(geometry)
                        leaves = list(geometry_leaves(fixed))
                        record["Discarded non-polygon pieces"] = sum(
                            part.geom_type != "Polygon" for part in leaves
                        )
                        parts = polygon_parts(fixed)
                        if parts:
                            valid = union_all(parts)
                        if valid is None or valid.is_empty or not valid.is_valid:
                            valid = None
                            record["Repair action"] = "No valid polygon area recovered"
                        else:
                            record["Repair action"] = "Repaired polygon copy"
                if (
                    repair
                    and valid is None
                    and record["Repair action"] == "Not requested"
                ):
                    record["Repair action"] = "Excluded from polygon output"
            if repair and valid is None and record["Repair action"] == "Not requested":
                record["Repair action"] = "Excluded from polygon output"
            if repair and valid is not None:
                record["Output type"] = valid.geom_type
                record["Output parts"] = len(polygon_parts(valid))
                record["Output area (m²)"] = geodesic_metrics(valid)[0]
                result.features.append(
                    Feature(
                        feature.name, valid, dict(feature.properties), feature.source
                    )
                )
        except Exception as exc:
            record.update(Status="Error", Reason=str(exc))
            if repair:
                record["Repair action"] = "Failed; excluded from output"
        counts[record["Status"]] += 1
        records.append(record)
        progress(
            index / max(1, len(originals)) * 100,
            f"Validating original feature {index:,}",
        )
    checkpoint(cancel)
    result.tables = {"Validation": records}
    result.summary = {
        "Input A features": len(originals),
        "Valid original features": counts["Valid"],
        "Features needing review": len(originals) - counts["Valid"],
        "Repair requested": repair,
        "Result features": len(result.features),
        "Basis": "Original parsed input, before automatic cleanup"
        if layer.raw_features is not None
        else "Current layer geometry",
    }
    result.warnings.append(
        "Validity is checked in 2D. Invalid polygon areas are not reported as reliable measurements. Repair copies retain valid polygons, repair recoverable polygons, and omit unrecoverable/non-polygon records; inspect every change before using them. Parser notices identify records that could not be constructed."
    )
    return result


def boundary_vertices(layer, options, cancel, progress):
    from shapely.geometry import Point, LineString
    from pyproj import Geod

    if options.boundary_output not in {"Boundaries", "Vertices"}:
        raise ValueError("Choose Boundaries or Vertices.")
    vertices = options.boundary_output == "Vertices"
    result = Result("Boundary / vertices", warnings=list(layer.warnings))
    records = []
    geod = Geod(ellps="WGS84")
    ring_count = 0
    for index, feature in enumerate(layer.features, 1):
        for part_number, part in enumerate(polygon_parts(feature.geometry), 1):
            rings = [("Exterior", 0, part.exterior)]
            if options.include_holes:
                rings.extend(
                    ("Hole", i, ring) for i, ring in enumerate(part.interiors, 1)
                )
            for kind, number, ring in rings:
                checkpoint(cancel)
                ring_count += 1
                base = {
                    "Input feature": index,
                    "Name": feature.name,
                    "Part": part_number,
                    "Ring type": kind,
                    "Ring": number,
                }
                if vertices:
                    coords = list(ring.coords)
                    closing = len(coords) - 1
                    if not options.include_closing_vertex:
                        coords = coords[:-1]
                    for vertex, (lon, lat) in enumerate(coords, 1):
                        if vertex % 256 == 1:
                            checkpoint(cancel)
                        record = {
                            **base,
                            "Vertex": vertex,
                            "Longitude": lon,
                            "Latitude": lat,
                            "Closing vertex": vertex - 1 == closing,
                            "Source": feature.source,
                        }
                        records.append(record)
                        result.features.append(
                            Feature(
                                f"{feature.name} / part {part_number} / {kind} {number} / vertex {vertex}",
                                Point(lon, lat),
                                with_metadata(feature, **record),
                                feature.source,
                            )
                        )
                else:
                    line = LineString(ring.coords)
                    record = {
                        **base,
                        "Length (m)": geod.geometry_length(line),
                        "Vertices excluding closure": len(ring.coords) - 1,
                        "Source": feature.source,
                    }
                    records.append(record)
                    result.features.append(
                        Feature(
                            f"{feature.name} / part {part_number} / {kind} {number}",
                            line,
                            with_metadata(feature, **record),
                            feature.source,
                        )
                    )
        progress(
            index / len(layer.features) * 100,
            f"Extracting {options.boundary_output.lower()}: polygon {index:,}",
        )
    checkpoint(cancel)
    result.tables["Vertices" if vertices else "Boundaries"] = records
    result.summary = {
        "Input A features": len(layer.features),
        "Result features": len(result.features),
        "Output": options.boundary_output,
        "Rings": ring_count,
        "Include holes": options.include_holes,
        "Include closing vertex": options.include_closing_vertex
        if vertices
        else "Always closed",
        "Coordinates": "WGS84, 2D",
    }
    return result


def run_operation(
    a: Layer,
    b: Layer | None,
    options: Options,
    table=None,
    cancel=None,
    progress: Progress = lambda value, text: None,
) -> Result:
    from shapely.ops import transform
    from shapely import union_all

    if options.operation not in OPERATIONS:
        raise ValueError("Unknown operation")
    checkpoint(cancel)
    if options.operation == "Geometry validation":
        return validation_report(a, options.repair_geometry, cancel, progress)
    if not a.features:
        raise ValueError(
            "No usable polygon area. Run Geometry validation to inspect the original inputs."
        )
    validate_extent(a.features)
    op = options.operation
    if op == "Split polygons":
        return split_polygons(a, b, options, cancel, progress)
    if op == "Fill holes":
        return fill_holes(a, options, cancel, progress)
    if op == "Boundary / vertices":
        return boundary_vertices(a, options, cancel, progress)
    binary = op in {"Intersection", "Difference", "SymDifference"}
    if binary and (b is None or not b.features):
        raise ValueError(f"{op} requires Layer B.")
    if binary or (op == "Combine" and b is not None):
        validate_extent(a.features + (b.features if b else []))
    if op == "Points inside polygon":
        if table is None:
            raise ValueError("Choose a point table first.")
        return match_points(a, table, options, cancel, progress)
    result = Result(op, warnings=list(a.warnings) + (list(b.warnings) if b else []))
    progress(15, f"Calculating {op}…")
    if op in {"Combine", "Intersection", "Difference", "SymDifference"}:
        left = union_all([f.geometry for f in a.features])
        checkpoint(cancel)
        right = union_all([f.geometry for f in b.features]) if b else None
        if op == "Combine":
            geometry = union_all([left, right]) if right is not None else left
        elif op == "Intersection":
            geometry = left.intersection(right)
        elif op == "Difference":
            geometry = left.difference(right)
        else:
            geometry = left.symmetric_difference(right)
        geometry, _ = clean_polygon(geometry)
        if geometry is not None and not geometry.is_empty:
            result.features = [Feature(op, geometry, {"operation": op}, "Calculated")]
        result.summary["Attribute policy"] = (
            "Merged footprint; source attribute tables are not overlaid"
        )
    elif op in {"Buffer", "Simplify"}:
        distance = options.distance * (1000 if options.unit == "Kilometers" else 1)
        if not math.isfinite(distance) or (op == "Simplify" and distance <= 0):
            raise ValueError(
                "Enter a finite distance; Simplify requires a positive tolerance."
            )
        forward, inverse, crs = metric_transformers(a.features)
        for index, feature in enumerate(a.features):
            checkpoint(cancel)
            projected = transform(forward, feature.geometry)
            changed = (
                projected.buffer(distance)
                if op == "Buffer"
                else projected.simplify(distance, preserve_topology=True)
            )
            if not changed.is_empty:
                geometry, _ = clean_polygon(transform(inverse, changed))
                if geometry is not None:
                    result.features.append(
                        Feature(
                            feature.name,
                            geometry,
                            dict(feature.properties),
                            feature.source,
                        )
                    )
            progress(
                15 + 75 * (index + 1) / len(a.features),
                f"Processing feature {index + 1:,}",
            )
        result.summary.update(
            {
                "Distance (m)": distance,
                "Working CRS": crs,
                "Removed empty features": len(a.features) - len(result.features),
            }
        )
    else:
        measured = []
        for index, feature in enumerate(a.features):
            checkpoint(cancel)
            area, perimeter = geodesic_metrics(feature.geometry)
            measured.append(
                {
                    "Feature": index + 1,
                    "Name": feature.name,
                    "Area (m²)": area,
                    "Area (km²)": area / 1e6,
                    "Area (acres)": area / 4046.8564224,
                    "Perimeter (m)": perimeter,
                    "Perimeter (km)": perimeter / 1000,
                    "Perimeter (miles)": perimeter / 1609.344,
                }
            )
        checkpoint(cancel)
        if options.measurement_scope == "Merged coverage":
            area, perimeter = geodesic_metrics(
                union_all([f.geometry for f in a.features])
            )
        else:
            area = sum(row["Area (m²)"] for row in measured)
            perimeter = sum(row["Perimeter (m)"] for row in measured)
        keys = (
            ("Feature", "Name", "Area (m²)", "Area (km²)", "Area (acres)")
            if op == "Area"
            else (
                "Feature",
                "Name",
                "Perimeter (m)",
                "Perimeter (km)",
                "Perimeter (miles)",
            )
        )
        result.tables["Measurements"] = [
            {key: row[key] for key in keys} for row in measured
        ]
        result.summary.update(
            {
                "Scope": options.measurement_scope,
                "Model": "WGS84 ellipsoid",
                "Total area (m²)" if op == "Area" else "Total perimeter (m)": area
                if op == "Area"
                else perimeter,
            }
        )
    checkpoint(cancel)
    result.summary.update({"Operation": op, "Input A features": len(a.features)})
    if op not in {"Area", "Perimeter"}:
        result.summary["Result features"] = len(result.features)
        result.tables["Features"] = [
            {
                "Feature": i,
                "Name": f.name,
                "Geometry": f.geometry.geom_type,
                "Source": f.source,
            }
            for i, f in enumerate(result.features, 1)
        ]
        if not result.features:
            result.warnings.append("The operation produced no polygon area.")
    progress(100, "Complete")
    return result


# ==========================================================================
# FILEIO
# ==========================================================================

MAX_MEMBER_BYTES = 128 * 1024 * 1024
MAX_ARCHIVE_BYTES = 512 * 1024 * 1024
KML_NS = "http://www.opengis.net/kml/2.2"


def decode_text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def choose_name(properties: dict, fallback: str) -> str:
    normalized = {
        re.sub(r"[\s_-]+", "", str(k)).lower(): v for k, v in properties.items()
    }
    for key in (
        "name",
        "polygonname",
        "sitename",
        "siteid",
        "site",
        "cellname",
        "cellid",
        "cell",
        "sectorname",
        "sectorid",
        "sector",
        "id",
    ):
        value = normalized.get(key)
        if (
            value is not None
            and str(value).strip()
            and str(value).strip().lower() not in {"nan", "none"}
        ):
            return str(value).strip()
    return fallback


def parse_kml(
    data: bytes, source: str, warnings: list[str], include_lines=False
) -> list[Feature]:
    from shapely.geometry import Polygon, LineString
    from defusedxml import ElementTree as SafeET

    root = SafeET.fromstring(data)
    for element in root.iter():
        element.tag = element.tag.rsplit("}", 1)[-1]
    features = []
    for number, placemark in enumerate(root.iter("Placemark"), 1):
        properties = {}
        for item in placemark.iter("SimpleData"):
            properties[item.get("name", "value")] = item.text or ""
        for item in placemark.iter("Data"):
            value = item.find("value")
            if value is None:
                value = item.find("Value")
            if value is not None:
                properties[item.get("name", "value")] = value.text or ""
        name = next(
            (
                value.strip()
                for tag in ("name", "Name")
                if (value := placemark.findtext(tag)) and value.strip()
            ),
            "",
        )
        name = name or choose_name(properties, "")
        if not name:
            # Some exporters prefix schema fields, e.g. network_sector_name.
            for key, value in properties.items():
                tokens = re.sub(r"([a-z])([A-Z])", r"\1_\2", key)
                tokens = re.findall(r"[a-z0-9]+", tokens.lower())
                if set(tokens) & {"name", "site", "cell", "sector", "id"}:
                    candidate = choose_name({"name": value}, "")
                    if candidate:
                        name = candidate
                        break
        if not name:
            name = next(
                (
                    str(value).strip()
                    for value in properties.values()
                    if choose_name({"name": value}, "")
                ),
                "",
            )
            if name:
                warnings.append(
                    f"{source}: placemark {number} has no recognized name field; used its first non-empty metadata value ({name})."
                )
        name = name or f"{Path(source).stem} / {number}"
        if include_lines:
            for element in placemark.iter("LineString"):
                try:
                    coordinates = [
                        tuple(float(v) for v in token.split(",")[:2])
                        for token in (element.findtext("coordinates") or "").split()
                    ]
                    features.append(
                        Feature(name, LineString(coordinates), properties, source)
                    )
                except (ValueError, TypeError) as exc:
                    warnings.append(f"{source}: skipped line {name}: {exc}")
        for element in placemark.iter("Polygon"):
            try:

                def ring(path):
                    text = element.findtext(path)
                    if not text:
                        raise ValueError("Missing ring coordinates")
                    coordinates = []
                    for token in text.split():
                        fields = token.split(",")
                        coordinates.append((float(fields[0]), float(fields[1])))
                    if len(set(coordinates)) < 3 or not all(
                        math.isfinite(v) for p in coordinates for v in p
                    ):
                        raise ValueError(
                            "Ring requires three distinct finite coordinates"
                        )
                    return coordinates

                exterior = ring("outerBoundaryIs/LinearRing/coordinates")
                holes = []
                for boundary in element.findall("innerBoundaryIs"):
                    text = boundary.findtext("LinearRing/coordinates")
                    if not text:
                        raise ValueError("Missing hole coordinates")
                    coordinates = [
                        tuple(float(v) for v in token.split(",")[:2])
                        for token in text.split()
                    ]
                    if len(set(coordinates)) < 3 or any(
                        len(p) != 2 or not all(math.isfinite(v) for v in p)
                        for p in coordinates
                    ):
                        raise ValueError("Invalid hole coordinates")
                    holes.append(coordinates)
                features.append(
                    Feature(name, Polygon(exterior, holes), properties, source)
                )
            except (ValueError, IndexError, TypeError) as exc:
                warnings.append(f"{source}: skipped {name}: {exc}")
    return features


def read_archive(
    data: bytes,
    source: str,
    warnings: list[str],
    budget=None,
    depth=0,
    include_lines=False,
) -> list[Feature]:
    if depth > 1:
        raise ValueError("Nested ZIP archives are not supported")
    budget = [0] if budget is None else budget
    features = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for member in archive.infolist():
            suffix = Path(member.filename).suffix.lower()
            if member.is_dir() or suffix not in {".kml", ".kmz"}:
                continue
            if (
                member.file_size > MAX_MEMBER_BYTES
                or budget[0] + member.file_size > MAX_ARCHIVE_BYTES
            ):
                raise ValueError(
                    "Archive exceeds the supported uncompressed size limit"
                )
            budget[0] += member.file_size
            content = archive.read(member)
            member_source = f"{source}/{member.filename}"
            if suffix == ".kmz":
                features.extend(
                    read_archive(
                        content,
                        member_source,
                        warnings,
                        budget,
                        depth + 1,
                        include_lines,
                    )
                )
            else:
                features.extend(
                    parse_kml(content, member_source, warnings, include_lines)
                )
    return features


def read_polygon_file(
    path: str, warnings: list[str], include_lines=False
) -> list[Feature]:
    from shapely.geometry import shape

    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".kml":
        return parse_kml(source.read_bytes(), source.name, warnings, include_lines)
    if suffix in {".kmz", ".zip"}:
        return read_archive(
            source.read_bytes(), source.name, warnings, include_lines=include_lines
        )
    if suffix in {".json", ".geojson"}:
        document = json.loads(source.read_text(encoding="utf-8-sig"))
        if "crs" in document:
            crs_name = document["crs"].get("properties", {}).get("name", "")
            if crs_name not in {
                "EPSG:4326",
                "urn:ogc:def:crs:OGC:1.3:CRS84",
                "urn:ogc:def:crs:EPSG::4326",
            }:
                raise ValueError(
                    "GeoJSON must use WGS84 longitude/latitude. Reproject this file first."
                )
        records = (
            document.get("features", [])
            if document.get("type") == "FeatureCollection"
            else [document]
        )
        features = []
        for index, record in enumerate(records, 1):
            properties = record.get("properties") or {}
            geometry = (
                record.get("geometry") if record.get("type") == "Feature" else record
            )
            if geometry:
                features.append(
                    Feature(
                        choose_name(properties, f"{source.stem} / {index}"),
                        shape(geometry),
                        properties,
                        source.name,
                    )
                )
            else:
                features.append(
                    Feature(
                        choose_name(properties, f"{source.stem} / {index}"),
                        None,
                        properties,
                        source.name,
                    )
                )
        return features
    if suffix in {".shp", ".gpkg"}:
        try:
            import geopandas as gpd
        except ImportError as exc:
            raise ValueError(
                "Shapefile/GeoPackage input needs the optional GIS dependencies. See requirements-gis.txt."
            ) from exc
        frames = []
        if suffix == ".gpkg":
            import pyogrio

            for name, geometry_type in pyogrio.list_layers(path):
                if geometry_type is not None:
                    frames.append((name, gpd.read_file(path, layer=name)))
        else:
            frames.append((source.stem, gpd.read_file(path)))
        features = []
        for layer_name, frame in frames:
            if frame.crs is None:
                raise ValueError(
                    f"{source.name}: missing CRS. Define the correct CRS before importing."
                )
            frame = frame.to_crs(4326)
            for index, (_, row) in enumerate(frame.iterrows(), 1):
                properties = row.drop(labels=[frame.geometry.name]).to_dict()
                features.append(
                    Feature(
                        choose_name(properties, f"{layer_name} / {index}"),
                        row.geometry,
                        properties,
                        f"{source.name}/{layer_name}",
                    )
                )
        return features
    raise ValueError(f"Unsupported polygon format: {suffix}")


def load_layer(paths, cancel=None, progress=lambda value, text: None) -> Layer:
    from shapely import normalize

    features, warnings, seen = [], [], set()
    original = []
    for index, path in enumerate(paths):
        checkpoint(cancel)
        progress(index / max(1, len(paths)) * 100, f"Reading {Path(path).name}")
        # Failed files abort the load. The UI keeps the previous layer intact.
        incoming = read_polygon_file(path, warnings)
        original.extend(incoming)
        for feature in incoming:
            checkpoint(cancel)
            try:
                geometry, repaired = (
                    clean_polygon(feature.geometry)
                    if feature.geometry is not None
                    else (None, False)
                )
            except Exception as exc:
                warnings.append(
                    f"{feature.source}: could not prepare {feature.name}: {exc}; inspect Geometry validation"
                )
                continue
            if geometry is None:
                warnings.append(
                    f"{feature.source}: skipped non-polygon or collapsed feature {feature.name}"
                )
                continue
            if repaired:
                warnings.append(
                    f"{feature.source}: repaired invalid geometry for {feature.name}; review its shape"
                )
            # Preserve semantic differences: names, properties, and interior rings.
            key = (
                feature.name,
                normalize(geometry).wkb,
                json.dumps(feature.properties, sort_keys=True, default=str),
            )
            if key in seen:
                warnings.append(
                    f"{feature.source}: removed an identical duplicate of {feature.name}"
                )
                continue
            seen.add(key)
            features.append(
                Feature(feature.name, geometry, feature.properties, feature.source)
            )
    if not original and not warnings:
        raise ValueError("No geometries found in the selected files.")
    if features:
        try:
            validate_extent(features)
        except ValueError as exc:
            warnings.append(
                str(exc) + " Inspect Geometry validation before processing."
            )
    else:
        warnings.append(
            "No usable polygons; the original records remain available to Geometry validation."
        )
    return Layer(features, tuple(paths), warnings, original)


def read_point_table(path: str):
    import pandas as pd

    if Path(path).suffix.lower() in {".csv", ".txt", ".tsv"}:
        text = decode_text(Path(path).read_bytes())
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
        table = pd.read_csv(
            io.StringIO(text), sep=delimiter, dtype=str, keep_default_na=False
        )
    elif Path(path).suffix.lower() == ".xlsx":
        table = pd.read_excel(path, dtype=str, keep_default_na=False, engine="openpyxl")
    else:
        raise ValueError("Point tables must be CSV, TXT, TSV, or XLSX.")
    table.columns = [str(column) for column in table.columns]
    if not table.columns.is_unique:
        raise ValueError("Point table headers must be unique.")
    return table


def suggest_columns(columns) -> tuple[str, str, str]:
    lookup = {re.sub(r"[\s_-]+", "", str(c)).lower(): c for c in columns}
    aliases = (
        ("pointname", "siteid", "sitename", "name", "site", "id"),
        ("latitude", "lat", "y", "lat.", "latitudedeg"),
        ("longitude", "lon", "long", "lng", "long2", "x", "lon.", "longitudedeg"),
    )
    return tuple(
        next((lookup[key] for key in group if key in lookup), "") for group in aliases
    )


def json_value(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, dict):
        return {str(k): json_value(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_value(v) for v in value]
    return str(value)


def kml_bytes(features: list[Feature]) -> bytes:
    ET.register_namespace("", KML_NS)

    def tag(name):
        return f"{{{KML_NS}}}{name}"

    root = ET.Element(tag("kml"))
    document = ET.SubElement(root, tag("Document"))
    for feature in features:
        placemark = ET.SubElement(document, tag("Placemark"))
        ET.SubElement(placemark, tag("name")).text = feature.name
        extended = ET.SubElement(placemark, tag("ExtendedData"))
        for key, value in feature.properties.items():
            node = ET.SubElement(extended, tag("Data"), name=str(key))
            ET.SubElement(node, tag("value")).text = str(json_value(value))
        parts = list(geometry_leaves(feature.geometry))
        parent = (
            ET.SubElement(placemark, tag("MultiGeometry"))
            if len(parts) > 1
            else placemark
        )
        for part in parts:
            if part.geom_type in {"Point", "LineString", "LinearRing"}:
                node = ET.SubElement(
                    parent, tag("Point" if part.geom_type == "Point" else "LineString")
                )
                ET.SubElement(node, tag("coordinates")).text = " ".join(
                    f"{coordinate[0]:.15g},{coordinate[1]:.15g}"
                    for coordinate in part.coords
                )
                continue
            if part.geom_type != "Polygon":
                raise ValueError(f"Unsupported KML geometry: {part.geom_type}")
            polygon = ET.SubElement(parent, tag("Polygon"))
            for boundary_name, ring in [("outerBoundaryIs", part.exterior)] + [
                ("innerBoundaryIs", hole) for hole in part.interiors
            ]:
                boundary = ET.SubElement(polygon, tag(boundary_name))
                linear = ET.SubElement(boundary, tag("LinearRing"))
                ET.SubElement(linear, tag("coordinates")).text = " ".join(
                    f"{x:.15g},{y:.15g}" for x, y in ring.coords
                )
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def match_fill(count: int) -> str | None:
    """V5 report colors, based on actual matches rather than name punctuation."""
    return (
        "FFEBEE"
        if count == 0
        else "FFF9C4"
        if count == 2
        else "FFE0B2"
        if count >= 3
        else None
    )


def write_workbook(result: Result, path: str, cancel=None) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    tables = {
        "Summary": [{"Item": k, "Value": v} for k, v in result.summary.items()],
        **result.tables,
    }
    if result.warnings:
        tables["Warnings"] = [{"Message": value} for value in result.warnings]
    for title, records in tables.items():
        checkpoint(cancel)
        sheet = workbook.create_sheet(title[:31])
        headers = list(records[0]) if records else ["No records"]
        if len(records) > 1_048_575:
            raise ValueError(
                f"{title} exceeds Excel's row limit. Split the input and run again."
            )
        sheet.append(headers)
        for index, record in enumerate(records):
            if index % 1000 == 0:
                checkpoint(cancel)
            sheet.append([json_value(record.get(header)) for header in headers])
            cells = tuple(
                sheet.cell(row=index + 2, column=column)
                for column in range(1, len(headers) + 1)
            )
            for cell in cells:
                cell.font = Font(name="Segoe UI", size=10)
                cell.alignment = Alignment(vertical="center")
                # User strings must remain text, including names starting with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if isinstance(cell.value, float):
                    cell.number_format = (
                        "0." + "0" * result.summary.get("Decimal places", 6)
                        if cell.column <= len(headers)
                        and headers[cell.column - 1] in {"Latitude", "Longitude"}
                        else "#,##0.0000"
                    )
                    if result.operation == "Coordinate Converter" and headers[
                        cell.column - 1
                    ].startswith("UTM "):
                        cell.number_format = "0.00"
            if title == "Matches":
                count = record.get("Match count", 0)
                color = match_fill(count)
                if color:
                    for cell in cells:
                        cell.fill = PatternFill("solid", fgColor=color)
        for cell in sheet[1]:
            cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2E5090")
            cell.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 25
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        from openpyxl.utils import get_column_letter

        for number, header in enumerate(headers, 1):
            width = max(
                [len(str(header))]
                + [len(str(row.get(header, ""))) for row in records[:200]]
            )
            sheet.column_dimensions[get_column_letter(number)].width = min(
                55, max(16, width + 2)
            )
    checkpoint(cancel)
    workbook.save(path)


def export_result(result: Result, destination: str, cancel=None) -> None:
    """Replace the destination only after a successful write in the same folder."""
    from shapely.geometry import mapping

    target = Path(destination)
    suffix = target.suffix.lower()
    if suffix not in {".xlsx", ".geojson", ".kml", ".kmz"}:
        raise ValueError("Choose XLSX, GeoJSON, KML, or KMZ.")
    if suffix != ".xlsx" and not result.features:
        raise ValueError(
            "This result has no geometry. Export an XLSX workbook instead."
        )
    fd, temporary = tempfile.mkstemp(
        prefix=".geotool-", suffix=suffix, dir=target.parent
    )
    os.close(fd)
    try:
        checkpoint(cancel)
        if suffix == ".xlsx":
            write_workbook(result, temporary, cancel)
        elif suffix == ".geojson":
            document = {
                "type": "FeatureCollection",
                "features": [
                    {
                        "type": "Feature",
                        "geometry": mapping(feature.geometry),
                        "properties": {
                            "name": feature.name,
                            "source": feature.source,
                            **json_value(feature.properties),
                        },
                    }
                    for feature in result.features
                ],
            }
            Path(temporary).write_text(
                json.dumps(document, ensure_ascii=False, allow_nan=False),
                encoding="utf-8",
            )
        elif suffix == ".kml":
            Path(temporary).write_bytes(kml_bytes(result.features))
        else:
            with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("doc.kml", kml_bytes(result.features))
        checkpoint(cancel)
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


# ==========================================================================
# UI
# ==========================================================================

APP_NAME = "Geo Operations Toolkit"
DATA_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "GeoOperationsToolkit"
POLYGON_TYPES = [
    ("Polygon files", "*.kml *.kmz *.zip *.geojson *.json *.shp *.gpkg"),
    ("All files", "*.*"),
]


GUIDE_PAGES = {
    "Coordinate Converter": (
        (
            "1 · Choose input",
            "Select Single point to enter a name (optional), latitude, and longitude, or Bulk table to load CSV, TXT, TSV, or the first XLSX worksheet. For bulk input, map latitude and longitude; a point-name column is optional. Missing names become Point_1, Point_2, etc. Inputs must represent WGS84 angular coordinates; this operation does not read UTM coordinates or convert arbitrary datums.",
        ),
        (
            "2 · Set conversion options",
            "DD, degrees and decimal minutes (DDM), and degrees/minutes/seconds (DMS) are recognized automatically. Use N/S for latitude and E/W for longitude. Choose 4, 6, or 8 decimal places. Include UTM adds zone/hemisphere, EPSG code, easting, and northing. N/S denotes hemisphere, not an MGRS latitude band. UTM uses Norway/Svalbard zone exceptions; polar points retain decimal coordinates with a UTM note.",
        ),
        (
            "3 · Review corrections",
            "Correct O/I/L typos is optional and off by default. It replaces those letters with 0/1/1; review the Corrections table before using the result. Invalid minutes, seconds, hemisphere letters, bounds, and missing coordinates are reported in Rejected rows. All valid rows remain, including duplicate names.",
        ),
        (
            "4 · Run and export",
            "Click Run operation. Inspect Converted coordinates, Summary, Rejected rows, and Corrections. Export XLSX for the complete report, or GeoJSON/KML/KMZ for valid point geometry. Angular report values use the selected precision; UTM values show two decimals. GIS geometry retains the original parsed precision. Use Export again for additional formats. View on Google Maps opens the last result when it contains exactly one point; the coordinate is sent to Google only when you click.",
        ),
    ),
    "Split polygons": (
        (
            "Grid mode",
            "Load polygon Layer A. Select Split polygons and Grid, then choose rows and columns. One grid spans the combined Layer A extent in a local UTM coordinate system. Rows start at the south and columns at the west. At most 10,000 grid cells and 100,000 output pieces are supported. The input must fit the local projection limits used by Buffer and Simplify. Pieces are not guaranteed equal area.",
        ),
        (
            "Cutting lines",
            "Choose Cutting lines, then load LineString data from GeoJSON, KML/KMZ/ZIP, Shapefile, or GeoPackage. Optional GIS packages are needed for the last two formats. Cutters are used without snapping or extension. A line must divide the polygon interior to produce additional pieces; touching boundaries or stopping inside may leave it unchanged.",
        ),
        (
            "Review and export",
            "Run, inspect the Pieces table, and export polygon geometry or an Excel report. Each piece retains source attributes. A collision-safe _geo_toolkit metadata field records its source feature, piece number, and grid cell where applicable. Source files are not modified.",
        ),
    ),
    "Fill holes": (
        (
            "Choose the scope",
            "Load Layer A and select Fill holes. Leave Fill all holes checked to remove every interior ring. Otherwise enter a finite, nonnegative maximum area in square meters; holes at or below that limit are filled. Areas are measured geodesically on WGS84.",
        ),
        (
            "Review changes",
            "The Changes table shows per-feature area before and after; Holes lists each ring, its area, and whether it was filled. An island inside a filled hole is merged within the same source feature to maintain valid output. Separate source features remain separate. Export the changed polygons or Excel report.",
        ),
    ),
    "Geometry validation": (
        (
            "Original input",
            "Load Layer A and select Geometry validation. The tool retains geometries as parsed from the input before automatic 2D cleanup, repair, or duplicate removal. GeoPackage/Shapefile coordinates are first transformed to WGS84. Parser notices are included for records that could not be constructed. An unreadable or structurally malformed file can still fail to load.",
        ),
        (
            "Report-only or repair copies",
            "Run with Create repaired polygon copies unchecked for a report only. Enable it to retain valid polygons and attempt repairs of invalid polygons. Source files and the loaded original records stay unchanged. Null, empty, unsupported-coordinate, and unrecoverable shapes are excluded from repaired polygon output.",
        ),
        (
            "Interpret the report",
            "Validation shows the original status and reason, original/output geometry types and part counts, repair action, output area, and any discarded line or point pieces. An invalid original polygon's area is left blank because it is unreliable. Validity is evaluated in 2D; repair output is 2D. Review the resulting shape before choosing Use as Layer A or exporting it.",
        ),
    ),
    "Boundary / vertices": (
        (
            "Boundary lines",
            "Load Layer A, select Boundary / vertices, and choose Boundaries. Each exterior ring becomes a closed line. Include hole rings is enabled by default. The table lists source feature, part, ring type/index, geodesic length, and vertex count excluding the closing duplicate.",
        ),
        (
            "Numbered vertices",
            "Choose Vertices for point geometry and a coordinate table. Vertex numbering starts at 1 for each ring; exterior ring index is 0 and holes start at 1. By default the repeated closing coordinate is omitted. Enable Include the repeated closing vertex when you need a complete ring sequence.",
        ),
        (
            "Export",
            "Export lines or points as GeoJSON, KML, or KMZ; export the report as XLSX. Coordinates are WGS84 longitude/latitude in 2D. Attributes are copied with collision-safe metadata identifying the source and ring. Line and point results cannot replace polygon Layer A.",
        ),
    ),
    "Distance Calculator": (
        (
            "Inputs",
            "Select Distance Calculator. Load a query-location table and a reference-site table (CSV or first XLSX worksheet). Map an ID, latitude, and longitude column for each. Coordinates must be WGS84. Polygon layers are not used. Swap latitude and longitude applies only to the query table.",
        ),
        (
            "Options",
            "Choose nearest N (1–1000). Optionally exclude reference rows with exactly the same trimmed ID as the query. Enable Ignore 0 distance to skip coincident reference points regardless of their IDs. This option is off by default and works independently of ID exclusion. Zero is checked before rounding: small positive distances are still eligible. Exclusions apply before choosing the nearest N, so the next eligible sites fill the ranks.",
        ),
        (
            "Results",
            "Run to obtain site IDs, source row numbers, rank, distance in meters and kilometers, and bearing clockwise from true north, measured from query to site. With Ignore 0 distance off, coincident points have zero distance and a blank bearing. With it on, they are omitted; a query with no remaining candidates appears in Unmatched queries. Equal distances follow reference input order. Repeated site IDs remain separate candidates. If fewer than N candidates qualify, all qualifying rows are returned.",
        ),
        (
            "Review and export",
            "Review Unmatched queries, Rejected queries, and Rejected sites. Export the complete Excel report. Distances use the WGS84 ellipsoid, including across the date line. Large combinations of queries and sites take longer; calculations run in cancellable batches without allocating an all-pairs distance matrix.",
        ),
    ),
    "Appearance": [
        (
            "Light and Dark themes",
            "Open Help → Appearance and choose Light or Dark. The choice is saved for the next launch and also updates any open help windows. Your inputs and results stay loaded.",
        ),
        (
            "Subtle motion",
            "Short fades introduce operation headings, result summaries, and help content. Enable Help → Appearance → Reduce motion to turn these transitions off immediately. This preference is saved. Processing does not wait for animations.",
        ),
    ],
    "Start here": (
        (
            "Your first result",
            "1. Choose an operation from the sidebar or the dropdown on narrow windows. Only the inputs relevant to that operation will appear. Inputs always stay on the left and results on the right. Drag the vertical divider to adjust their widths. Each panel scrolls independently; narrow windows may also need horizontal scrolling.\n\n2. For polygon operations, click Choose files beside Layer A and select polygon files. For Distance Calculator, load Query locations and Reference sites instead. Wait for loading to finish; each new selection replaces that input.\n\n3. Add the other inputs requested by the operation: Layer B, a point table, or a distance.\n\n4. Check the settings, then click Run operation (Ctrl+Enter). Progress appears at the bottom of the window.\n\n5. Read the Summary and result tabs. Review Warnings or Rejected rows when present. All result rows are available in the table. Hover to read a full value, or double-click a row to view its details.\n\n6. Click Export, choose a format and filename, and save. Use Open export to open the saved file. On narrower panels, Open export and Use result as Layer A are available under More actions.",
        ),
        (
            "Continue working",
            "For geometry results, Use result as Layer A lets you perform the next operation on that result. Loading other inputs does not erase the previous result; the result heading identifies the operation that produced it.",
        ),
        (
            "Shortcuts and cancellation",
            "Ctrl+O: load Layer A.\nCtrl+Enter: run the selected operation.\nF1: open this guide.\nEscape: close a help window.\n\nCancel stops at the next safe checkpoint. A geometry calculation or workbook write already in progress must return first. Help remains available during processing.",
        ),
    ),
    "Point matching": (
        (
            "1 · Load polygons and points",
            "Select Points inside polygon. Load the polygon files into Layer A, then choose a CSV or XLSX point table. XLSX uses its first worksheet. The three required fields are a point identifier, latitude, and longitude.",
        ),
        (
            "2 · Check the column mapping",
            "Choose the correct Point name, Latitude, and Longitude columns. Common headers such as Site ID, lat, and lon are suggested automatically. The three selections must be different.\n\nCoordinates may be decimal degrees or supported degrees/minutes/seconds notation. Unrecognized values are reported rather than silently changed. Input rows retain their order and duplicate point names are kept.",
        ),
        (
            "3 · Choose the matching rules",
            "Include points on polygon and hole boundaries is enabled by default. Turn it off for strict interior matching. Points inside a hole do not match that polygon.\n\nUse Swap latitude and longitude only if the mapped coordinate values are reversed. The tool does not guess and swap them automatically. Full numeric precision is used for matching.",
        ),
        (
            "4 · Run and inspect",
            "Click Run operation. A point can match several overlapping polygons; their names are joined with semicolons. Match count records the actual number of matching names. Multipart polygons sharing exactly the same name appear once per point.\n\nUnmatched points remain in Matches as Not in any polygon. Invalid rows appear in Rejected rows with their input row number and rejection reason.\n\nReport colors: pink means no match, yellow means two matching names, and orange means three or more. One match has no special fill. Dark uses muted equivalents; selected rows use teal. Excel keeps the original light report colors in either theme.",
        ),
        (
            "5 · Export the workbook",
            "Export as XLSX. The workbook includes Summary, Matches, Empty polygons, Polygon counts, Rejected rows, and Warnings when applicable. Counts and empty status are grouped by polygon name. Two polygons with the same name therefore share one count.\n\nExtra columns from the original point table are not copied into the report. Coordinates display to six decimal places without rounding the values used by matching.",
        ),
    ),
    "Polygon operations": (
        (
            "Combine",
            "Load A. To include a second layer, select Include Layer B and load B. Run Combine to merge all selected polygons into one footprint. Disconnected pieces remain in one multipart feature.",
        ),
        (
            "Intersection",
            "Load A and B, then run Intersection. The result contains the area shared by their merged footprints. A shared edge or point alone produces no polygon area.",
        ),
        (
            "Difference",
            "Load the polygons to keep in A and the polygons to subtract in B. Run Difference. The operation is A minus B, so layer order matters.",
        ),
        (
            "SymDifference",
            "Load A and B, then run SymDifference. The result retains areas belonging to either layer and removes their overlapping area.",
        ),
        (
            "Buffer",
            "Load A and enter a distance in Meters or Kilometers. Positive values expand each feature; negative values shrink it. Run Buffer and inspect the result. A large negative distance may remove a feature completely. Export or reuse the result as A.",
        ),
        (
            "Simplify",
            "Load A and enter a positive tolerance in Meters or Kilometers. A smaller tolerance retains more detail; a larger tolerance removes more vertices. Run Simplify and review the output. Each feature's topology is preserved, but neighboring features are simplified independently and their shared borders may differ.",
        ),
        (
            "Geometry and attribute rules",
            "Combine, Intersection, Difference, and SymDifference work on merged geographic footprints. They do not overlay the original attribute tables. Buffer and Simplify retain each feature's name and properties.\n\nBuffer and Simplify use a local UTM projection. Their input extent must be at most 6° wide and 8° high, within 80°S to 84°N. Split broader datasets into local regions first.",
        ),
    ),
    "Measurements": (
        (
            "Measure area",
            "1. Load Layer A.\n2. Select Area.\n3. Choose the total measurement scope.\n4. Click Run operation.\n5. Read the total in Summary and individual rows in Measurements.\n6. Export an XLSX report.\n\nPer-feature areas are available in square meters, square kilometers, and acres. Polygon holes are excluded from area.",
        ),
        (
            "Measure perimeter",
            "1. Load Layer A.\n2. Select Perimeter.\n3. Choose the total measurement scope.\n4. Click Run operation.\n5. Read Summary and Measurements.\n6. Export an XLSX report.\n\nResults use meters, kilometers, and miles. Perimeter includes both exterior boundaries and the boundaries of holes.",
        ),
        (
            "Choose how overlaps are counted",
            "Merged coverage measures the union of A's features. It avoids counting overlapping area more than once and removes internal shared boundaries from the total perimeter.\n\nSum of features adds each feature's measurement separately. Overlapping areas and shared boundaries can contribute more than once.\n\nPer-feature table rows always describe individual features. Their sum may differ from the merged-coverage total, which is expected.",
        ),
        (
            "Measurement model",
            "Area and Perimeter are calculated on the WGS84 ellipsoid. They produce tabular results, not a new geometry layer. This tool is intended for regional geographic analysis; dateline-crossing and greater-than-hemisphere datasets must be split before use.",
        ),
    ),
    "Files & export": (
        (
            "Supported inputs",
            "Polygons: KML, KMZ, ZIP containing KML/KMZ, and WGS84 GeoJSON. Shapefile and GeoPackage input require the optional geopandas and pyogrio packages. All spatial layers in a GeoPackage are loaded.\n\nPoints: CSV or XLSX, using the first Excel worksheet. CSV supports common delimiters.\n\nAll KML documents inside a KMZ are read. ZIP files inside ZIP files are not supported. Analysis uses two dimensions; altitude is ignored.",
        ),
        (
            "Coordinate systems",
            "KML and GeoJSON use WGS84 longitude/latitude. Shapefile and GeoPackage inputs are reprojected to WGS84 from their declared CRS. Files without a declared CRS are rejected rather than assigned an assumed coordinate system.\n\nPolygon coordinates outside geographic limits or extents crossing the dateline are rejected. Review notices about repaired or skipped geometries before using results.",
        ),
        (
            "Choose an output format",
            "Geometry results: GeoJSON, KML, or KMZ. These retain polygon holes, multipart geometry, names, and available properties.\n\nPoint matching and measurements: XLSX.\n\nXLSX export of a geometry operation contains a summary and feature inventory; it does not store coordinate geometry. A result with no polygon area can still be exported as an XLSX report.",
        ),
        (
            "Saving and reopening",
            "Choose a new output filename. The interface prevents overwriting a currently loaded source file. An export is written to a temporary file and replaces the destination only after successful completion.\n\nOpen export uses your system's default application. Results are kept in memory during the session; export anything you want to retain before closing the tool. This version does not save workspaces.",
        ),
    ),
    "Troubleshooting": (
        (
            "Required packages",
            "The window uses PySide6 Qt Widgets. It rearranges inputs and results to fit the window. Processing requires shapely, pyproj, pandas, openpyxl, and defusedxml. Install them with the same Python interpreter that launches the tool:\n\npython -m pip install PySide6-Essentials shapely pyproj pandas openpyxl defusedxml\n\nOptional Shapefile/GeoPackage input:\npython -m pip install geopandas pyogrio",
        ),
        (
            "Points do not match",
            "Check the column mapping and the latitude/longitude order. Confirm that the point coordinates and polygon data describe the same region in WGS84. Check holes and the boundary checkbox. Review Rejected rows for invalid values. Zero valid rows still produce a report you can inspect.",
        ),
        (
            "An operation has no polygon output",
            "A and B may not overlap, Difference may remove all of A, or a negative Buffer may shrink every feature away. Intersection of edges alone has no polygon area. Check the inputs and operation order before rerunning.",
        ),
        (
            "An export fails",
            "Close the destination workbook if Excel has locked it. Choose a writable folder and a supported filename extension. Excel supports up to 1,048,575 data rows per sheet; split larger point inputs. The previous destination file remains intact if writing fails.",
        ),
        (
            "Processing takes time",
            "Complex boundaries, many overlapping polygons, and large files require more work. The interface remains available while a worker processes the data. Cancel is cooperative; it cannot interrupt the middle of a native geometry calculation. Very large inputs must fit in memory.",
        ),
        (
            "Request support",
            "Open Help → Contact & Support. Copy the support details and describe the selected operation, input formats, reproduction steps, expected result, and actual result or error message. Email opens a draft in your mail application; you review and send it yourself.",
        ),
    ),
}


def support_details() -> str:
    """Collect useful version information only when the user requests it."""
    import platform
    from importlib.metadata import PackageNotFoundError, version

    lines = [
        APP_NAME,
        f"Version: {__version__}",
        f"Build: {BUILD_NUMBER}",
        f"Release: {RELEASE_DATE}",
        f"Author: {AUTHOR_NAME}",
        f"Python: {platform.python_version()}",
        f"System: {platform.system()} {platform.release()} ({platform.machine()})",
    ]
    for package in (
        "PySide6-Essentials",
        "shapely",
        "pyproj",
        "pandas",
        "openpyxl",
        "defusedxml",
        "geopandas",
        "pyogrio",
    ):
        try:
            installed = version(package)
        except PackageNotFoundError:
            installed = "Not installed"
        lines.append(f"{package}: {installed}")
    return "\n".join(lines)


# ===========================================================================
# RESPONSIVE PYSIDE6 INTERFACE
# ===========================================================================

from html import escape
from shiboken6 import isValid
from PySide6.QtCore import (
    QPropertyAnimation,
    QEasingCurve,
    QAbstractTableModel,
    QEvent,
    QModelIndex,
    QObject,
    QPointF,
    QRect,
    QRectF,
    QSize,
    Qt,
    QThread,
    QTimer,
    QUrl,
    Signal,
    Slot,
)
from PySide6.QtGui import (
    QActionGroup,
    QPalette,
    QColor,
    QCursor,
    QDesktopServices,
    QFont,
    QIcon,
    QKeySequence,
    QPainter,
    QPainterPath,
    QPen,
    QPixmap,
    QPolygonF,
    QShortcut,
)
from PySide6.QtWidgets import (
    QGraphicsOpacityEffect,
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QComboBox,
    QDialog,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QListWidget,
    QMainWindow,
    QMenu,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QSplitter,
    QStackedWidget,
    QTabBar,
    QTableView,
    QTextBrowser,
    QToolButton,
    QVBoxLayout,
    QWidget,
    QProgressBar,
    QSpinBox,
    QStyle,
    QStyleOptionButton,
    QStyleOptionSpinBox,
    QMessageBox,
)

STYLE = """
QWidget { color: #223748; background-color: transparent; }
QMainWindow, QDialog { background: #F3F6F9; }
QFrame#Header { background: #FFFFFF; border-bottom: 1px solid #DEE6EC; }
QFrame#Sidebar { background: #EAF0F4; border: 1px solid #E0E8EF; border-radius: 12px; }
QFrame#Panel { background: #FFFFFF; border: 1px solid #DEE6EC; border-radius: 12px; }
QFrame#FileCard { background: #F8FAFC; border: 1px solid #DFE7ED; border-radius: 9px; }
QFrame#Metric { background: #F2F7F8; border: 1px solid #DFECEE; border-radius: 9px; }
QLabel[role="muted"], QLabel[role="caption"] { color: #5C7081; }
QLabel[role="section"] { color: #5A7183; }
QLabel[role="metric"] { color: #116978; }
QLabel#Error { color: #A22B38; background: #FFF0F0; padding: 8px; border-radius: 6px; }
QPushButton, QToolButton { background: #FFFFFF; border: 1px solid #CDD9E2; border-radius: 7px; padding: 8px 12px; }
QPushButton:hover, QToolButton:hover { background: #EAF4F5; border-color: #7AA7AE; }
QPushButton:pressed, QToolButton:pressed { background: #DCECEF; }
QPushButton:checked { background: #D4E8EB; border-color: #4A959E; color: #075A67; }
QPushButton:focus, QToolButton:focus { border: 2px solid #15798C; }
QPushButton:disabled, QToolButton:disabled { color: #8595A0; background: #F0F3F5; border-color: #E0E6EA; }
QPushButton#Primary { background: #176D7B; color: white; border-color: #176D7B; font-weight: 600; }
QPushButton#Primary:hover { background: #125966; }
QPushButton#Primary:disabled { background: #C0CFD3; border-color: #C0CFD3; color: #526972; }
QPushButton#Nav { text-align: left; background: transparent; border: 2px solid transparent; padding: 9px 8px; }
QPushButton#Nav:hover { background: #DCE9EE; }
QPushButton#Nav:checked { background: #D4E8EB; color: #075A67; border-left: 3px solid #137889; font-weight: 600; }
QPushButton#Nav:focus { border: 2px solid #137889; }
QLineEdit, QComboBox, QSpinBox { background: white; border: 1px solid #C7D5DF; border-radius: 6px; padding: 8px; selection-background-color: #176D7B; selection-color: white; }
QComboBox { padding-right: 22px; }
QSpinBox { padding-right: 28px; }
QSpinBox::up-button { subcontrol-origin: border; subcontrol-position: top right; width: 24px; border: none; }
QSpinBox::down-button { subcontrol-origin: border; subcontrol-position: bottom right; width: 24px; border: none; }
QSpinBox::up-arrow, QSpinBox::down-arrow { image: none; }
QLineEdit { placeholder-text-color: #5C7081; }
QLineEdit:disabled { placeholder-text-color: #8294A6; }
QComboBox::drop-down { border: none; width: 22px; }
QComboBox QAbstractItemView { background: white; selection-background-color: #D8EBEE; selection-color: #183F49; }
QLineEdit:focus, QComboBox:focus, QSpinBox:focus { border: 2px solid #15798C; }
QLineEdit[invalid="true"], QComboBox[invalid="true"] { border: 2px solid #B43A47; }
QLineEdit:disabled, QComboBox:disabled, QSpinBox:disabled { background: #EDF1F4; color: #788B98; }
QCheckBox { spacing: 8px; padding-top: 4px; padding-bottom: 4px; }
QCheckBox::indicator { width: 16px; height: 16px; }
QScrollArea { background: transparent; border: none; }
QScrollBar:vertical { background: #F0F3F5; width: 11px; margin: 0; }
QScrollBar::handle:vertical { background: #B4C5CF; min-height: 24px; border-radius: 5px; }
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0; }
QScrollBar:horizontal { background: #F0F3F5; height: 11px; margin: 0; }
QScrollBar::handle:horizontal { background: #B4C5CF; min-width: 24px; border-radius: 5px; }
QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal { width: 0; }
QSplitter::handle { background: transparent; }
QSplitter::handle:hover { background: #B8D9DE; border-radius: 3px; }
QTableView { background: white; alternate-background-color: #F7FAFC; border: 1px solid #E0E7ED; gridline-color: #EDF1F5; selection-background-color: #D6EBEF; selection-color: #193A47; }
QTableView::item { padding: 0px 8px; }
QTableView::item:selected { background: #176D7B; color: white; }
QTableView::item:selected:!active { background: #176D7B; color: white; }
QHeaderView::section { background: #EDF3F7; color: #3A5265; border: none; border-bottom: 1px solid #D7E2E9; padding: 9px 10px; font-weight: 600; }
QTabBar::tab { background: transparent; color: #5B7080; padding: 9px 12px; border-bottom: 3px solid transparent; }
QTabBar::tab:selected { color: #116978; border-bottom: 3px solid #19788A; font-weight: 600; }
QTabBar::tab:hover { background: #EEF6F7; }
QTextBrowser { background: white; border: none; padding: 8px; }
QListWidget { background: #EAF0F4; border: none; border-radius: 8px; padding: 6px; }
QListWidget::item { padding: 10px; border-radius: 6px; }
QListWidget::item:selected { background: #176D7B; color: white; }
QListWidget#GuideTopics { outline: 0; }
QListWidget#GuideTopics::item { padding: 0px 10px; border: 1px solid transparent; }
QListWidget#GuideTopics::item:hover { background: #DCE9EE; }
QListWidget#GuideTopics::item:selected { background: #176D7B; color: white; }
QListWidget#GuideTopics::item:selected:!active { background: #176D7B; color: white; }
QListWidget#GuideTopics::item:focus { border: 1px solid #6CE0DE; }
QProgressBar { border: none; border-radius: 2px; background: #DCE8ED; min-height: 4px; max-height: 4px; }
QProgressBar::chunk { background: #158092; }
QMenu { background: white; border: 1px solid #D4DFE7; padding: 5px; }
QMenu::item { padding: 8px 22px; }
QMenu::item:selected { background: #DBEEF0; color: #155766; }
QToolTip { background: #233D4C; color: white; border: none; padding: 7px; }
"""


# CellScope V11.1.0 brand palette. Semantic roles keep custom painting and Qt
# controls consistent, including widgets created after a theme change.
DARK_COLORS = {}
for target, sources in {
    "#111c2a": "#F3F6F9 #F0F3F5 #EDF1F4",
    "#1c2b3d": "#FFFFFF #F8FAFC #EAF0F4 #233D4C",
    "#223347": "#F2F7F8 #F7FAFC #EDF3F7 #EAF4F5 #DCECEF #DCE9EE #EEF6F7 #DCE8ED",
    "#3c5064": "#DEE6EC #E0E8EF #DFE7ED #DFECEE #CDD9E2 #E0E6EA #C7D5DF #E0E7ED #EDF1F5 #D7E2E9 #D4DFE7 #B4C5CF #B8D9DE #C0CFD3 #93A8B5 #B9C6CF",
    "#edf4fa": "#223748 #243F50 #075A67 #183F49 #193A47 #155766",
    "#a8bccd": "#5C7081 #5A7183 #3A5265 #5B7080 #52717F",
    "#8294a6": "#8595A0 #526972 #788B98 #9CACB4",
    "#6ce0de": "#116978 #7AA7AE #4A959E",
    "#128d94": "#15798C #137889 #19788A #158092 #125966",
    "#086e76": "#176D7B #D4E8EB #D8EBEE #D6EBEF #DBEEF0",
    "#ff9b9b": "#A22B38 #B43A47",
    "#402a35": "#FFF0F0 #FFF0EF #FFEBEE",
    "#493326": "#FFE0B2",
    "#3d3524": "#FFF6DF #FFF9C4",
}.items():
    DARK_COLORS.update({source.lower(): target for source in sources.split()})


DARK_COLORS["#93a8b5"] = "#a8bccd"  # Unchecked indicators need a clear outline.


def theme_color(value):
    app = QApplication.instance()
    return (
        DARK_COLORS.get(value.lower(), value)
        if app and app.property("darkTheme")
        else value
    )


def apply_theme(name):
    app = QApplication.instance()
    dark = name == "dark"
    app.setProperty("darkTheme", dark)
    palette = app.style().standardPalette()
    if dark:
        roles = {
            "Window": "#111c2a",
            "WindowText": "#edf4fa",
            "Base": "#1c2b3d",
            "AlternateBase": "#223347",
            "Text": "#edf4fa",
            "Button": "#1c2b3d",
            "ButtonText": "#edf4fa",
            "ToolTipBase": "#1c2b3d",
            "ToolTipText": "#edf4fa",
            "Highlight": "#086e76",
            "HighlightedText": "#ffffff",
            "Link": "#6ce0de",
            "LinkVisited": "#c58af9",
            "PlaceholderText": "#a8bccd",
        }
        for role, color in roles.items():
            palette.setColor(getattr(QPalette.ColorRole, role), QColor(color))
        for role in (
            QPalette.ColorRole.WindowText,
            QPalette.ColorRole.Text,
            QPalette.ColorRole.ButtonText,
        ):
            palette.setColor(QPalette.ColorGroup.Disabled, role, QColor("#8294a6"))
    if dark:
        for role, value in {
            "Light": "#3c5064",
            "Midlight": "#3c5064",
            "Mid": "#223347",
            "Dark": "#111c2a",
            "Shadow": "#111c2a",
        }.items():
            palette.setColor(getattr(QPalette.ColorRole, role), QColor(value))
        palette.setColor(
            QPalette.ColorGroup.Disabled,
            QPalette.ColorRole.Highlight,
            QColor("#3c5064"),
        )
        palette.setColor(
            QPalette.ColorGroup.Disabled,
            QPalette.ColorRole.HighlightedText,
            QColor("#edf4fa"),
        )
    app.setPalette(palette)
    stylesheet = re.sub(r"#[0-9a-fA-F]{6}", lambda m: theme_color(m.group()), STYLE)
    if dark:
        stylesheet = stylesheet.replace("background: white", "background: #1c2b3d")
    app.setStyleSheet(stylesheet)
    for widget in app.allWidgets():
        widget.update()


class GentleMotion(QObject):
    """Temporary fades with explicit cleanup on hide, deletion, and interruption."""

    def __init__(self, owner, reduced=False):
        super().__init__(owner)
        self.reduced = reduced
        self.active = {}

    def stop_all(self):
        for widget in list(self.active):
            self._finish(widget)

    def _discard(self, widget):
        animation = self.active.pop(widget, None)
        if animation is not None and isValid(animation):
            animation.stop()
            animation.deleteLater()

    def _finish(self, widget):
        active = widget in self.active
        self._discard(widget)
        if active and isValid(widget):
            widget.removeEventFilter(self)
            widget.setGraphicsEffect(None)

    def eventFilter(self, widget, event):
        if event.type() in (QEvent.Type.Hide, QEvent.Type.DeferredDelete):
            self._finish(widget)
        return False

    def reveal(self, widget):
        self._finish(widget)
        if self.reduced or not isValid(widget) or not widget.isVisible():
            return
        if not widget.property("motionTracked"):
            widget.setProperty("motionTracked", True)
            widget.destroyed.connect(lambda: self._discard(widget))
        effect = QGraphicsOpacityEffect(widget)
        effect.setOpacity(0.65)
        widget.setGraphicsEffect(effect)
        widget.installEventFilter(self)
        animation = QPropertyAnimation(effect, b"opacity", self)
        animation.setDuration(150)
        animation.setStartValue(0.65)
        animation.setEndValue(1.0)
        animation.setEasingCurve(QEasingCurve.Type.OutCubic)
        self.active[widget] = animation
        animation.finished.connect(lambda: self._finish(widget))
        animation.start()


def make_icon(operation=None):
    """Vector artwork rendered at several sizes; no external icon assets."""
    icon = QIcon()
    for size in (16, 24, 32, 48, 64, 128):
        pixmap = QPixmap(size, size)
        pixmap.fill(Qt.GlobalColor.transparent)
        painter = QPainter(pixmap)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.scale(size / 64, size / 64)
        a = QPainterPath()
        a.addPolygon(
            QPolygonF(
                [
                    QPointF(x, y)
                    for x, y in ((12, 17), (34, 11), (46, 28), (35, 48), (13, 42))
                ]
            )
        )
        a.closeSubpath()
        b = QPainterPath()
        b.addPolygon(
            QPolygonF(
                [
                    QPointF(x, y)
                    for x, y in ((29, 24), (50, 18), (55, 43), (37, 54), (22, 42))
                ]
            )
        )
        b.closeSubpath()
        if operation is None:
            painter.setPen(Qt.PenStyle.NoPen)
            painter.setBrush(QColor("#193141"))
            painter.drawRoundedRect(QRectF(2, 2, 60, 60), 10, 10)
        painter.setPen(QPen(QColor("#CEECEF" if operation is None else "#397C92"), 2.5))
        painter.setBrush(QColor("#4084B0"))
        if operation in ("Area", "Perimeter", "Simplify", "Buffer"):
            if operation == "Buffer":
                painter.setPen(QPen(QColor("#94CBCB"), 8))
            if operation == "Perimeter":
                painter.setBrush(Qt.BrushStyle.NoBrush)
                painter.setPen(
                    QPen(
                        QColor(
                            "#6ce0de"
                            if operation
                            and QApplication.instance().property("darkTheme")
                            else "#176D7B"
                        ),
                        5,
                    )
                )
            painter.drawPath(a)
            if operation == "Area":
                painter.setPen(QPen(QColor("#D2EFEF"), 2))
                for x in (22, 29, 36):
                    painter.drawLine(QPointF(x, 23), QPointF(x, 38))
        else:
            painter.drawPath(a)
            painter.setBrush(QColor("#299F98"))
            painter.drawPath(b)
            painter.setPen(Qt.PenStyle.NoPen)
            painter.setBrush(QColor("#70D0BE"))
            if operation == "Difference":
                painter.setBrush(QColor("#F0F5F8"))
            painter.drawPath(a.intersected(b))
        if operation in (None, "Points inside polygon"):
            painter.setBrush(QColor("#F6B043"))
            painter.setPen(QPen(QColor("white"), 2.3))
            painter.drawEllipse(QPointF(34, 33), 5, 5)
        painter.end()
        icon.addPixmap(pixmap)
    return icon


def label(text="", role=None, wrap=False):
    widget = QLabel(text)
    widget.setWordWrap(wrap)
    widget.setSizePolicy(
        QSizePolicy.Policy.Ignored if wrap else QSizePolicy.Policy.Preferred,
        QSizePolicy.Policy.Preferred,
    )
    if role:
        widget.setProperty("role", role)
    return widget


def typography(root):
    base = QApplication.font()
    points = base.pointSizeF() if base.pointSizeF() > 0 else 10
    factors = {
        "title": 1.9,
        "heading": 1.35,
        "metric": 1.9,
        "section": 0.9,
        "caption": 0.9,
    }
    for widget in root.findChildren(QLabel):
        role = widget.property("role")
        if role in factors:
            font = QFont(base)
            font.setPointSizeF(points * factors[role])
            font.setBold(role != "caption")
            widget.setFont(font)


def scroll_area(content):
    area = QScrollArea()
    area.setWidgetResizable(True)
    area.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
    area.setWidget(content)
    area.setMinimumSize(0, 0)
    return area


class ClearCheckBox(QCheckBox):
    """Keep both checked and unchecked states visible on every theme."""

    def __init__(self, text="", parent=None):
        super().__init__(text, parent)
        policy = QSizePolicy(QSizePolicy.Policy.Preferred, QSizePolicy.Policy.Minimum)
        policy.setHeightForWidth(True)
        self.setSizePolicy(policy)

    def minimumSizeHint(self):
        return QSize(120, self.fontMetrics().height() + 8)

    def heightForWidth(self, width):
        bounds = self.fontMetrics().boundingRect(
            QRect(0, 0, max(20, width - 30), 10000),
            Qt.TextFlag.TextWordWrap,
            self.text(),
        )
        return max(26, bounds.height() + 8)

    def paintEvent(self, event):
        option = QStyleOptionButton()
        self.initStyleOption(option)
        option.text = ""
        background = QPainter(self)
        self.style().drawControl(
            QStyle.ControlElement.CE_CheckBox, option, background, self
        )
        background.end()
        rect = self.style().subElementRect(
            QStyle.SubElement.SE_CheckBoxIndicator, option, self
        )
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setPen(
            QColor(theme_color("#243F50" if self.isEnabled() else "#9CACB4"))
        )
        painter.drawText(
            self.rect().adjusted(30, 0, -2, 0),
            Qt.AlignmentFlag.AlignVCenter | Qt.TextFlag.TextWordWrap,
            self.text(),
        )
        color = "#176D7B" if self.isChecked() else "#93A8B5"
        painter.setPen(
            QPen(QColor(theme_color(color if self.isEnabled() else "#B9C6CF")), 1.3)
        )
        painter.setBrush(QColor(theme_color(color if self.isChecked() else "#FFFFFF")))
        painter.drawRoundedRect(QRectF(rect).adjusted(1, 1, -1, -1), 3, 3)
        if self.isChecked():
            painter.setPen(
                QPen(
                    QColor("white"),
                    2,
                    Qt.PenStyle.SolidLine,
                    Qt.PenCapStyle.RoundCap,
                    Qt.PenJoinStyle.RoundJoin,
                )
            )
            path = QPainterPath(QPointF(rect.left() + 4, rect.center().y()))
            path.lineTo(rect.left() + 7, rect.bottom() - 4)
            path.lineTo(rect.right() - 3, rect.top() + 4)
            painter.drawPath(path)


class CountBox(QSpinBox):
    """Draw legible step arrows independently of platform spinbox glyphs."""

    def paintEvent(self, event):
        super().paintEvent(event)
        option = QStyleOptionSpinBox()
        self.initStyleOption(option)
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setPen(Qt.PenStyle.NoPen)
        for control, direction, enabled in (
            (QStyle.SubControl.SC_SpinBoxUp, -1, self.value() < self.maximum()),
            (QStyle.SubControl.SC_SpinBoxDown, 1, self.value() > self.minimum()),
        ):
            rect = self.style().subControlRect(
                QStyle.ComplexControl.CC_SpinBox, option, control, self
            )
            center = rect.center()
            x, y = center.x(), center.y()
            painter.setBrush(
                QColor(
                    theme_color(
                        "#52717F" if self.isEnabled() and enabled else "#9CACB4"
                    )
                )
            )
            painter.drawPolygon(
                QPolygonF(
                    [
                        QPointF(x - 4, y - direction * 2),
                        QPointF(x + 4, y - direction * 2),
                        QPointF(x, y + direction * 3),
                    ]
                )
            )


class ChoiceBox(QComboBox):
    """A clear arrow remains visible with the custom control border style."""

    def paintEvent(self, event):
        super().paintEvent(event)
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.setPen(Qt.PenStyle.NoPen)
        painter.setBrush(
            QColor(theme_color("#52717F" if self.isEnabled() else "#9CACB4"))
        )
        x, y = self.width() - 14, self.height() / 2
        painter.drawPolygon(
            QPolygonF([QPointF(x - 4, y - 2), QPointF(x + 4, y - 2), QPointF(x, y + 3)])
        )


class ElidedLabel(QLabel):
    def __init__(self, text="", parent=None):
        super().__init__(parent)
        self.full_text = text
        self.setMinimumWidth(0)
        self.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred)
        self.setText(text)

    def set_full_text(self, text, tooltip=None):
        self.full_text = text
        self.setToolTip(tooltip or text)
        self._elide()

    def setText(self, text):
        self.set_full_text(text)

    def _elide(self):
        super().setText(
            self.fontMetrics().elidedText(
                self.full_text, Qt.TextElideMode.ElideMiddle, max(10, self.width())
            )
        )

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._elide()


class FileCard(QFrame):
    def __init__(self, title, hint, callback):
        super().__init__()
        self.setObjectName("FileCard")
        self.setMinimumWidth(0)
        self.layout_grid = QGridLayout(self)
        self.layout_grid.setContentsMargins(14, 12, 14, 12)
        self.layout_grid.setSpacing(8)
        self.title = label(title, "section")
        self.filename = ElidedLabel("No file selected")
        font = QFont(self.filename.font())
        font.setBold(True)
        self.filename.setFont(font)
        self.meta = label(hint, "muted", True)
        self.button = QPushButton(
            "Choose files…" if title != "POINT TABLE" else "Choose table…"
        )
        self.button.setAccessibleName(f"Choose {title.lower()}")
        self.button.clicked.connect(callback)
        self.layout_grid.addWidget(self.title, 0, 0)
        self.layout_grid.addWidget(self.filename, 1, 0)
        self.layout_grid.addWidget(self.meta, 2, 0)
        self.layout_grid.setColumnStretch(0, 1)
        self.layout_grid.addWidget(self.button, 0, 1, 3, 1)
        self.narrow = None

    def set_loaded(self, paths, count, kind="features"):
        names = [Path(path).name for path in paths]
        self.filename.set_full_text(
            names[0] if len(names) == 1 else f"{len(names)} files selected",
            "\n".join(map(str, paths)),
        )
        self.meta.setText(f"Loaded · {count:,} {kind}")

    def resizeEvent(self, event):
        super().resizeEvent(event)
        narrow = self.width() < max(370, self.fontMetrics().averageCharWidth() * 48)
        if narrow != self.narrow:
            self.narrow = narrow
            self.layout_grid.removeWidget(self.button)
            if narrow:
                self.layout_grid.addWidget(
                    self.button, 3, 0, 1, 2, Qt.AlignmentFlag.AlignLeft
                )
            else:
                self.layout_grid.addWidget(self.button, 0, 1, 3, 1)


class RecordModel(QAbstractTableModel):
    """Qt requests visible cells; no per-cell widgets or truncated preview."""

    def __init__(self, records, parent=None):
        super().__init__(parent)
        self.records = records
        self.headers = list(records[0]) if records else ["No records"]
        self.coordinate_precision = 6

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.records)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.headers)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        value = self.records[index.row()].get(self.headers[index.column()], "")
        if role in (Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.ToolTipRole):
            if role == Qt.ItemDataRole.DisplayRole and isinstance(value, float):
                if self.headers[index.column()].startswith("UTM "):
                    return f"{value:,.2f}"
                return (
                    f"{value:,.{self.coordinate_precision}f}"
                    if self.headers[index.column()]
                    .lower()
                    .endswith(("latitude", "longitude"))
                    else f"{value:,.4f}"
                )
            return str(value)
        if role == Qt.ItemDataRole.TextAlignmentRole:
            return int(
                (
                    Qt.AlignmentFlag.AlignRight
                    if isinstance(value, (int, float))
                    else Qt.AlignmentFlag.AlignLeft
                )
                | Qt.AlignmentFlag.AlignVCenter
            )
        if (
            role == Qt.ItemDataRole.BackgroundRole
            and "Match count" in self.records[index.row()]
        ):
            count = self.records[index.row()]["Match count"]
            color = match_fill(count)
            if color:
                return QColor(theme_color("#" + color))
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole:
            return (
                self.headers[section]
                if orientation == Qt.Orientation.Horizontal
                else str(section + 1)
            )
        return None


class TaskWorker(QObject):
    progress = Signal(float, str)
    succeeded = Signal(object)
    failed = Signal(str)
    cancelled = Signal(str)
    finished = Signal()

    def __init__(self, task, cancel, logger):
        super().__init__()
        self.task, self.cancel, self.logger = task, cancel, logger

    @Slot()
    def run(self):
        try:
            self.succeeded.emit(self.task(self.cancel, self.progress.emit))
        except Cancelled as exc:
            self.cancelled.emit(str(exc))
        except Exception as exc:
            self.logger.exception("Background task failed")
            message = str(exc)
            if isinstance(exc, ImportError):
                message += "\nInstall the processing dependencies listed at the top of the Python file using the interpreter that runs this application."
            self.failed.emit(message)
        finally:
            self.finished.emit()


class ProductDialog(QDialog):
    """Scrollable content with a fixed footer and screen-bounded dimensions."""

    def __init__(self, owner, title, subtitle, width=680, height=560):
        super().__init__(owner)
        self.setWindowTitle(f"{title} — {APP_NAME}")
        self.setWindowIcon(owner.windowIcon())
        self.setAttribute(Qt.WidgetAttribute.WA_DeleteOnClose)
        self.setModal(False)
        self.setMinimumSize(300, 280)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 16)
        layout.setSpacing(16)
        head = QHBoxLayout()
        picture = QLabel()
        picture.setPixmap(owner.windowIcon().pixmap(40, 40))
        head.addWidget(picture)
        texts = QVBoxLayout()
        self.title_label = label(title, "title", True)
        texts.addWidget(self.title_label)
        texts.addWidget(label(subtitle, "muted", True))
        head.addLayout(texts, 1)
        layout.addLayout(head)
        self.body = QWidget()
        self.body_layout = QVBoxLayout(self.body)
        self.body_layout.setContentsMargins(0, 0, 0, 0)
        self.body_layout.setSpacing(14)
        self.scroller = scroll_area(self.body)
        layout.addWidget(self.scroller, 1)
        footer = QHBoxLayout()
        self.feedback = label(f"v{__version__} · {AUTHOR_NAME}", "muted", True)
        footer.addWidget(self.feedback, 1)
        close = QPushButton("Close")
        close.clicked.connect(self.close)
        footer.addWidget(close)
        layout.addLayout(footer)
        screen = owner.screen() or QApplication.primaryScreen()
        available = screen.availableGeometry()
        scale = max(1.0, QApplication.font().pointSizeF() / 10)
        self.resize(
            min(round(width * scale), available.width() - 32),
            min(round(height * scale), available.height() - 48),
        )
        geometry = self.geometry()
        geometry.moveCenter(owner.frameGeometry().center())
        geometry.moveLeft(
            max(
                available.left() + 12,
                min(geometry.left(), available.right() - geometry.width() - 12),
            )
        )
        geometry.moveTop(
            max(
                available.top() + 12,
                min(geometry.top(), available.bottom() - geometry.height() - 12),
            )
        )
        self.move(geometry.topLeft())

    def fit_content_height(self):
        """Fit short information dialogs while keeping small-screen scrolling."""
        if not isValid(self):
            return
        self.ensurePolished()
        self.layout().activate()
        self.body_layout.activate()
        width = self.scroller.viewport().width()
        content_height = (
            self.body_layout.totalHeightForWidth(width)
            if self.body_layout.hasHeightForWidth()
            else self.body_layout.sizeHint().height()
        )
        chrome_height = self.height() - self.scroller.height()
        available = self.screen().availableGeometry()
        height = min(content_height + chrome_height + 6, available.height() - 48)
        self.resize(self.width(), max(self.minimumHeight(), height))
        self.move(
            self.x(),
            max(
                available.top() + 12,
                min(self.y(), available.bottom() - self.height() - 36),
            ),
        )

    def copy_text(self, text, confirmation="Copied to clipboard."):
        QApplication.clipboard().setText(text)
        self.feedback.setText(confirmation)

    def open_link(self, url):
        accepted = QDesktopServices.openUrl(QUrl(url))
        self.feedback.setText(
            "Opened in your default application."
            if accepted
            else "Could not open the link. Copy the address and open it manually."
        )


class GuideDialog(ProductDialog):
    def __init__(self, owner):
        super().__init__(
            owner,
            "Quick Guide",
            "Complete steps, result meanings, and troubleshooting.",
            900,
            660,
        )
        self.pages = {"Start here": GUIDE_PAGES["Start here"]}
        for operation in OPERATIONS:
            if operation == "Points inside polygon":
                self.pages[operation] = GUIDE_PAGES["Point matching"]
            elif operation in {"Area", "Perimeter"}:
                index = 0 if operation == "Area" else 1
                self.pages[operation] = (
                    GUIDE_PAGES["Measurements"][index],
                    *GUIDE_PAGES["Measurements"][2:],
                )
            elif operation in GUIDE_PAGES:
                self.pages[operation] = GUIDE_PAGES[operation]
            else:
                self.pages[operation] = tuple(
                    page
                    for page in GUIDE_PAGES["Polygon operations"]
                    if page[0] in {operation, "Geometry and attribute rules"}
                )
        for topic in ("Appearance", "Files & export", "Troubleshooting"):
            self.pages[topic] = GUIDE_PAGES[topic]
        self.topics = list(self.pages)
        self.topic_combo = ChoiceBox()
        self.topic_combo.addItems(self.topics)
        self.body_layout.addWidget(self.topic_combo)
        self.columns = QHBoxLayout()
        self.navigation = QListWidget()
        self.navigation.setObjectName("GuideTopics")
        self.navigation.addItems(self.topics)
        self.navigation.setSpacing(3)
        self.navigation.setUniformItemSizes(True)
        self.navigation.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.navigation.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOff
        )
        self.navigation.setMaximumWidth(220)
        self.columns.addWidget(self.navigation)
        self.reader = QTextBrowser()
        self.reader.setMinimumHeight(180)
        self.columns.addWidget(self.reader, 1)
        self.body_layout.addLayout(self.columns, 1)
        self.navigation.currentRowChanged.connect(self.show_page)
        self.topic_combo.currentIndexChanged.connect(self.show_page)
        self.navigation.setCurrentRow(0)
        self.show_page(0)
        typography(self)
        self._size_topics()

    def _size_topics(self):
        height = self.navigation.fontMetrics().height() + 22
        for index in range(self.navigation.count()):
            self.navigation.item(index).setSizeHint(QSize(0, height))
        self.navigation.doItemsLayout()

    def changeEvent(self, event):
        super().changeEvent(event)
        if hasattr(self, "navigation") and event.type() in (
            QEvent.Type.FontChange,
            QEvent.Type.ApplicationFontChange,
            QEvent.Type.StyleChange,
        ):
            self._size_topics()

    def show_page(self, index):
        if index < 0:
            return
        self.navigation.blockSignals(True)
        self.topic_combo.blockSignals(True)
        self.navigation.setCurrentRow(index)
        self.topic_combo.setCurrentIndex(index)
        self.navigation.blockSignals(False)
        self.topic_combo.blockSignals(False)
        html = "".join(
            f"<h3 style='color:{'#6ce0de' if QApplication.instance().property('darkTheme') else '#176D7B'}'>{escape(heading)}</h3><p>{escape(paragraph).replace(chr(10), '<br>')}</p>"
            for heading, paragraph in self.pages[self.topics[index]]
        )
        self.reader.setHtml(html)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, "navigation"):
            compact = self.width() < 740 * max(
                1.0, QApplication.font().pointSizeF() / 10
            )
            self.navigation.setVisible(not compact)
            self.topic_combo.setVisible(compact)


class ToolkitApp(QMainWindow):
    def __init__(self, settings_dir=DATA_DIR):
        super().__init__()
        self.setWindowTitle(f"{APP_NAME} · {__version__}")
        self.setWindowIcon(make_icon())
        self.settings_dir = Path(settings_dir)
        self.settings = self._read_settings()
        self.theme = self.settings.get("theme", "light")
        if self.theme not in {"light", "dark"}:
            self.theme = "light"
        apply_theme(self.theme)
        self.motion = GentleMotion(
            self, bool(self.settings.get("reduce_motion", False))
        )
        self.logger = self._create_logger()
        self.layers = {"A": None, "B": None}
        self.point_table = None
        self.point_path = None
        self.site_table = None
        self.site_path = None
        self.cutter_layer = None
        self.result = None
        self.last_export = None
        self.busy = False
        self.closing = False
        self.cancel_event = threading.Event()
        self._thread = None
        self._worker = None
        self.help_windows = {}
        self.controls = []
        self.layout_mode = None
        self._reflow_pending = False
        self._restoring_split = False
        self.table_views = []
        self._build_ui()
        self._restore_window()
        self._select_operation(OPERATIONS[0])
        self._apply_typography()
        self._schedule_reflow()
        for shortcut, callback in (
            ("Ctrl+O", lambda: self._browse_layer("A")),
            ("Ctrl+Return", self._run),
            ("F1", self._show_quick_guide),
        ):
            QShortcut(QKeySequence(shortcut), self, activated=callback)
        QApplication.instance().screenRemoved.connect(
            lambda screen: QTimer.singleShot(0, self._fit_to_screen)
        )

    def _set_theme(self, name):
        self.motion.stop_all()
        self.theme = name
        self.settings["theme"] = name
        self.theme_actions[name].setChecked(True)
        apply_theme(name)
        for operation, button in self.nav_buttons.items():
            button.setIcon(make_icon(operation))
        for window in self.help_windows.values():
            if isinstance(window, GuideDialog):
                window.show_page(window.topic_combo.currentIndex())
            for item in window.findChildren(QLabel):
                if item.property("linkHtml"):
                    item.setText(
                        item.property("linkHtml").replace(
                            "#176D7B", "#6ce0de" if name == "dark" else "#176D7B"
                        )
                    )
        for view in self.table_views:
            view.viewport().update()
        self._save_settings()

    def _set_reduced_motion(self, reduced):
        self.motion.reduced = reduced
        self.motion.stop_all()
        self.settings["reduce_motion"] = reduced
        self._save_settings()

    def _read_settings(self):
        try:
            value = json.loads(
                (self.settings_dir / "settings.json").read_text(encoding="utf-8")
            )
            return value if isinstance(value, dict) else {}
        except (OSError, ValueError):
            return {}

    def _create_logger(self):
        logger = logging.getLogger(f"geotoolkit.qt.{id(self)}")
        logger.setLevel(logging.INFO)
        try:
            self.settings_dir.mkdir(parents=True, exist_ok=True)
            handler = RotatingFileHandler(
                self.settings_dir / "toolkit.log",
                maxBytes=2_000_000,
                backupCount=2,
                encoding="utf-8",
            )
            handler.setFormatter(
                logging.Formatter("%(asctime)s %(levelname)s %(message)s")
            )
            logger.addHandler(handler)
        except OSError:
            logger.addHandler(logging.NullHandler())
        return logger

    def _save_settings(self):
        try:
            self.settings_dir.mkdir(parents=True, exist_ok=True)
            temporary = self.settings_dir / "settings.tmp"
            temporary.write_text(json.dumps(self.settings, indent=2), encoding="utf-8")
            temporary.replace(self.settings_dir / "settings.json")
        except OSError:
            self.logger.warning("Could not save preferences", exc_info=True)

    def _restore_window(self):
        screen = QApplication.screenAt(QCursor.pos()) or QApplication.primaryScreen()
        available = screen.availableGeometry()
        saved = self.settings.get("qt_window_rect")
        if (
            isinstance(saved, list)
            and len(saved) == 4
            and all(isinstance(v, int) for v in saved)
        ):
            rect = QRect(*saved)
            if rect.width() > 0 and rect.height() > 0:
                screen = next(
                    (
                        s
                        for s in QApplication.screens()
                        if s.availableGeometry().intersects(rect)
                    ),
                    screen,
                )
                available = screen.availableGeometry()
            else:
                rect = QRect(0, 0, 1220, 820)
        else:
            rect = QRect(
                0,
                0,
                min(1220, round(available.width() * 0.90)),
                min(820, round(available.height() * 0.88)),
            )
            rect.moveCenter(available.center())
        rect.setWidth(min(max(420, rect.width()), max(240, available.width() - 24)))
        rect.setHeight(min(max(360, rect.height()), max(260, available.height() - 48)))
        rect.moveLeft(
            max(
                available.left() + 12,
                min(rect.left(), available.right() - rect.width() - 12),
            )
        )
        rect.moveTop(
            max(
                available.top() + 12,
                min(rect.top(), available.bottom() - rect.height() - 36),
            )
        )
        self.setMinimumSize(min(420, rect.width()), min(360, rect.height()))
        self.setGeometry(rect)

    def _fit_to_screen(self):
        if self.isMaximized() or self.isMinimized():
            return
        available = (self.screen() or QApplication.primaryScreen()).availableGeometry()
        rect = self.geometry()
        self.setMinimumSize(
            min(420, available.width() - 24), min(360, available.height() - 48)
        )
        rect.setWidth(min(rect.width(), available.width() - 24))
        rect.setHeight(min(rect.height(), available.height() - 48))
        rect.moveLeft(
            max(
                available.left() + 12,
                min(rect.left(), available.right() - rect.width() - 12),
            )
        )
        rect.moveTop(
            max(
                available.top() + 12,
                min(rect.top(), available.bottom() - rect.height() - 36),
            )
        )
        self.setGeometry(rect)

    def showEvent(self, event):
        super().showEvent(event)
        if self.windowHandle() and not getattr(self, "_screen_connected", False):
            self._screen_connected = True
            self.windowHandle().screenChanged.connect(
                lambda screen: QTimer.singleShot(0, self._fit_to_screen)
            )
        self._schedule_reflow()

    def _track(self, widget):
        self.controls.append(widget)
        return widget

    def _build_ui(self):
        central = QWidget()
        self.setCentralWidget(central)
        self.outer = QVBoxLayout(central)
        self.outer.setContentsMargins(0, 0, 0, 0)
        self.outer.setSpacing(0)
        header = QFrame()
        header.setObjectName("Header")
        self.header_layout = QHBoxLayout(header)
        self.header_layout.setContentsMargins(20, 14, 20, 14)
        logo = QLabel()
        logo.setPixmap(self.windowIcon().pixmap(40, 40))
        self.header_layout.addWidget(logo)
        brand = QVBoxLayout()
        brand.setSpacing(2)
        self.brand_title = label(APP_NAME, "title")
        self.brand_title.setToolTip(APP_NAME)
        self.brand_title.setMinimumWidth(0)
        self.brand_subtitle = label("Polygon analysis. Clear results.", "muted")
        brand.addWidget(self.brand_title)
        brand.addWidget(self.brand_subtitle)
        self.header_layout.addLayout(brand, 1)
        self.version_badge = label(f"v{__version__}", "caption")
        self.header_layout.addWidget(self.version_badge)
        self.help_button = QToolButton()
        self.help_button.setText("Help")
        self.help_button.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        appearance = QMenu("Appearance", self)
        self.theme_actions = {}
        group = QActionGroup(self)
        group.setExclusive(True)
        for title, name in (("Light", "light"), ("Dark", "dark")):
            action = appearance.addAction(title)
            action.setCheckable(True)
            action.setChecked(self.theme == name)
            group.addAction(action)
            action.triggered.connect(
                lambda checked=False, value=name: self._set_theme(value)
            )
            self.theme_actions[name] = action
        appearance.addSeparator()
        self.reduce_motion_action = appearance.addAction("Reduce motion")
        self.reduce_motion_action.setCheckable(True)
        self.reduce_motion_action.setChecked(self.motion.reduced)
        self.reduce_motion_action.toggled.connect(self._set_reduced_motion)
        help_menu = QMenu(self)
        help_menu.addMenu(appearance)
        help_menu.addSeparator()
        help_menu.addAction("About", self._show_about)
        help_menu.addAction("Quick Guide  ·  F1", self._show_quick_guide)
        help_menu.addSeparator()
        help_menu.addAction("Contact & Support", self._show_support)
        self.help_button.setMenu(help_menu)
        self.header_layout.addWidget(self.help_button)
        self.outer.addWidget(header)
        workspace = QWidget()
        self.workspace_layout = QVBoxLayout(workspace)
        self.workspace_layout.setContentsMargins(16, 16, 16, 10)
        self.workspace_layout.setSpacing(12)
        self.operation_combo = self._track(ChoiceBox())
        self.operation_combo.addItems(OPERATIONS)
        self.operation_combo.setAccessibleName("Operation")
        self.operation_combo.setSizePolicy(
            QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Fixed
        )
        self.operation_combo.currentTextChanged.connect(self._select_operation)
        self.workspace_layout.addWidget(self.operation_combo)
        self.workspace_row = QHBoxLayout()
        self.workspace_row.setSpacing(12)
        self.sidebar = QFrame()
        self.sidebar.setObjectName("Sidebar")
        navbox = QVBoxLayout(self.sidebar)
        navbox.setContentsMargins(10, 14, 10, 14)
        navcontent = QWidget()
        nav = QVBoxLayout(navcontent)
        nav.setContentsMargins(0, 0, 0, 0)
        nav.setSpacing(5)
        self.nav_buttons = {}
        for category, operations in (
            ("MATCH", OPERATIONS[:1]),
            ("OVERLAY", OPERATIONS[1:5]),
            ("TRANSFORM", OPERATIONS[5:7]),
            ("MEASURE", OPERATIONS[7:9]),
            ("PROXIMITY", OPERATIONS[9:10]),
            ("POLYGON TOOLS", OPERATIONS[10:14]),
            ("COORDINATES", OPERATIONS[14:]),
        ):
            nav.addWidget(label(category, "section"))
            for operation in operations:
                button = self._track(QPushButton(operation))
                button.setObjectName("Nav")
                button.setIcon(make_icon(operation))
                button.setIconSize(QSize(22, 22))
                button.setCheckable(True)
                button.setToolTip(DESCRIPTIONS[operation])
                button.clicked.connect(
                    lambda checked=False, op=operation: self._select_operation(op)
                )
                nav.addWidget(button)
                self.nav_buttons[operation] = button
            nav.addSpacing(10)
        nav.addStretch()
        self.nav_scroll = scroll_area(navcontent)
        navbox.addWidget(self.nav_scroll, 1)
        navbox.addWidget(
            label("Ctrl+O  Load polygons\nCtrl+Enter  Run", "caption", True)
        )
        self.workspace_row.addWidget(self.sidebar)
        self.splitter = QSplitter(Qt.Orientation.Horizontal)
        self.splitter.setHandleWidth(10)
        self.splitter.setChildrenCollapsible(False)
        self.splitter.splitterMoved.connect(self._remember_split)
        self.settings_panel = self._build_settings()
        self.results_panel = self._build_results()
        self.splitter.addWidget(self.settings_panel)
        self.splitter.addWidget(self.results_panel)
        self.splitter.setStretchFactor(0, 0)
        self.splitter.setStretchFactor(1, 1)
        self.workspace_row.addWidget(self.splitter, 1)
        self.workspace_layout.addLayout(self.workspace_row, 1)
        self.outer.addWidget(workspace, 1)
        footer = QWidget()
        footer_layout = QVBoxLayout(footer)
        footer_layout.setContentsMargins(20, 4, 20, 10)
        footer_layout.setSpacing(5)
        self.progress_bar = QProgressBar()
        self.progress_bar.setTextVisible(False)
        self.progress_bar.hide()
        footer_layout.addWidget(self.progress_bar)
        self.status = ElidedLabel("Ready · Choose an operation and load Layer A.")
        self.status.setProperty("role", "muted")
        footer_layout.addWidget(self.status)
        self.outer.addWidget(footer)

    def _build_settings(self):
        panel = QFrame()
        panel.setObjectName("Panel")
        panel.setMinimumSize(180, 80)
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(0, 0, 0, 0)
        content = QWidget()
        box = QVBoxLayout(content)
        box.setContentsMargins(18, 18, 18, 12)
        box.setSpacing(12)
        self.operation_heading = label("", "heading", True)
        self.operation_description = label("", "muted", True)
        box.addWidget(self.operation_heading)
        box.addWidget(self.operation_description)
        self.inputs_heading = label("01  INPUTS", "section")
        box.addWidget(self.inputs_heading)
        self.a_card = FileCard(
            "LAYER A", "Main polygon layer", lambda: self._browse_layer("A")
        )
        self.b_card = FileCard(
            "LAYER B", "Second polygon layer", lambda: self._browse_layer("B")
        )
        self.point_card = FileCard(
            "POINT TABLE", "CSV or XLSX · first worksheet", self._browse_points
        )
        for card in (self.a_card, self.b_card, self.point_card):
            self._track(card.button)
        box.addWidget(self.a_card)
        self.include_b = self._track(ClearCheckBox("Include Layer B"))
        self.include_b.toggled.connect(
            lambda checked: self._select_operation(self.operation)
        )
        box.addWidget(self.include_b)
        box.addWidget(self.b_card)
        box.addWidget(self.point_card)
        self.site_card = FileCard(
            "REFERENCE SITES", "CSV or XLSX · first worksheet", self._browse_sites
        )
        self.site_card.button.setText("Choose table…")
        self._track(self.site_card.button)
        box.addWidget(self.site_card)
        self.cutter_card = FileCard(
            "CUTTING LINES",
            "LineString input · KML, GeoJSON, or GIS",
            self._browse_cutters,
        )
        self._track(self.cutter_card.button)
        box.addWidget(self.cutter_card)
        self.options_heading = label("02  OPERATION SETTINGS", "section", True)
        box.addWidget(self.options_heading)
        self.point_options = QWidget()
        pointbox = QVBoxLayout(self.point_options)
        pointbox.setContentsMargins(0, 0, 0, 0)
        pointbox.setSpacing(10)
        self.mapping_widget = QWidget()
        self.mapping_grid = QGridLayout(self.mapping_widget)
        self.mapping_grid.setContentsMargins(0, 0, 0, 0)
        self.mapping_grid.setSpacing(10)
        self.mapping_fields = []
        self.column_combos = []
        for name in ("Point name", "Latitude", "Longitude"):
            field = QWidget()
            fieldbox = QVBoxLayout(field)
            fieldbox.setContentsMargins(0, 0, 0, 0)
            fieldbox.setSpacing(6)
            fieldbox.addWidget(label(name))
            combo = self._track(ChoiceBox())
            combo.setAccessibleName(name + " column")
            combo.setSizeAdjustPolicy(
                QComboBox.SizeAdjustPolicy.AdjustToMinimumContentsLengthWithIcon
            )
            combo.setMinimumContentsLength(1)
            combo.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Fixed)
            combo.currentIndexChanged.connect(self._clear_validation)
            fieldbox.addWidget(combo)
            self.mapping_fields.append(field)
            self.column_combos.append(combo)
        pointbox.addWidget(self.mapping_widget)
        self.boundary = self._track(ClearCheckBox("Include boundary points"))
        self.boundary.setChecked(True)
        self.boundary.setToolTip(
            "Includes points on polygon and hole boundaries. Points inside holes remain outside."
        )
        pointbox.addWidget(self.boundary)
        self.swap = self._track(ClearCheckBox("Swap latitude and longitude"))
        self.swap.setToolTip(
            "Explicitly reverse the selected coordinate columns for this run."
        )
        pointbox.addWidget(self.swap)
        pointbox.addWidget(
            label(
                "All input rows are kept. Invalid rows are explained in the report.",
                "muted",
                True,
            )
        )
        box.addWidget(self.point_options)
        self.converter_options = QWidget()
        converter_layout = QVBoxLayout(self.converter_options)
        converter_layout.setContentsMargins(0, 0, 0, 0)
        converter_layout.setSpacing(10)
        converter_layout.addWidget(label("CONVERSION INPUT", "section"))
        self.converter_mode = self._track(ChoiceBox())
        self.converter_mode.addItems(["Single point", "Bulk table"])
        converter_layout.addWidget(self.converter_mode)
        self.converter_single = QWidget()
        single_layout = QVBoxLayout(self.converter_single)
        single_layout.setContentsMargins(0, 0, 0, 0)
        self.converter_entries = []
        for title, hint in (
            ("Point name (optional)", "Point_1"),
            ("Latitude", "30° 12′ 30″ N"),
            ("Longitude", "31.2357"),
        ):
            single_layout.addWidget(label(title))
            entry = self._track(QLineEdit())
            entry.setPlaceholderText(hint)
            entry.setAccessibleName(title)
            single_layout.addWidget(entry)
            self.converter_entries.append(entry)
        converter_layout.addWidget(self.converter_single)
        converter_layout.addWidget(label("Decimal places"))
        self.converter_precision = self._track(ChoiceBox())
        self.converter_precision.addItems(["4", "6", "8"])
        self.converter_precision.setCurrentText("6")
        converter_layout.addWidget(self.converter_precision)
        self.converter_utm = self._track(ClearCheckBox("Include UTM"))
        self.converter_utm.setChecked(True)
        self.converter_typos = self._track(
            ClearCheckBox("Correct O/I/L typos (review Corrections)")
        )
        converter_layout.addWidget(self.converter_utm)
        converter_layout.addWidget(self.converter_typos)
        self.converter_map = self._track(QPushButton("View on Google Maps"))
        self.converter_map.clicked.connect(self._converter_map)
        converter_layout.addWidget(self.converter_map)
        converter_layout.addWidget(
            label(
                "Bulk names are optional. Missing coordinates are listed in Rejected rows.",
                "muted",
                True,
            )
        )
        self.converter_mode.currentTextChanged.connect(
            lambda: self._select_operation(self.operation)
        )
        box.addWidget(self.converter_options)
        self.nearest_options = QWidget()
        nearbox = QVBoxLayout(self.nearest_options)
        nearbox.setContentsMargins(0, 0, 0, 0)
        nearbox.setSpacing(10)
        nearbox.addWidget(label("REFERENCE SITE COLUMNS", "section", True))
        self.site_combos = []
        for title in ("Site ID", "Site latitude", "Site longitude"):
            nearbox.addWidget(label(title))
            combo = self._track(ChoiceBox())
            combo.setAccessibleName(title + " column")
            combo.setSizeAdjustPolicy(
                QComboBox.SizeAdjustPolicy.AdjustToMinimumContentsLengthWithIcon
            )
            combo.setMinimumContentsLength(1)
            combo.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Fixed)
            self.site_combos.append(combo)
            nearbox.addWidget(combo)
        nearbox.addWidget(label("Nearest N sites per location"))
        self.nearest_count = self._track(CountBox())
        self.nearest_count.setRange(1, 1000)
        self.nearest_count.setValue(3)
        self.nearest_count.setAccessibleName("Nearest N")
        nearbox.addWidget(self.nearest_count)
        self.exclude_ids = self._track(
            ClearCheckBox("Exclude reference rows with the same ID")
        )
        nearbox.addWidget(self.exclude_ids)
        self.ignore_zero_distance = self._track(ClearCheckBox("Ignore 0 distance"))
        self.ignore_zero_distance.setToolTip(
            "Skip reference points with exactly zero WGS84 distance, regardless of ID. "
            "Positive distances are kept even when display rounding shows zero."
        )
        nearbox.addWidget(self.ignore_zero_distance)
        box.addWidget(self.nearest_options)
        self.split_options = QWidget()
        splitbox = QVBoxLayout(self.split_options)
        splitbox.setContentsMargins(0, 0, 0, 0)
        splitbox.addWidget(label("Split method"))
        self.split_mode = self._track(ChoiceBox())
        self.split_mode.addItems(("Grid", "Cutting lines"))
        splitbox.addWidget(self.split_mode)
        self.grid_options = QWidget()
        gridbox = QVBoxLayout(self.grid_options)
        gridbox.setContentsMargins(0, 0, 0, 0)
        self.grid_rows = self._track(CountBox())
        self.grid_columns = self._track(CountBox())
        for title, control in (
            ("Grid rows", self.grid_rows),
            ("Grid columns", self.grid_columns),
        ):
            control.setRange(1, 1000)
            control.setValue(2)
            control.setAccessibleName(title)
            gridbox.addWidget(label(title))
            gridbox.addWidget(control)
        gridbox.addWidget(
            label(
                "One local UTM grid covers Layer A. At most 10,000 cells. Rows run south to north; columns west to east. Grid pieces need not have equal area.",
                "muted",
                True,
            )
        )
        splitbox.addWidget(self.grid_options)
        self.split_hint = label(
            "Cutting lines must cross a polygon interior; no automatic snapping or extension is applied.",
            "muted",
            True,
        )
        splitbox.addWidget(self.split_hint)
        box.addWidget(self.split_options)
        self.fill_options = QWidget()
        fillbox = QVBoxLayout(self.fill_options)
        fillbox.setContentsMargins(0, 0, 0, 0)
        self.fill_all = self._track(ClearCheckBox("Fill all holes"))
        self.fill_all.setChecked(True)
        self.hole_limit_label = label("Maximum hole area (m²)")
        self.hole_limit = self._track(QLineEdit("1000"))
        self.hole_limit.setAccessibleName("Maximum hole area in square meters")
        for control in (self.fill_all, self.hole_limit_label, self.hole_limit):
            fillbox.addWidget(control)
        fillbox.addWidget(
            label(
                "Area limits use WGS84 geodesic area. Existing islands inside filled holes merge within each source feature. Other features remain separate.",
                "muted",
                True,
            )
        )
        box.addWidget(self.fill_options)
        self.validation_options = QWidget()
        auditbox = QVBoxLayout(self.validation_options)
        auditbox.setContentsMargins(0, 0, 0, 0)
        self.audit_basis = label(
            "Inspect original parsed records before import cleanup.", "muted", True
        )
        auditbox.addWidget(self.audit_basis)
        self.repair_checkbox = self._track(
            ClearCheckBox("Create repaired polygon copies")
        )
        auditbox.addWidget(self.repair_checkbox)
        auditbox.addWidget(
            label(
                "Report-only by default. Repair retains valid polygons and recovers polygon area where possible. Review changed and excluded records before using the output.",
                "muted",
                True,
            )
        )
        box.addWidget(self.validation_options)
        self.extract_options = QWidget()
        extractbox = QVBoxLayout(self.extract_options)
        extractbox.setContentsMargins(0, 0, 0, 0)
        extractbox.addWidget(label("Output"))
        self.extract_mode = self._track(ChoiceBox())
        self.extract_mode.addItems(("Boundaries", "Vertices"))
        self.extract_holes = self._track(ClearCheckBox("Include hole rings"))
        self.extract_holes.setChecked(True)
        self.closing_vertex = self._track(
            ClearCheckBox("Include the repeated closing vertex")
        )
        for control in (self.extract_mode, self.extract_holes, self.closing_vertex):
            extractbox.addWidget(control)
        extractbox.addWidget(
            label(
                "Export geometry as GeoJSON/KML/KMZ, or the numbered table as Excel. Ring 0 is exterior; hole rings start at 1. Point and line outputs cannot replace polygon Layer A.",
                "muted",
                True,
            )
        )
        box.addWidget(self.extract_options)
        for control in (self.split_mode, self.extract_mode):
            control.currentIndexChanged.connect(
                lambda index: self._select_operation(self.operation)
            )
        self.fill_all.toggled.connect(
            lambda checked: self._select_operation(self.operation)
        )
        self.distance_options = QWidget()
        distancebox = QVBoxLayout(self.distance_options)
        distancebox.setContentsMargins(0, 0, 0, 0)
        self.distance_label = label("Distance")
        distancebox.addWidget(self.distance_label)
        row = QHBoxLayout()
        self.distance = self._track(QLineEdit("50"))
        self.distance.setAccessibleName("Distance or tolerance")
        self.distance.setMinimumWidth(0)
        self.distance.textChanged.connect(self._clear_validation)
        row.addWidget(self.distance, 1)
        self.unit = self._track(ChoiceBox())
        self.unit.addItems(("Meters", "Kilometers"))
        row.addWidget(self.unit)
        distancebox.addLayout(row)
        self.distance_hint = label("", "muted", True)
        distancebox.addWidget(self.distance_hint)
        box.addWidget(self.distance_options)
        self.scope_options = QWidget()
        scopebox = QVBoxLayout(self.scope_options)
        scopebox.setContentsMargins(0, 0, 0, 0)
        scopebox.addWidget(label("Total measurement"))
        self.scope = self._track(ChoiceBox())
        self.scope.addItems(("Merged coverage", "Sum of features"))
        scopebox.addWidget(self.scope)
        scopebox.addWidget(
            label(
                "Merged coverage avoids counting overlaps or shared internal boundaries repeatedly.",
                "muted",
                True,
            )
        )
        box.addWidget(self.scope_options)
        self.validation = label("", None, True)
        self.validation.setObjectName("Error")
        self.validation.hide()
        box.addWidget(self.validation)
        box.addStretch()
        self.input_scroll = scroll_area(content)
        self.input_scroll.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAsNeeded
        )
        self.input_scroll.viewport().installEventFilter(self)
        layout.addWidget(self.input_scroll, 1)
        action = QWidget()
        actions = QHBoxLayout(action)
        actions.setContentsMargins(18, 10, 18, 14)
        self.run_button = self._track(QPushButton("Run operation"))
        self.run_button.setObjectName("Primary")
        self.run_button.clicked.connect(self._run)
        actions.addWidget(self.run_button, 1)
        self.cancel_button = QPushButton("Cancel")
        self.cancel_button.clicked.connect(self._cancel)
        self.cancel_button.hide()
        actions.addWidget(self.cancel_button)
        layout.addWidget(action)
        return panel

    def _build_results(self):
        panel = QFrame()
        panel.setObjectName("Panel")
        panel.setMinimumSize(180, 80)
        panel.installEventFilter(self)
        box = self.results_layout = QVBoxLayout(panel)
        box.setContentsMargins(18, 16, 18, 16)
        box.setSpacing(10)
        head = QHBoxLayout()
        self.results_heading = label("RESULTS", "section")
        head.addWidget(self.results_heading)
        head.addStretch()
        self.open_button = QPushButton("Open export")
        self.open_button.clicked.connect(self._open_export)
        self.promote_button = QPushButton("Use as Layer A")
        self.promote_button.clicked.connect(self._promote)
        self.export_button = QPushButton("Export…")
        self.export_button.clicked.connect(self._export)
        self.more_button = QToolButton()
        self.more_button.setText("More actions")
        self.more_button.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
        menu = QMenu(self)
        self.open_action = menu.addAction("Open export", self._open_export)
        self.promote_action = menu.addAction("Use result as Layer A", self._promote)
        self.more_button.setMenu(menu)
        for button in (
            self.open_button,
            self.promote_button,
            self.more_button,
            self.export_button,
        ):
            head.addWidget(button)
        box.addLayout(head)
        self.result_subtitle = label("Your next result starts here.", "muted", True)
        box.addWidget(self.result_subtitle)
        self.metrics = QWidget()
        self.metrics_layout = QHBoxLayout(self.metrics)
        self.metrics_layout.setContentsMargins(0, 0, 0, 0)
        self.metrics_layout.setSpacing(10)
        self.metric_values = []
        self.metric_titles = []
        for _ in range(3):
            card = QFrame()
            card.setObjectName("Metric")
            cardbox = QVBoxLayout(card)
            cardbox.setContentsMargins(12, 10, 12, 10)
            title = label("", "caption", True)
            value = ElidedLabel("—")
            value.setProperty("role", "metric")
            value.setSizePolicy(
                QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Preferred
            )
            cardbox.addWidget(title)
            cardbox.addWidget(value)
            self.metric_titles.append(title)
            self.metric_values.append(value)
            self.metrics_layout.addWidget(card, 1)
        box.addWidget(self.metrics)
        self.compact_summary = label("", "muted", True)
        box.addWidget(self.compact_summary)
        self.tabs = QTabBar()
        self.tabs.setDrawBase(False)
        self.tabs.setExpanding(False)
        self.tabs.setUsesScrollButtons(True)
        self.tabs.setElideMode(Qt.TextElideMode.ElideRight)
        self.tabs.currentChanged.connect(self._select_result_table)
        box.addWidget(self.tabs)
        self.table_selector = ChoiceBox()
        self.table_selector.setSizePolicy(
            QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Fixed
        )
        self.table_selector.currentIndexChanged.connect(self._select_result_table)
        self.table_selector.setAccessibleName("Result table")
        box.addWidget(self.table_selector)
        self.result_stack = QStackedWidget()
        self.result_stack.setMinimumSize(0, 0)
        placeholder = QWidget()
        empty = QVBoxLayout(placeholder)
        empty.setContentsMargins(16, 12, 16, 12)
        empty.addStretch()
        self.empty_icon = empty_icon = QLabel()
        empty_icon.setPixmap(self.windowIcon().pixmap(64, 64))
        empty_icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        empty.addWidget(empty_icon)
        self.empty_hint = label(
            "Load your inputs, choose the settings,\nand run an operation to inspect the result.",
            "muted",
            True,
        )
        self.empty_hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        empty.addWidget(self.empty_hint)
        empty.addStretch()
        self.result_stack.addWidget(placeholder)
        box.addWidget(self.result_stack, 1)
        self._update_result_actions()
        return panel

    def _schedule_reflow(self):
        if not self._reflow_pending:
            self._reflow_pending = True
            QTimer.singleShot(0, self._reflow)

    def eventFilter(self, watched, event):
        if event.type() == QEvent.Type.Resize:
            self._schedule_reflow()
        return super().eventFilter(watched, event)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, "splitter"):
            self._schedule_reflow()

    def changeEvent(self, event):
        super().changeEvent(event)
        if event.type() in (
            QEvent.Type.FontChange,
            QEvent.Type.ApplicationFontChange,
        ) and hasattr(self, "splitter"):
            self._apply_typography()
            self._schedule_reflow()

    def _apply_typography(self):
        typography(self)
        height = max(32, self.fontMetrics().height() + 14)
        for view in self.table_views:
            view.verticalHeader().setDefaultSectionSize(height)

    def _remember_split(self, *args):
        if not self._restoring_split:
            sizes = self.splitter.sizes()
            if sum(sizes) > 0:
                self.settings["qt_split_side_by_side"] = sizes[0] / sum(sizes)
        self._schedule_reflow()

    def _restore_split_ratio(self):
        ratio = self.settings.get("qt_split_side_by_side", 0.45)
        if not isinstance(ratio, (float, int)) or not math.isfinite(ratio):
            ratio = 0.45
        ratio = max(0.2, min(0.8, ratio))
        self._restoring_split = True
        total = max(1, self.splitter.width() - self.splitter.handleWidth())
        self.splitter.setSizes([round(total * ratio), round(total * (1 - ratio))])
        self._restoring_split = False
        self._schedule_reflow()

    def _reflow(self):
        self._reflow_pending = False
        if not hasattr(self, "results_panel"):
            return
        scale = max(1.0, QApplication.font().pointSizeF() / 10)
        effective = self.width() / scale
        mode = (
            "compact" if effective < 900 else "medium" if effective < 1280 else "wide"
        )
        # Measure the selected (bold) label too; a scrollbar must not steal
        # the space reserved for the final characters of a long operation.
        nav_font = QFont(self.font())
        nav_font.setWeight(QFont.Weight.DemiBold)
        from PySide6.QtGui import QFontMetrics

        label_width = max(
            QFontMetrics(nav_font).horizontalAdvance(op) for op in OPERATIONS
        )
        self.sidebar.setFixedWidth(
            max(round(250 * scale), label_width + round(90 * scale))
        )
        if mode != self.layout_mode:
            self.layout_mode = mode
            self._restoring_split = True
            self.sidebar.setVisible(mode != "compact")

            self.operation_combo.setVisible(mode == "compact")
            self.splitter.setOrientation(Qt.Orientation.Horizontal)
            # Wait for the sidebar's layout change before applying the saved
            # proportion. A temporary narrow layout must not overwrite it.
            QTimer.singleShot(0, self._restore_split_ratio)
            QTimer.singleShot(0, self._ensure_navigation_visible)
        self.brand_title.setText("Geo Toolkit" if effective < 620 else APP_NAME)
        self.brand_subtitle.setVisible(effective >= 620 and self.height() > 500 * scale)
        self.version_badge.setVisible(effective >= 800)
        margins = 10 if mode == "compact" else 16
        self.workspace_layout.setContentsMargins(margins, margins, margins, 8)
        width = self.input_scroll.viewport().width() / scale - 36
        columns = 3 if width >= 650 else 2 if width >= 450 else 1
        if getattr(self, "_mapping_columns", None) != columns:
            self._mapping_columns = columns
            for field in self.mapping_fields:
                self.mapping_grid.removeWidget(field)
            for index, field in enumerate(self.mapping_fields):
                self.mapping_grid.addWidget(field, index // columns, index % columns)
            for index in range(3):
                self.mapping_grid.setColumnStretch(index, 1 if index < columns else 0)
        self.settings_panel.show()
        self.results_panel.show()
        self.run_button.setText(
            "Run" if self.settings_panel.width() / scale < 330 else "Run operation"
        )
        self.empty_icon.setVisible(
            self.results_panel.height() > 260 * scale
        ) if not self.result else None
        result_width = self.results_panel.width() / scale
        narrow_results = result_width < 360
        self.results_heading.setVisible(not narrow_results)
        self.more_button.setText("⋯" if narrow_results else "More actions")
        self.more_button.setToolTip(
            "More actions: open export or use result as Layer A"
        )
        self.export_button.setText("Export" if narrow_results else "Export…")
        compact_actions = result_width < 780
        self.open_button.setVisible(not compact_actions)
        self.promote_button.setVisible(not compact_actions)
        self.more_button.setVisible(compact_actions)
        cards = (
            bool(self.result)
            and result_width >= 560
            and self.results_panel.height() > 350 * scale
        )
        self.metrics.setVisible(cards)
        self.compact_summary.setVisible(
            bool(self.result)
            and not cards
            and self.results_panel.height() > 245 * scale
            and result_width >= 300
        )
        self.result_subtitle.setVisible(
            self.results_panel.height() > 230 * scale and result_width >= 300
        )
        tight = self.height() < 460 * scale or result_width < 360
        self.header_layout.setContentsMargins(
            20, 4 if tight else 14, 20, 4 if tight else 14
        )
        self.workspace_layout.setSpacing(6 if tight else 12)
        self.results_layout.setContentsMargins(
            12 if tight else 18,
            8 if tight else 16,
            12 if tight else 18,
            8 if tight else 16,
        )
        self.results_layout.setSpacing(6 if tight else 10)
        compact_tabs = result_width < 750
        self.tabs.setVisible(bool(self.result) and not compact_tabs)
        self.table_selector.setVisible(bool(self.result) and compact_tabs)
        self._size_table_columns()

    def _ensure_navigation_visible(self):
        if isValid(self) and self.sidebar.isVisible():
            self.nav_scroll.ensureWidgetVisible(self.nav_buttons[self.operation])

    def _select_operation(self, operation):
        if operation not in OPERATIONS:
            return
        self.operation = operation
        self.operation_combo.blockSignals(True)
        self.operation_combo.setCurrentText(operation)
        self.operation_combo.blockSignals(False)
        for name, button in self.nav_buttons.items():
            button.setChecked(name == operation)
        QTimer.singleShot(0, self._ensure_navigation_visible)
        self.operation_heading.setText(operation)
        self.motion.reveal(self.operation_heading)
        self.operation_description.setText(DESCRIPTIONS[operation])
        needs_b = operation in {"Intersection", "Difference", "SymDifference"} or (
            operation == "Combine" and self.include_b.isChecked()
        )
        self.include_b.setVisible(operation == "Combine")
        self.b_card.setVisible(needs_b)
        nearest = operation == "Distance Calculator"
        self.mapping_fields[0].findChild(QLabel).setText(
            "Query ID" if nearest else "Point name"
        )
        converting = operation == "Coordinate Converter"
        bulk_conversion = (
            converting and self.converter_mode.currentText() == "Bulk table"
        )
        self.inputs_heading.setVisible(not converting or bulk_conversion)
        self.options_heading.setText("CONVERSION" if converting else "02  OPERATION SETTINGS")
        if converting:
            self.mapping_fields[0].findChild(QLabel).setText("Point name (optional)")
        self.converter_options.setVisible(converting)
        self.converter_single.setVisible(not bulk_conversion)
        self.a_card.setVisible(not (nearest or converting))
        splitting = operation == "Split polygons"
        lines = self.split_mode.currentText() == "Cutting lines"
        self.split_options.setVisible(splitting)
        self.cutter_card.setVisible(splitting and lines)
        self.grid_options.setVisible(not lines)
        self.split_hint.setVisible(lines)
        self.fill_options.setVisible(operation == "Fill holes")
        self.hole_limit.setEnabled(not self.busy and not self.fill_all.isChecked())
        self.hole_limit_label.setEnabled(not self.fill_all.isChecked())
        self.validation_options.setVisible(operation == "Geometry validation")
        self.extract_options.setVisible(operation == "Boundary / vertices")
        self.closing_vertex.setVisible(self.extract_mode.currentText() == "Vertices")
        layer = self.layers["A"]
        if layer is not None:
            original = (
                layer.raw_features if layer.raw_features is not None else layer.features
            )
            self.audit_basis.setText(
                f"Inspecting {len(original):,} original records; {len(layer.features):,} prepared polygon features. Parser notices are included in the report."
            )
        self.site_card.setVisible(nearest)
        self.nearest_options.setVisible(nearest)
        self.boundary.setVisible(not (nearest or converting))
        self.swap.setVisible(not converting)
        self.point_card.title.setText("QUERY LOCATIONS" if nearest else "POINT TABLE")
        self.point_card.setVisible(
            operation in {"Points inside polygon", "Distance Calculator"}
            or bulk_conversion
        )
        self.point_options.setVisible(
            operation in {"Points inside polygon", "Distance Calculator"}
            or bulk_conversion
        )
        self.distance_options.setVisible(operation in {"Buffer", "Simplify"})
        self.scope_options.setVisible(operation in {"Area", "Perimeter"})
        self.options_heading.setVisible(
            operation
            in {
                "Points inside polygon",
                "Distance Calculator",
                "Coordinate Converter",
                "Split polygons",
                "Fill holes",
                "Geometry validation",
                "Boundary / vertices",
                "Buffer",
                "Simplify",
                "Area",
                "Perimeter",
            }
        )
        self.distance_label.setText(
            "Tolerance" if operation == "Simplify" else "Distance"
        )
        self.distance_hint.setText(
            "Enter a positive tolerance. Larger values remove more detail."
            if operation == "Simplify"
            else "Positive expands; negative shrinks. Distance is applied to each feature."
        )
        self._clear_validation()
        self._schedule_reflow()

    def _clear_validation(self, *args):
        if not hasattr(self, "validation"):
            return
        self.validation.hide()
        for widget in self.column_combos + [self.distance, self.hole_limit]:
            if widget.property("invalid"):
                widget.setProperty("invalid", False)
                widget.style().unpolish(widget)
                widget.style().polish(widget)

    def _invalid(self, message, widgets=()):
        self.validation.setText(message)
        self.validation.show()
        for widget in widgets:
            widget.setProperty("invalid", True)
            widget.style().unpolish(widget)
            widget.style().polish(widget)
        if widgets:
            widgets[0].setFocus()
        self.input_scroll.ensureWidgetVisible(self.validation)
        self.status.setText(message)

    def _set_busy(self, value):
        self.busy = value
        for widget in self.controls:
            widget.setEnabled(not value)
        self.hole_limit.setEnabled(not value and not self.fill_all.isChecked())
        self.cancel_button.setVisible(value)
        self.cancel_button.setEnabled(value)
        self.progress_bar.setVisible(value)
        if value:
            self.progress_bar.setRange(0, 0)
        self._update_result_actions()

    def _update_result_actions(self):
        has_result = self.result is not None
        has_geometry = has_result and bool(self.result.features)
        has_polygons = has_geometry and all(
            f.geometry.geom_type in {"Polygon", "MultiPolygon"}
            for f in self.result.features
        )
        self.export_button.setEnabled(not self.busy and has_result)
        for widget in (self.open_button, self.open_action):
            widget.setEnabled(not self.busy and bool(self.last_export))
        for widget in (self.promote_button, self.promote_action):
            widget.setEnabled(not self.busy and has_polygons)
        self.more_button.setEnabled(
            not self.busy and (has_polygons or bool(self.last_export))
        )

    def _submit(self, title, task, callback):
        if self.busy:
            return
        self.cancel_event.clear()
        self._task_callback = callback
        self._task_started = time.perf_counter()
        self.status.setText(title)
        self._set_busy(True)
        self._thread = QThread(self)
        self._worker = TaskWorker(task, self.cancel_event, self.logger)
        self._worker.moveToThread(self._thread)
        self._thread.started.connect(self._worker.run)
        self._worker.progress.connect(self._on_progress)
        self._worker.succeeded.connect(self._on_success)
        self._worker.failed.connect(self._on_error)
        self._worker.cancelled.connect(self._on_cancelled)
        self._worker.finished.connect(self._thread.quit)
        self._worker.finished.connect(self._worker.deleteLater)
        self._thread.finished.connect(self._thread_done)
        self._thread.finished.connect(self._thread.deleteLater)
        self._thread.start()

    @Slot(float, str)
    def _on_progress(self, value, text):
        if not self.cancel_event.is_set():
            self.status.setText(text)

    @Slot(object)
    def _on_success(self, payload):
        if not self.closing:
            try:
                self._task_callback(payload)
                self.status.setText(
                    f"Complete · {time.perf_counter() - self._task_started:.2f} s · All result rows are available in the table and export."
                )
            except Exception as exc:
                self.logger.exception("Could not display the result")
                self._on_error(str(exc))

    @Slot(str)
    def _on_error(self, message):
        self.status.setText(
            "Task failed. Your previous inputs and result remain available."
        )
        if not self.closing:
            self._show_text_dialog("Task error", message)

    @Slot(str)
    def _on_cancelled(self, message):
        self.status.setText(message)

    @Slot()
    def _thread_done(self):
        self._worker = None
        self._thread = None
        self._set_busy(False)
        if self.closing:
            self.close()

    def _browse_layer(self, key):
        if self.busy:
            return
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            f"Choose Layer {key} polygons",
            self.settings.get("polygon_dir", ""),
            "Polygon files (*.kml *.kmz *.zip *.geojson *.json *.shp *.gpkg)",
        )
        if not paths:
            return

        def loaded(layer):
            self.layers[key] = layer
            (self.a_card if key == "A" else self.b_card).set_loaded(
                paths, len(layer.features)
            )
            self.settings["polygon_dir"] = str(Path(paths[0]).parent)
            self._save_settings()
            self._clear_validation()
            self._select_operation(self.operation)
            if layer.warnings:
                self._show_text_dialog("Input notices", "\n\n".join(layer.warnings))

        self._submit(
            f"Loading Layer {key}…",
            lambda cancel, progress: load_layer(paths, cancel, progress),
            loaded,
        )

    def _browse_cutters(self):
        if self.busy:
            return
        paths, _ = QFileDialog.getOpenFileNames(
            self,
            "Choose cutting lines",
            self.settings.get("polygon_dir", ""),
            "Line files (*.geojson *.json *.kml *.kmz *.zip *.shp *.gpkg)",
        )
        if not paths:
            return

        def loaded(layer):
            self.cutter_layer = layer
            self.cutter_card.set_loaded(paths, len(layer.features), "lines")
            self._clear_validation()
            if layer.warnings:
                self._show_text_dialog(
                    "Cutting line notices", "\n\n".join(layer.warnings)
                )

        self._submit(
            "Reading cutting lines…",
            lambda cancel, progress: load_cutting_lines(paths, cancel, progress),
            loaded,
        )

    def _browse_points(self):
        if self.busy:
            return
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Choose point table",
            self.settings.get("points_dir", ""),
            "Point tables (*.csv *.txt *.tsv *.xlsx)",
        )
        if not path:
            return

        def read(cancel, progress):
            table = read_point_table(path)
            checkpoint(cancel)
            return table

        def loaded(table):
            self.point_table, self.point_path = table, path
            self.point_card.set_loaded([path], len(table), "rows")
            for combo, suggestion in zip(
                self.column_combos, suggest_columns(table.columns)
            ):
                combo.clear()
                combo.addItem("Choose a column…", None)
                for column in table.columns:
                    combo.addItem(str(column), str(column))
                combo.setCurrentIndex(combo.findData(suggestion) if suggestion else 0)
            self.swap.setChecked(False)
            self.settings["points_dir"] = str(Path(path).parent)
            self._save_settings()
            self._clear_validation()

        self._submit("Reading point table…", read, loaded)

    def _browse_sites(self):
        if self.busy:
            return
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Choose reference sites",
            self.settings.get("points_dir", ""),
            "Site tables (*.csv *.xlsx)",
        )
        if not path:
            return

        def loaded(table):
            self.site_table, self.site_path = table, path
            self.site_card.set_loaded([path], len(table), "rows")
            for combo, suggestion in zip(
                self.site_combos, suggest_columns(table.columns)
            ):
                combo.clear()
                combo.addItem("Choose a column…", None)
                for column in table.columns:
                    combo.addItem(str(column), str(column))
                combo.setCurrentIndex(combo.findData(suggestion) if suggestion else 0)
            self.settings["points_dir"] = str(Path(path).parent)
            self._save_settings()
            self._clear_validation()

        def read(cancel, progress):
            table = read_point_table(path)
            checkpoint(cancel)
            return table

        self._submit("Reading reference sites…", read, loaded)

    def _converter_map(self):
        if (
            self.result
            and self.result.operation == "Coordinate Converter"
            and len(self.result.features) == 1
        ):
            point = self.result.features[0].geometry
            QDesktopServices.openUrl(
                QUrl(
                    f"https://www.google.com/maps/search/?api=1&query={point.y},{point.x}"
                )
            )
        else:
            self._invalid("Convert exactly one valid point before opening its map.")

    def _run_converter(self):
        import pandas as pd

        if self.converter_mode.currentText() == "Single point":
            name, latitude, longitude = [
                entry.text() for entry in self.converter_entries
            ]
            table = pd.DataFrame(
                [{"Name": name, "Latitude": latitude, "Longitude": longitude}]
            )
            name_col, lat_col, lon_col = "Name", "Latitude", "Longitude"
        else:
            table = self.point_table
            name_col, lat_col, lon_col = [
                combo.currentData() for combo in self.column_combos
            ]
            if table is None or not lat_col or not lon_col:
                self._invalid(
                    "Load a table and map latitude and longitude; point name is optional."
                )
                return
        precision = int(self.converter_precision.currentText())
        include_utm = self.converter_utm.isChecked()
        correct_typos = self.converter_typos.isChecked()
        self._submit(
            "Converting coordinates…",
            lambda cancel, progress: convert_coordinates(
                table,
                lat_col,
                lon_col,
                name_col,
                precision,
                include_utm,
                correct_typos,
                cancel,
                progress,
            ),
            self._display_result,
        )

    def _run_nearest(self):
        if self.point_table is None or self.site_table is None:
            self._invalid(
                "Load query locations and reference sites. Polygon input is not required."
            )
            return
        query_cols = tuple(c.currentData() for c in self.column_combos)
        site_cols = tuple(c.currentData() for c in self.site_combos)
        if any(
            not all(cols) or len(set(cols)) != 3 for cols in (query_cols, site_cols)
        ):
            self._invalid(
                "Map three distinct ID, latitude, and longitude columns for each table."
            )
            return
        queries, sites = self.point_table, self.site_table
        count, exclude, swap, ignore_zero = (
            self.nearest_count.value(),
            self.exclude_ids.isChecked(),
            self.swap.isChecked(),
            self.ignore_zero_distance.isChecked(),
        )
        self._submit(
            "Finding nearest sites…",
            lambda cancel, progress: nearest_sites(
                queries,
                sites,
                query_cols,
                site_cols,
                count,
                ignore_zero,
                exclude,
                swap,
                cancel,
                progress,
            ),
            self._display_result,
        )

    def _run(self):
        if self.busy:
            return
        self._clear_validation()
        op = self.operation
        if op == "Coordinate Converter":
            self._run_converter()
            return
        if op == "Distance Calculator":
            self._run_nearest()
            return
        a = self.layers["A"]
        needs_b = op in {"Intersection", "Difference", "SymDifference"} or (
            op == "Combine" and self.include_b.isChecked()
        )
        b = self.layers["B"] if needs_b else None
        if op == "Split polygons" and self.split_mode.currentText() == "Cutting lines":
            b = self.cutter_layer
            if b is None:
                self._invalid("Choose a cutting-line file.")
                return
        if a is None or (needs_b and b is None):
            self._invalid(
                "Load Layer A and Layer B."
                if needs_b
                else "Load polygons into Layer A."
            )
            return
        names = [combo.currentData() for combo in self.column_combos]
        if op == "Points inside polygon":
            if self.point_table is None:
                self._invalid("Choose a point table, then map its three columns.")
                return
            if not all(names) or len(set(names)) != 3:
                self._invalid(
                    "Choose three different columns for point name, latitude, and longitude.",
                    self.column_combos,
                )
                return
        try:
            distance = (
                float(self.distance.text()) if op in {"Buffer", "Simplify"} else 0.0
            )
            if not math.isfinite(distance) or (op == "Simplify" and distance <= 0):
                raise ValueError
        except ValueError:
            self._invalid(
                "Enter a finite number. Simplify requires a positive tolerance.",
                [self.distance],
            )
            return
        try:
            hole_limit = (
                float(self.hole_limit.text())
                if op == "Fill holes" and not self.fill_all.isChecked()
                else 0.0
            )
            if not math.isfinite(hole_limit) or hole_limit < 0:
                raise ValueError
        except ValueError:
            self._invalid(
                "Enter a finite, nonnegative hole area in square meters.",
                [self.hole_limit],
            )
            return
        if (
            op == "Split polygons"
            and self.split_mode.currentText() == "Grid"
            and self.grid_rows.value() * self.grid_columns.value() > 10000
        ):
            self._invalid("Choose a grid with at most 10,000 cells.")
            return
        options = Options(
            operation=op,
            split_mode=self.split_mode.currentText(),
            grid_rows=self.grid_rows.value(),
            grid_columns=self.grid_columns.value(),
            fill_all_holes=self.fill_all.isChecked(),
            hole_area_limit=hole_limit,
            repair_geometry=self.repair_checkbox.isChecked(),
            boundary_output=self.extract_mode.currentText(),
            include_holes=self.extract_holes.isChecked(),
            include_closing_vertex=self.closing_vertex.isChecked(),
            distance=distance,
            unit=self.unit.currentText(),
            include_boundary=self.boundary.isChecked(),
            swap_coordinates=self.swap.isChecked(),
            name_column=names[0] or "",
            latitude_column=names[1] or "",
            longitude_column=names[2] or "",
            measurement_scope=self.scope.currentText(),
        )
        table = self.point_table
        self._submit(
            f"Running {op}…",
            lambda cancel, progress: run_operation(
                a, b, options, table, cancel, progress
            ),
            self._display_result,
        )

    def _display_result(self, result):
        self.result = result
        self.last_export = None
        while self.result_stack.count():
            page = self.result_stack.widget(0)
            self.result_stack.removeWidget(page)
            page.deleteLater()
        self.table_views = []
        self.tabs.blockSignals(True)
        self.table_selector.blockSignals(True)
        while self.tabs.count():
            self.tabs.removeTab(0)
        self.table_selector.clear()
        tables = {
            "Summary": [
                {"Item": key, "Value": value} for key, value in result.summary.items()
            ],
            **result.tables,
        }
        if result.warnings:
            tables["Warnings"] = [{"Message": text} for text in result.warnings]
        for title, records in tables.items():
            view = QTableView()
            view.setModel(RecordModel(records, view))
            if result.operation == "Coordinate Converter":
                view.model().coordinate_precision = result.summary["Decimal places"]
            view.setAlternatingRowColors(True)
            view.setShowGrid(False)
            view.setWordWrap(False)
            view.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
            view.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            view.verticalHeader().hide()
            view.horizontalHeader().setSectionResizeMode(
                QHeaderView.ResizeMode.Interactive
            )
            view.horizontalHeader().setMinimumSectionSize(60)
            view._manual_columns = set()
            view._setting_columns = False
            view.horizontalHeader().sectionResized.connect(
                lambda index, old, new, table=view: (
                    table._manual_columns.add(index)
                    if not table._setting_columns
                    else None
                )
            )
            view.doubleClicked.connect(
                lambda index, table=view: self._show_row_details(table, index.row())
            )
            view.setToolTip(
                "Hover for the full value. Double-click a row to open its details."
            )
            self.table_views.append(view)
            self.result_stack.addWidget(view)
            self.tabs.addTab(title)
            self.table_selector.addItem(title)
        self.tabs.blockSignals(False)
        self.table_selector.blockSignals(False)
        self._select_result_table(1 if len(tables) > 1 else 0)
        if result.operation == "Points inside polygon":
            metrics = [
                ("Matched", result.summary["Matched rows"]),
                ("Outside", result.summary["Outside rows"]),
                ("Rejected", result.summary["Rejected rows"]),
            ]
        elif result.operation == "Coordinate Converter":
            metrics = [
                ("Converted", result.summary["Converted rows"]),
                ("Rejected", result.summary["Rejected rows"]),
                ("Corrected", result.summary["Corrected rows"]),
            ]
        elif result.operation == "Geometry validation":
            metrics = [
                ("Valid", result.summary["Valid original features"]),
                ("Review", result.summary["Features needing review"]),
                ("Output copies", len(result.features)),
            ]
        elif result.operation == "Fill holes":
            metrics = [
                ("Holes filled", result.summary["Holes filled"]),
                ("Inspected", result.summary["Holes inspected"]),
                ("Features", len(result.features)),
            ]
        elif result.operation == "Distance Calculator":
            metrics = [
                ("Matched queries", result.summary["Matched queries"]),
                ("Unmatched", result.summary["Unmatched queries"]),
                ("Result pairs", result.summary["Result pairs"]),
            ]
        elif result.operation in {"Area", "Perimeter"}:
            key = (
                "Total area (m²)"
                if result.operation == "Area"
                else "Total perimeter (m)"
            )
            metrics = [
                (key, f"{result.summary[key]:,.3f}"),
                ("Scope", result.summary["Scope"]),
                ("Features", result.summary["Input A features"]),
            ]
        else:
            metrics = [
                ("Result features", len(result.features)),
                ("Input A", result.summary.get("Input A features", 0)),
                ("Notices", len(result.warnings)),
            ]
        for title, value, pair in zip(self.metric_titles, self.metric_values, metrics):
            title.setText(pair[0])
            text = f"{pair[1]:,}" if isinstance(pair[1], int) else str(pair[1])
            value.setText(text)
            value.setProperty("role", "heading" if pair[0] == "Scope" else "metric")
            value.setToolTip(text)
        self.compact_summary.setText(
            " · ".join(
                f"{name}: {value:,}" if isinstance(value, int) else f"{name}: {value}"
                for name, value in metrics
            )
        )
        self.result_subtitle.setText(
            f"{result.operation} · Double-click a row for full details."
        )
        self._apply_typography()
        self._update_result_actions()
        self._schedule_reflow()
        QTimer.singleShot(0, self._reveal_result_summary)

    def _reveal_result_summary(self):
        if not self.closing and self.isVisible():
            self.motion.reveal(
                self.metrics if self.metrics.isVisible() else self.compact_summary
            )

    def _select_result_table(self, index):
        if index < 0 or index >= self.result_stack.count():
            return
        self.tabs.blockSignals(True)
        self.table_selector.blockSignals(True)
        self.tabs.setCurrentIndex(index)
        self.table_selector.setCurrentIndex(index)
        self.tabs.blockSignals(False)
        self.table_selector.blockSignals(False)
        self.result_stack.setCurrentIndex(index)
        self.motion.reveal(self.result_subtitle)
        self._size_table_columns()

    def _size_table_columns(self):
        for view in self.table_views:
            model = view.model()
            width = view.viewport().width()
            if getattr(view, "_sized_for", None) == width:
                continue
            view._sized_for = width
            fm = view.fontMetrics()
            view._setting_columns = True
            for index, header in enumerate(model.headers):
                if index in view._manual_columns:
                    continue
                sample = [str(row.get(header, "")) for row in model.records[:100]]
                natural = max(
                    [fm.horizontalAdvance(header) + 30]
                    + [fm.horizontalAdvance(value[:120]) + 24 for value in sample]
                )
                if header in {
                    "Input row",
                    "Reference row",
                    "Rank",
                    "Feature",
                    "Match count",
                    "Point rows",
                }:
                    target = min(140, max(85, natural))
                elif header in {"Latitude", "Longitude"}:
                    target = max(140, fm.horizontalAdvance("-180.000000") + 30)
                else:
                    target = min(
                        max(180, int(width * 0.45)), min(380, max(160, natural))
                    )
                view.setColumnWidth(index, target)
            view._setting_columns = False

    def _show_row_details(self, view, row):
        records = view.model().records
        if 0 <= row < len(records):
            self._show_text_dialog(
                "Row details",
                "\n\n".join(f"{name}\n{value}" for name, value in records[row].items()),
            )

    def _promote(self):
        if (
            not self.busy
            and self.result
            and self.result.features
            and all(
                f.geometry.geom_type in {"Polygon", "MultiPolygon"}
                for f in self.result.features
            )
        ):
            self.layers["A"] = Layer(
                list(self.result.features),
                ("Previous result",),
                list(self.result.warnings),
            )
            self.a_card.set_loaded(["Previous result"], len(self.result.features))
            self.status.setText(
                "Result copied into Layer A. Choose the next operation."
            )

    def _export(self):
        if self.busy or self.result is None:
            return
        filters = "Excel report (*.xlsx)"
        suffix = ".xlsx"
        if self.result.features:
            filters = "GeoJSON (*.geojson);;Google Earth KML (*.kml);;Google Earth KMZ (*.kmz);;Excel report (*.xlsx)"
            suffix = ".geojson"
        initial = str(
            Path(self.settings.get("output_dir", "."))
            / (self.result.operation.replace(" ", "_") + suffix)
        )
        path, selected = QFileDialog.getSaveFileName(
            self, "Export result", initial, filters
        )
        if not path:
            return
        expected = {
            "GeoJSON (*.geojson)": ".geojson",
            "Google Earth KML (*.kml)": ".kml",
            "Google Earth KMZ (*.kmz)": ".kmz",
            "Excel report (*.xlsx)": ".xlsx",
        }.get(selected, suffix)
        original_path = path
        if Path(path).suffix.lower() != expected:
            path = str(Path(path).with_suffix(expected))
        sources = [
            source
            for layer in self.layers.values()
            if layer
            for source in layer.sources
        ]
        if self.point_path:
            sources.append(self.point_path)
        if self.site_path:
            sources.append(self.site_path)
        if self.cutter_layer:
            sources.extend(self.cutter_layer.sources)
        if any(Path(path).resolve() == Path(source).resolve() for source in sources):
            self._show_text_dialog(
                "Choose another filename",
                "Use a new output filename to keep your loaded source files intact.",
            )
            return
        if path != original_path and Path(path).exists():
            answer = QMessageBox.question(
                self,
                "Replace existing export?",
                f"The selected format uses {expected}. Replace this existing file?\n\n{path}",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No,
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
        result = self.result

        def save(cancel, progress):
            export_result(result, path, cancel)
            return path

        def saved(destination):
            self.last_export = destination
            self.settings["output_dir"] = str(Path(destination).parent)
            self._save_settings()

        self._submit("Saving result…", save, saved)

    def _open_export(self):
        if not self.busy and self.last_export:
            if not QDesktopServices.openUrl(
                QUrl.fromLocalFile(str(Path(self.last_export).resolve()))
            ):
                self._show_text_dialog(
                    "Open export",
                    f"Could not open the default application. Your export is saved at:\n\n{self.last_export}",
                )

    def _cancel(self):
        self.cancel_event.set()
        self.cancel_button.setEnabled(False)
        self.status.setText(
            "Cancelling at the next safe checkpoint. A running geometry calculation must finish first."
        )

    def _dialog(self, key, factory):
        existing = self.help_windows.get(key)
        if existing is not None:
            existing.showNormal()
            existing.raise_()
            existing.activateWindow()
            return None
        window = factory()
        QTimer.singleShot(
            0,
            lambda: self.motion.reveal(window.title_label) if isValid(window) else None,
        )
        self.help_windows[key] = window
        window.destroyed.connect(
            lambda obj=None, name=key: self.help_windows.pop(name, None)
        )
        return window

    def _show_about(self):
        window = self._dialog(
            "about",
            lambda: ProductDialog(
                self,
                "About",
                "Polygon analysis and point matching in one desktop tool.",
                610,
                560,
            ),
        )
        if window is None:
            return
        window.body_layout.addWidget(label(APP_NAME, "heading", True))
        for name, value in (
            ("Version", __version__),
            ("Build", BUILD_NUMBER),
            ("Released", RELEASE_DATE),
            ("Author", AUTHOR_NAME),
            ("Distribution", "exe Desktop application"),
        ):
            card = QFrame()
            card.setObjectName("FileCard")
            row = QHBoxLayout(card)
            row.setContentsMargins(12, 8, 12, 8)
            field_label = label(name, "caption")
            field_label.setMinimumWidth(100)
            row.addWidget(field_label)
            row.addWidget(label(value, None, True), 1)
            window.body_layout.addWidget(card)
        for title, callback in (
            ("Copy version details", lambda: window.copy_text(support_details())),
            ("Quick Guide", self._show_quick_guide),
            ("Contact & Support", self._show_support),
        ):
            button = QPushButton(title.replace("&", "&&"))
            button.clicked.connect(callback)
            window.body_layout.addWidget(button)
        window.body_layout.setSpacing(10)
        window.body_layout.setAlignment(Qt.AlignmentFlag.AlignTop)
        typography(window)
        window.show()
        QTimer.singleShot(0, window.fit_content_height)

    def _show_quick_guide(self):
        window = self._dialog("guide", lambda: GuideDialog(self))
        if window:
            window.show()

    def _show_support(self):
        from urllib.parse import quote

        window = self._dialog(
            "support",
            lambda: ProductDialog(
                self,
                "Contact & Support",
                "Questions, bug reports, and suggestions.",
                610,
                660,
            ),
        )
        if window is None:
            return
        window.body_layout.addWidget(
            label(f"Developer · {AUTHOR_NAME}", "heading", True)
        )
        email = f"mailto:{SUPPORT_EMAIL}?subject={quote(APP_NAME + ' ' + __version__ + ' — Support request')}"
        for title, value, url in (
            ("Email", SUPPORT_EMAIL, email),
            ("LinkedIn", LINKEDIN_URL, LINKEDIN_URL),
        ):
            card = QFrame()
            card.setObjectName("FileCard")
            box = QVBoxLayout(card)
            box.setContentsMargins(14, 12, 14, 12)
            box.addWidget(label(title, "section"))
            address = label(
                f'<a style="color:#176D7B" href="{escape(url, quote=True)}">{escape(value)}</a>',
                None,
                True,
            )
            address.setProperty("linkHtml", address.text())
            if self.theme == "dark":
                address.setText(address.text().replace("#176D7B", "#6ce0de"))
            address.setTextInteractionFlags(
                Qt.TextInteractionFlag.TextBrowserInteraction
            )
            address.linkActivated.connect(window.open_link)
            box.addWidget(address)
            button = QPushButton(
                "Write an email" if title == "Email" else "Open LinkedIn"
            )
            button.clicked.connect(
                lambda checked=False, link=url: window.open_link(link)
            )
            actions = QHBoxLayout()
            actions.addWidget(button, 1)
            copy = QPushButton(f"Copy {title}")
            copy.clicked.connect(
                lambda checked=False, text=value: window.copy_text(text)
            )
            actions.addWidget(copy, 1)
            box.addLayout(actions)
            window.body_layout.addWidget(card)
        window.body_layout.addWidget(
            label("For a useful support request", "heading", True)
        )
        window.body_layout.addWidget(
            label(
                "Include the operation, input formats, steps to reproduce, expected result, and actual result or error message.",
                "muted",
                True,
            )
        )
        copy = QPushButton("Copy support details")
        copy.clicked.connect(
            lambda: window.copy_text(
                support_details(), "Version and environment details copied."
            )
        )
        window.body_layout.addWidget(copy)
        window.body_layout.setSpacing(10)
        window.body_layout.setAlignment(Qt.AlignmentFlag.AlignTop)
        typography(window)
        window.show()
        QTimer.singleShot(0, window.fit_content_height)

    def _show_text_dialog(self, title, text):
        window = ProductDialog(self, title, APP_NAME, 700, 500)
        reader = QTextBrowser()
        reader.setPlainText(text)
        reader.setMinimumHeight(160)
        window.body_layout.addWidget(reader, 1)
        typography(window)
        window.show()
        return window

    def closeEvent(self, event):
        if self.busy:
            self.closing = True
            self._cancel()
            self.status.setText("Finishing the current task safely, then closing…")
            event.ignore()
            return
        self.motion.stop_all()
        self._remember_split()
        rect = self.normalGeometry() if self.isMaximized() else self.geometry()
        self.settings["qt_window_rect"] = [
            rect.x(),
            rect.y(),
            rect.width(),
            rect.height(),
        ]
        self._save_settings()
        for window in list(self.help_windows.values()):
            window.close()
        for handler in self.logger.handlers[:]:
            handler.close()
            self.logger.removeHandler(handler)
        event.accept()


def create_application(argv=None):
    app = QApplication.instance() or QApplication(
        argv if argv is not None else sys.argv
    )
    app.setApplicationName(APP_NAME)
    app.setOrganizationName("Mohamed Ashraf")
    app.setApplicationVersion(__version__)
    app.setStyle("Fusion")
    apply_theme("light")
    return app


def main():
    if sys.platform == "win32":
        try:
            import ctypes

            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
                "GeoOperationsToolkit.Desktop.2"
            )
        except (AttributeError, OSError):
            pass
    app = create_application()
    window = ToolkitApp()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
